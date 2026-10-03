"""canon-gate: check pipetting configs / Params sheets against the canonical-constants registry.

WHY: a decision about an experimental constant (a cofactor stock of 10 mM with a final 0.015 mM, an
enzyme batch label) was made but never propagated -- an older workbook's Params sheet still said
10 uM / 0.015 uM and a new config still named the previous batch. Nobody noticed until an independent
reviewer diffed the workbooks. This gate makes that diff mechanical.

Contract (shared by the toolkit's mechanical gates):
    exit 0  clean (OK; WARN lines may be present -- they are "undecided, please confirm", not failures)
    exit 1  at least one FAIL (a value contradicts a canonical entry)
    exit 2  cannot check (unreadable file / unreadable registry / no registry field found) = BLIND,
            never a pass
`--strict` turns WARN into exit 1.

Inputs
    *.json   reaction_matrix.py config: stocks[name].conc_mM / final_mM, enzymes[name].batch / stock_gL
    *.xlsx   every sheet whose name contains "param": rows are (label, value, unit, note) and a
             header row ("Component", "Stock conc"|"Final"|"Target", ...) switches the section;
             "Setting/Value" rows use the label suffix _STOCK / _EXISTING_STOCK / _FINAL.
             _WORKING_STOCK rows are derived dilutions and are ignored.

Registry lookup (first hit wins; --registry PATH / registry= overrides all of them):
    1. the file named by the environment variable LAB_CANON_REGISTRY (set but missing = BLIND)
    2. canonical_constants.toml next to the checked file / config
    3. canonical_constants.toml next to this script (the lab's own registry)
    4. canonical_constants.example.toml next to this script (synthetic example entries)
    none found -> BLIND, exit 2.

CLI:  python canon_gate.py FILE [FILE ...] [--registry PATH] [--json] [--strict]
API:  check_file(path, registry=None) -> GateResult ;  check_config(dict, ...) -> GateResult
"""
from __future__ import annotations

import json
import math
import os
import re
import sys
import tomllib
from dataclasses import dataclass, field
from pathlib import Path

REGISTRY_NAME = "canonical_constants.toml"
ENV_REGISTRY = "LAB_CANON_REGISTRY"
DEFAULT_REGISTRY = Path(__file__).with_name(REGISTRY_NAME)
EXAMPLE_REGISTRY = Path(__file__).with_name("canonical_constants.example.toml")

_UNIT_TO_MM = {"m": 1000.0, "mm": 1.0, "um": 1e-3, "µm": 1e-3, "μm": 1e-3, "nm": 1e-6}
RTOL = 1e-6


@dataclass
class Finding:
    level: str  # FAIL | WARN | OK
    entry: str  # registry id
    where: str  # "config stocks['CofactorX'].conc_mM" or "Params!B54"
    found: str
    canonical: str
    message: str

    def line(self) -> str:
        return f"{self.level} {self.entry}: {self.message} [{self.where}]"


@dataclass
class Observation:
    name: str  # name as written in the file
    field: str  # stock_conc | final_conc | batch | stock_gL
    value: object  # float (mM or g/L) or str (batch)
    where: str
    raw: str  # original value+unit text for messages


@dataclass
class GateResult:
    source: str
    findings: list[Finding] = field(default_factory=list)
    blind_reason: str = ""

    @property
    def n_fail(self) -> int:
        return sum(f.level == "FAIL" for f in self.findings)

    @property
    def n_warn(self) -> int:
        return sum(f.level == "WARN" for f in self.findings)

    @property
    def n_ok(self) -> int:
        return sum(f.level == "OK" for f in self.findings)

    def exit_code(self, strict: bool = False) -> int:
        if self.blind_reason:
            return 2
        if self.n_fail or (strict and self.n_warn):
            return 1
        return 0

    @property
    def status(self) -> str:
        if self.blind_reason:
            return "BLIND"
        if self.n_fail:
            return "FAIL"
        return "OK" if not self.n_warn else "OK+WARN"

    def summary(self) -> str:
        if self.blind_reason:
            return f"canon-gate {self.source}: BLIND (exit 2, NOT a pass): {self.blind_reason}"
        return (f"canon-gate {self.source}: {self.status} "
                f"({self.n_fail} FAIL, {self.n_warn} WARN, {self.n_ok} OK)")


# ------------------------------------------------------------------ registry
class RegistryError(Exception):
    pass


def resolve_registry(near: Path | str | None = None) -> Path:
    """The registry file by the lookup order in the module docstring; RegistryError when none exists."""
    env = os.environ.get(ENV_REGISTRY, "").strip()
    if env:
        p = Path(env).expanduser()
        if not p.is_file():
            raise RegistryError(f"{ENV_REGISTRY}={env} does not name a readable file")
        return p
    tried = []
    if near is not None:
        d = Path(near)
        d = d if d.is_dir() else d.parent
        tried.append(d / REGISTRY_NAME)
    tried += [DEFAULT_REGISTRY, EXAMPLE_REGISTRY]
    for p in tried:
        if p.is_file():
            return p
    raise RegistryError("no canonical-constants registry found (looked at "
                        + ", ".join(str(p) for p in tried) + f"; or set {ENV_REGISTRY})")


def load_registry(path: Path | str | None = None, near: Path | str | None = None) -> dict:
    p = Path(path) if path else resolve_registry(near)
    try:
        with open(p, "rb") as fh:
            reg = tomllib.load(fh)
    except (OSError, tomllib.TOMLDecodeError) as exc:
        raise RegistryError(f"cannot read registry {p}: {exc}") from exc
    if not reg.get("canon"):
        raise RegistryError(f"registry {p} declares no [[canon]] entries (empty declaration)")
    return reg


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9,]+", "", str(s).lower())


def _name_matches(entry: dict, name: str) -> bool:
    n = _norm(name)
    return any(n == _norm(x) for x in entry.get("names", []))


def to_mM(value: float, unit: str | None) -> float | None:
    key = (unit or "").strip().lower().replace("mol/l", "m").replace(" ", "")
    if key not in _UNIT_TO_MM:
        return None
    return float(value) * _UNIT_TO_MM[key]


def _close(a: float, b: float) -> bool:
    return math.isclose(a, b, rel_tol=RTOL, abs_tol=0.0)


def _fmt(v) -> str:
    return f"{v:g}" if isinstance(v, float) else str(v)


# ------------------------------------------------------------------ extraction
def observations_from_config(cfg: dict) -> list[Observation]:
    obs: list[Observation] = []
    for name, info in (cfg.get("stocks") or {}).items():
        if not isinstance(info, dict):
            continue
        if isinstance(info.get("conc_mM"), (int, float)):
            obs.append(Observation(name, "stock_conc", float(info["conc_mM"]),
                                   f"config stocks['{name}'].conc_mM", f"{info['conc_mM']} mM"))
        if isinstance(info.get("final_mM"), (int, float)):
            obs.append(Observation(name, "final_conc", float(info["final_mM"]),
                                   f"config stocks['{name}'].final_mM", f"{info['final_mM']} mM"))
    for name, info in (cfg.get("enzymes") or {}).items():
        if not isinstance(info, dict):
            continue
        # A diluted working stock is pipetted at `stock_gL`; the canonical check applies to the
        # UNDILUTED lab stock, given as `source_stock_gL` when the config uses a dilution.
        _key = "source_stock_gL" if isinstance(info.get("source_stock_gL"), (int, float)) else "stock_gL"
        name = info.get("enzyme", name)   # pseudo-enzymes from equal_volume.py carry the real name
        if isinstance(info.get(_key), (int, float)):
            obs.append(Observation(name, "stock_gL", float(info[_key]),
                                   f"config enzymes['{name}'].{_key}", f"{info[_key]} g/L"))
        batch = info.get("batch")
        if isinstance(batch, float) and batch.is_integer():
            batch = int(batch)
        if isinstance(batch, (str, int)) and not isinstance(batch, bool):
            text = str(batch).strip()          # 100001 (int) and ' 100001' are the batch 100001
            obs.append(Observation(name, "batch", text, f"config enzymes['{name}'].batch", text))
    return obs


_SUFFIX = re.compile(r"^(?P<name>.+?)[_ ]+(?P<kind>EXISTING_STOCK|WORKING_STOCK|STOCK|FINAL)$", re.I)
_BATCH_IN_NOTE = re.compile(r"batch\s*(?:no\.?\s*)?(\d{6})", re.I)


def observations_from_xlsx(path: Path) -> tuple[list[Observation], list[str]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets = [n for n in wb.sheetnames if "param" in n.lower()]
    obs: list[Observation] = []
    for sn in sheets:
        ws = wb[sn]
        section = "stock"
        for row in ws.iter_rows():
            cells = [c.value for c in row[:4]] + [None] * (4 - len(row[:4]))
            label, val, unit, note = cells
            if label is None:
                continue
            ltxt = str(label).strip()
            vtxt = str(val).strip().lower() if isinstance(val, str) else ""
            if vtxt in ("stock conc", "stock", "stock concentration"):
                section = "stock"
                continue
            if vtxt in ("final", "target", "final conc"):
                section = "final"
                continue
            if vtxt == "value":  # Setting/Value header: kind comes from label suffix
                section = "setting"
                continue
            if not isinstance(val, (int, float)) or isinstance(val, bool):
                continue
            where = f"{sn}!B{row[0].row}"
            kind = None
            name = ltxt
            m = _SUFFIX.match(ltxt) if section == "setting" or _SUFFIX.match(ltxt) else None
            if m:
                name, k = m.group("name"), m.group("kind").upper()
                if k == "WORKING_STOCK":
                    continue
                kind = "stock" if k in ("STOCK", "EXISTING_STOCK") else "final"
            else:
                if section == "setting":
                    continue
                kind = section
            u = (unit or "").strip() if isinstance(unit, str) else ""
            raw = f"{val:g} {u}".strip()
            if u.lower() in ("g/l",):
                if kind == "stock":
                    obs.append(Observation(name, "stock_gL", float(val), where, raw))
                    bm = _BATCH_IN_NOTE.search(str(note or ""))
                    if bm:
                        obs.append(Observation(name, "batch", bm.group(1),
                                               f"{sn}!D{row[0].row} (note)", bm.group(1)))
                continue
            mm = to_mM(val, u)
            if mm is None:
                continue
            obs.append(Observation(name, "stock_conc" if kind == "stock" else "final_conc",
                                   mm, where, raw))
    return obs, sheets


# ------------------------------------------------------------------ comparison
def _superseded_hit(entry: dict, found) -> str:
    for s in entry.get("supersedes", []):
        if s.get("kind") != "file":
            continue
        sv = s.get("value")
        if isinstance(found, str) or isinstance(sv, str):
            if str(sv) in str(found):
                return f" -- matches SUPERSEDED value ({s.get('where')})"
            continue
        smm = to_mM(sv, s.get("unit")) if entry["field"] in ("stock_conc", "final_conc") else float(sv)
        if smm is not None and _close(float(found), smm):
            return f" -- matches SUPERSEDED value {_fmt(sv)} {s.get('unit', '')} ({s.get('where')})"
    return ""


def compare(observations: list[Observation], reg: dict, source_name: str) -> list[Finding]:
    findings: list[Finding] = []
    for o in observations:
        for e in reg["canon"]:
            if e["field"] != o.field or not _name_matches(e, o.name):
                continue
            canon = e["value"]
            cdisp = f"{_fmt(canon)} {e.get('unit', '')}".strip()
            if o.field == "batch":
                if str(canon) in str(o.value):
                    findings.append(Finding("OK", e["id"], o.where, o.raw, cdisp, f"{o.name} batch {o.value}"))
                elif re.search(r"\d{6}", str(o.value)):
                    findings.append(Finding(
                        "FAIL", e["id"], o.where, o.raw, cdisp,
                        f"{o.name} batch is {o.value}, canonical {canon} "
                        f"(decided {e['decided']} by {e['decided_by']}){_superseded_hit(e, o.value)}"))
                else:
                    findings.append(Finding(
                        "WARN", e["id"], o.where, o.raw, cdisp,
                        f"{o.name} batch label '{o.value}' carries no 6-digit batch; cannot compare to {canon}"))
                continue
            canon_cmp = float(canon) if o.field == "stock_gL" else to_mM(canon, e["unit"])
            if _close(float(o.value), canon_cmp):
                findings.append(Finding("OK", e["id"], o.where, o.raw, cdisp, f"{o.name} {o.field} = {o.raw}"))
            else:
                findings.append(Finding(
                    "FAIL", e["id"], o.where, o.raw, cdisp,
                    f"{o.name} {o.field} is {o.raw}, canonical {cdisp} "
                    f"(decided {e['decided']} by {e['decided_by']}){_superseded_hit(e, o.value)}"))
    for u in reg.get("undecided", []):
        for o in observations:
            if o.field != u["field"] or not _name_matches(u, o.name):
                continue
            others = [x for x in u.get("observed", []) if x["file"] != source_name]
            diff = [x for x in others if not _close(float(o.value), float(x["value"]))]
            for x in diff:
                findings.append(Finding(
                    "WARN", u["id"], o.where, o.raw, f"{_fmt(x['value'])} {u['unit']} in {x['file']}",
                    f"{o.name} {o.field} {o.raw} differs from {_fmt(x['value'])} {u['unit']} in {x['file']} "
                    f"(working stock, NOT canonical -- confirm the dilution is intentional)"))
    return findings


# ------------------------------------------------------------------ public API
def _finish(res: GateResult, obs: list[Observation], reg: dict, name: str, what: str) -> GateResult:
    res.findings = compare(obs, reg, name)
    canon_hits = [f for f in res.findings if f.entry in {e["id"] for e in reg["canon"]}]
    if not canon_hits:
        res.blind_reason = (f"no registry field found in {what} (looked for "
                            f"{sorted({n for e in reg['canon'] for n in e['names']})[:6]}...); "
                            f"nothing was checked")
    return res


def check_config(cfg: dict, registry: dict | None = None, source: str = "<config>") -> GateResult:
    """source = the config's path when known: its folder is searched for a registry (lookup step 2)."""
    res = GateResult(source)
    near = Path(source).parent if source and not source.startswith("<") else None
    try:
        reg = registry or load_registry(near=near)
    except RegistryError as exc:
        res.blind_reason = str(exc)
        return res
    return _finish(res, observations_from_config(cfg), reg, Path(source).name, "config stocks/enzymes")


def check_file(path: Path | str, registry: dict | None = None) -> GateResult:
    p = Path(path)
    res = GateResult(str(p))
    try:
        reg = registry or load_registry(near=p.parent)
    except RegistryError as exc:
        res.blind_reason = str(exc)
        return res
    try:
        if p.suffix.lower() == ".json":
            with open(p, encoding="utf-8") as fh:
                obs = observations_from_config(json.load(fh))
            what = "config stocks/enzymes"
        elif p.suffix.lower() in (".xlsx", ".xlsm"):
            obs, sheets = observations_from_xlsx(p)
            if not sheets:
                res.blind_reason = "workbook has no Params-like sheet (name containing 'param')"
                return res
            what = f"sheet(s) {sheets}"
        else:
            res.blind_reason = f"unsupported file type {p.suffix!r}"
            return res
    except Exception as exc:  # unreadable / corrupt / not a dict
        res.blind_reason = f"cannot read {p.name}: {type(exc).__name__}: {exc}"
        return res
    return _finish(res, obs, reg, p.name, what)


def cli(argv: list[str] | None = None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description="canon-gate: check files against canonical_constants.toml")
    ap.add_argument("files", nargs="+")
    ap.add_argument("--registry", help=f"registry TOML (default: ${ENV_REGISTRY}, then {REGISTRY_NAME} next to "
                                       f"each file, then the skill's own, then the example)")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--strict", action="store_true", help="WARN -> exit 1")
    a = ap.parse_args(argv)
    reg = None
    if a.registry:
        try:
            reg = load_registry(a.registry)
        except RegistryError as exc:
            print(f"canon-gate: BLIND (exit 2): {exc}")
            return 2
    worst = 0
    out = []
    for f in a.files:
        r = check_file(f, reg)       # reg None -> per-file lookup (BLIND when no registry is found)
        rc = r.exit_code(a.strict)
        # 1 beats 2 only if no BLIND is hidden: report the max numerically but keep both visible
        worst = max(worst, rc)
        if a.json:
            out.append({"file": f, "status": r.status, "exit": rc, "blind": r.blind_reason,
                        "findings": [f_.__dict__ for f_ in r.findings]})
        else:
            print(r.summary())
            for fd in r.findings:
                if fd.level != "OK":
                    print("  " + fd.line())
    if a.json:
        print(json.dumps(out, indent=2, ensure_ascii=False))
    return worst


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    sys.exit(cli())

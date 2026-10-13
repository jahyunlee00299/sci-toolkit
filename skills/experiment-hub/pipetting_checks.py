"""Pure-Python (openpyxl-only) validation for pipetting-protocol workbooks.

WHY THIS EXISTS: repeated verification sessions on one workbook series
re-opened Excel via win32com.client for every single check — volume closure,
concentration re-derivation, sampling headroom — even though none of those
checks touch a live Excel calculation engine. Each COM round-trip costs
seconds and adds one more chance to hang on a modal dialog (measured: one
COM session sat Not-Responding for 44+ min and required taskkill, which also
rolled the file back to its last save). These checks read the workbook with
openpyxl only and never open Excel; run excel_com_guard.recalc_and_scan()
separately, once, only for the final formula-recalculation + error-cell scan.

Two read modes matter and are NOT interchangeable:
    - data_only=False (formulas): needed to detect a "leading-=" text cell
      (see leading_equals_guard) and to know which cells are formulas at all.
    - data_only=True (cached values): needed for the actual numeric checks.
      This returns None for any formula Excel has never calculated (e.g. a
      value written by openpyxl and never opened in Excel) — callers must
      treat None as "not yet computed", not as zero.

None of this is a general workbook-QC tool. It encodes the three checks this
project's pipetting sheets actually need: does DW-based volume closure sum to
the target volume, does the mM-from-stock-and-uL re-derivation match the
sheet's stated concentration, and is there enough sampling headroom left in
the well/tube after all timepoints are drawn.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class Finding:
    sheet: str
    cell: str
    message: str
    severity: str = "error"  # "error" | "warning"


@dataclass
class CheckReport:
    findings: list[Finding] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not any(f.severity == "error" for f in self.findings)

    def add(self, sheet: str, cell: str, message: str, severity: str = "error"):
        self.findings.append(Finding(sheet, cell, message, severity))

    def merge(self, other: "CheckReport"):
        self.findings.extend(other.findings)

    def to_dict(self) -> dict:
        return {
            "ok": self.ok,
            "findings": [
                {"sheet": f.sheet, "cell": f.cell, "message": f.message, "severity": f.severity}
                for f in self.findings
            ],
        }


def leading_equals_guard(path: Path, sheet_names: list[str] | None = None) -> CheckReport:
    """Flag text cells whose VALUE starts with '=' but which openpyxl parsed
    as a plain string, not a formula.

    This happens when a formula is typed into a cell that was previously
    formatted/entered as Text (number_format == '@'), or pasted as values
    with a stray leading '='. Excel then displays the literal string "=B2*2"
    instead of evaluating it — silently wrong data that LOOKS like a formula
    at a glance and will NOT show up as a #VALUE!/#REF! error during a COM
    recalculation pass, because Excel never treats it as a formula to
    recalculate in the first place. This is exactly the defect class the
    earlier verification sessions needed a manual eyeball pass to catch before this
    guard existed.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=False)
    try:
        names = sheet_names or wb.sheetnames
        for name in names:
            if name not in wb.sheetnames:
                report.add(name, "-", f"sheet '{name}' not found", severity="warning")
                continue
            ws = wb[name]
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.startswith("=") and cell.data_type != "f":
                        report.add(
                            name,
                            cell.coordinate,
                            f"leading-= text cell (not a live formula): {v!r}",
                        )
    finally:
        wb.close()
    return report


def volume_closure_check(
    path: Path,
    sheet: str,
    component_cols: list[str],
    total_col: str,
    row_range: tuple[int, int],
    target_volume: float,
    tolerance: float = 1e-6,
) -> CheckReport:
    """Verify SUM(component_cols) == target_volume for every row in row_range,
    using Excel's own cached calculated values (data_only=True) — this is the
    "volume closure" check: every component's uL plus diluent/water (DW) must
    sum to the reaction's total working volume, or the recipe silently under-
    or over-fills the tube.

    A component cell holding None (formula never recalculated by Excel) is
    reported as its own finding rather than silently treated as 0 — a zero
    read here would make a genuinely-broken sheet look closed by accident.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            report.add(sheet, "-", f"sheet '{sheet}' not found")
            return report
        ws = wb[sheet]
        lo, hi = row_range
        for r in range(lo, hi + 1):
            values = []
            uncomputed = []
            for col in component_cols:
                cell = ws[f"{col}{r}"]
                if cell.value is None:
                    uncomputed.append(cell.coordinate)
                elif isinstance(cell.value, (int, float)):
                    values.append(cell.value)
                else:
                    report.add(sheet, cell.coordinate, f"non-numeric component value: {cell.value!r}")
            if uncomputed:
                report.add(
                    sheet, f"{sheet}!{r}", f"uncomputed formula cell(s) {uncomputed} — run a COM recalc first"
                )
                continue
            total = sum(values)
            diff = total - target_volume
            if abs(diff) > tolerance:
                report.add(
                    sheet,
                    f"{total_col}{r}",
                    f"volume closure off by {diff:+.4f} uL (sum={total:.4f}, target={target_volume:.4f})",
                )
    finally:
        wb.close()
    return report


def concentration_rederivation_check(
    path: Path,
    sheet: str,
    checks: list[dict],
    tolerance_rel: float = 1e-3,
) -> CheckReport:
    """Independently re-derive final concentration from stock_conc * volume_uL
    / total_volume_uL and compare against the sheet's stated final
    concentration cell — the check called "concentration
    re-derivation": trust nothing the workbook computed for itself, redo the
    dilution arithmetic in plain Python from the same three inputs the sheet
    used, and flag any mismatch beyond tolerance_rel (relative).

    Each entry of `checks` is a dict:
        {"row": int, "stock_col": str, "vol_col": str, "final_col": str,
         "total_volume_uL": float}
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            report.add(sheet, "-", f"sheet '{sheet}' not found")
            return report
        ws = wb[sheet]
        for spec in checks:
            r = spec["row"]
            stock = ws[f'{spec["stock_col"]}{r}'].value
            vol = ws[f'{spec["vol_col"]}{r}'].value
            final_cell = ws[f'{spec["final_col"]}{r}']
            final = final_cell.value
            total_vol = spec["total_volume_uL"]
            missing = [
                name
                for name, val in (("stock", stock), ("vol", vol), ("final", final))
                if val is None
            ]
            if missing:
                report.add(sheet, final_cell.coordinate, f"missing value(s) for re-derivation: {missing}")
                continue
            expected = stock * vol / total_vol
            if expected == 0:
                rel_diff = abs(final - expected)
            else:
                rel_diff = abs(final - expected) / abs(expected)
            if rel_diff > tolerance_rel:
                report.add(
                    sheet,
                    final_cell.coordinate,
                    f"concentration mismatch: sheet={final}, re-derived={expected:.6g} "
                    f"(rel diff {rel_diff:.2%})",
                )
    finally:
        wb.close()
    return report


def sampling_headroom_check(
    path: Path,
    sheet: str,
    total_volume_uL: float,
    sample_volume_uL: float,
    n_timepoints_col: str,
    row_range: tuple[int, int],
    min_headroom_uL: float = 0.0,
) -> CheckReport:
    """Verify total_volume_uL - (sample_volume_uL * n_timepoints) stays above
    min_headroom_uL for every row — catches a protocol that draws more sample
    across its timepoints than the tube/well actually holds after the last
    draw, which otherwise only surfaces mid-experiment at the bench.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            report.add(sheet, "-", f"sheet '{sheet}' not found")
            return report
        ws = wb[sheet]
        lo, hi = row_range
        for r in range(lo, hi + 1):
            n_cell = ws[f"{n_timepoints_col}{r}"]
            n = n_cell.value
            if n is None:
                continue
            if not isinstance(n, (int, float)):
                report.add(sheet, n_cell.coordinate, f"non-numeric timepoint count: {n!r}")
                continue
            headroom = total_volume_uL - sample_volume_uL * n
            if headroom < min_headroom_uL:
                report.add(
                    sheet,
                    n_cell.coordinate,
                    f"sampling headroom {headroom:.2f} uL below minimum {min_headroom_uL:.2f} uL "
                    f"({n:g} timepoints x {sample_volume_uL:g} uL from {total_volume_uL:g} uL)",
                )
    finally:
        wb.close()
    return report

# ---------------------------------------------------------------------------
# Structural guards (added after repeated independent-review rounds on one
# workbook series). Three regex attempts at these defects produced one miss and 23 false
# positives; every one of them is a STRUCTURAL property, so these check
# structure instead of chasing symptom text.
#
# `leading_equals_guard` above and `formula_validity_guard` /
# `constant_formula_guard` below are a PAIR, each covering the other's blind
# spot:
#   - leading_equals_guard : text that LOOKS like a formula
#   - formula_validity / constant_formula : a formula where text or a constant
#     belonged
# The cell that once made a workbook unopenable passed leading_equals_guard
# with 0 findings.
#
# Run order (structure before arithmetic): cache_freshness -> formula_validity
# -> constant_formula -> dead_input_cell -> volume_closure -> concentration ->
# sampling. `volume_closure_check` reports "uncomputed formula cell(s) - run a
# COM recalc first" when caches are empty, which is a confusing proxy for
# "nobody recalculated"; cache_freshness_guard says that directly, so it leads.
# ---------------------------------------------------------------------------

import re as _re
import zipfile as _zipfile

_FUNC = _re.compile(r"[A-Z][A-Z0-9.]*\s*\(")
_REF = _re.compile(r"(?:'[^']+'!)?\$?[A-Z]{1,3}\$?\d+")
_STRLIT = _re.compile(r'"[^"]*"')


def cache_freshness_guard(path: Path, min_filled_ratio: float = 0.5) -> CheckReport:
    """Report whether Excel has actually recalculated this file.

    An openpyxl save strips calcChain.xml and empties every cached <v>, so a
    "COM recalculation passed" claim made before such a save is void. The file
    still opens correctly (calcPr fullCalcOnLoad), but every openpyxl-based
    check downstream then reads None instead of a number.
    """
    report = CheckReport()
    with _zipfile.ZipFile(path) as z:
        names = z.namelist()
        has_chain = any("calcChain" in n for n in names)
        formulas = filled = 0
        for n in (n for n in names if n.startswith("xl/worksheets/sheet")):
            xml = z.read(n).decode("utf-8", "replace")
            formulas += len(_re.findall(r"<f[ >]", xml))
            filled += len(_re.findall(r"</f><v>[^<]", xml))  # openpyxl writes an EMPTY <v></v>
    ratio = (filled / formulas) if formulas else 1.0
    if not has_chain:
        report.add("-", "-", "calcChain.xml missing: Excel has not recalculated this file "
                             "since the last non-Excel save")
    if formulas and ratio < min_filled_ratio:
        report.add("-", "-", f"only {filled}/{formulas} formula caches filled ({ratio:.0%}) "
                             f"- openpyxl-based checks will read None")
    return report


def formula_validity_guard(path: Path) -> CheckReport:
    """Flag formulas holding tokens that are neither function, reference, number
    nor operator.

    The motivating case was a note written into a cell as
    '= 100/1000 x 50e-3 x 112.09': 'x' is not an Excel operator and the spaces
    act as intersection operators, so the file would not open at all.
    String literals are stripped first, or prose inside quotes would be flagged.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=False)
    try:
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    body = _STRLIT.sub(" ", v[1:])
                    body = _FUNC.sub(" ", body)
                    body = _REF.sub(" ", body)
                    leftover = _re.sub(r"[\d.,:;()\[\]{}+\-*/^%&<>=!\s$]", "", body)
                    if leftover:
                        report.add(name, cell.coordinate,
                                   f"formula holds non-operator token(s) {leftover!r}: {v!r}")
    finally:
        wb.close()
    return report


def constant_formula_guard(path: Path) -> CheckReport:
    """Flag formulas whose body is a bare number, such as '=9' or '=27'.

    They look live and are not: a reader edits one and nothing changes, while
    the number that actually drives the sheet stays hard-coded elsewhere.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=False)
    try:
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and _re.fullmatch(r"=\s*[\d.]+\s*", v):
                        report.add(name, cell.coordinate,
                                   f"constant disguised as a formula: {v!r}")
    finally:
        wb.close()
    return report


def dead_input_cell_guard(path: Path, max_populated: int = 4) -> CheckReport:
    """Flag labelled input cells that no formula references (in-degree 0).

    One workbook set a premix-scale cell to 36 while every volume below it still
    multiplied by the literal 13, then added a note telling the reader to change
    that cell - an actively wrong instruction.

    Two things keep the false-positive rate usable:
      1. cross-sheet prefixes ('Reaction Matrix'!$B$39) must be parsed, or the
         parameter cells they point at get reported as dead themselves;
      2. rows with more than `max_populated` filled cells are skipped - those are
         data tables, whose numbers are legitimately unreferenced. This single
         rule took the prototype from 16 findings to 5.
    Advisory by default; promote to error once a per-workbook whitelist exists.
    """
    report = CheckReport()
    wb = load_workbook(path, data_only=False)
    try:
        referenced: set = set()
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    v = cell.value
                    if not (isinstance(v, str) and v.startswith("=")):
                        continue
                    for m in _REF.finditer(_STRLIT.sub(" ", v)):
                        ref = m.group(0)
                        sheet, _, addr = ref.rpartition("!")
                        sheet = sheet.strip("'") or name
                        referenced.add((sheet, addr.replace("$", "")))
        for name in wb.sheetnames:
            ws = wb[name]
            for row in ws.iter_rows():
                populated = [c for c in row if c.value is not None]
                if len(populated) > max_populated:
                    continue
                for cell in populated:
                    if not isinstance(cell.value, (int, float)):
                        continue
                    label = ws.cell(cell.row, 1).value
                    if not isinstance(label, str) or not label.strip():
                        continue
                    if (name, cell.coordinate) not in referenced:
                        report.add(name, cell.coordinate,
                                   f"input cell {label.strip()!r} = {cell.value} is referenced by "
                                   f"no formula - changing it does nothing",
                                   severity="warning")
    finally:
        wb.close()
    return report

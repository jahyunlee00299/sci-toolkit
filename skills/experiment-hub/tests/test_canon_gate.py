"""Tests for canon_gate.py against the SYNTHETIC example registry (canonical_constants.example.toml) and the
synthetic configs / Params workbook built in synthetic_fixtures.py."""
import json
import subprocess
import sys
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import canon_gate as g  # noqa: E402
import synthetic_fixtures as sf  # noqa: E402

REG = g.load_registry(g.EXAMPLE_REGISTRY)


def fails(res):
    return {f.entry for f in res.findings if f.level == "FAIL"}


@pytest.fixture
def params_xlsx(tmp_path):
    return sf.write_params_xlsx(tmp_path / "example_params.xlsx")


@pytest.fixture
def canon_json(tmp_path):
    return sf.dump(sf.canon_config(), tmp_path / "example_config.json")


def test_fail_params_sheet_with_cofactor_in_uM(params_xlsx):
    r = g.check_file(params_xlsx, REG)
    assert r.exit_code() == 1 and r.status == "FAIL"
    assert fails(r) == {"cofactorx.stock", "cofactorx.final"}
    assert any("SUPERSEDED" in f.message for f in r.findings if f.level == "FAIL")
    assert any(f.entry == "enza.batch" and f.level == "OK" for f in r.findings)   # batch read from the note


def test_synthetic_workbook_metadata_is_neutral(params_xlsx):
    import openpyxl
    props = openpyxl.load_workbook(params_xlsx).properties
    assert props.creator == "synthetic-fixture" and props.lastModifiedBy == "synthetic-fixture"


def test_ok_config_with_working_stocks_warns_only(canon_json):
    r = g.check_file(canon_json, REG)
    assert r.exit_code() == 0 and r.n_fail == 0
    assert r.n_warn == 2  # SaltM / SaltC working stocks differ from the Params sheet (undecided, not failures)
    assert g.check_file(canon_json, REG).exit_code(strict=True) == 1


def test_ok_pooled_config(tmp_path):
    r = g.check_file(sf.dump(sf.pooled_config(), tmp_path / "pooled.json"), REG)
    assert r.exit_code() == 0 and r.n_ok >= 4


def test_fail_superseded_batch():
    cfg = sf.canon_config()
    cfg["enzymes"]["EnzA"]["batch"] = "100009"
    r = g.check_config(cfg, REG, source="x.json")
    assert r.exit_code() == 1 and fails(r) == {"enza.batch"}
    assert any("SUPERSEDED" in f.message for f in r.findings if f.level == "FAIL")


def test_fail_enzyme_stock_changed():
    cfg = sf.canon_config()
    cfg["enzymes"]["EnzB"]["stock_gL"] = 70.0
    assert fails(g.check_config(cfg, REG, source="x.json")) == {"enzb.stock_gL"}


def test_fail_cofactor_in_wrong_unit_config():
    cfg = sf.canon_config()
    cfg["stocks"]["CofactorX"]["conc_mM"] = 0.01  # 10 uM written as mM
    assert "cofactorx.stock" in fails(g.check_config(cfg, REG, source="x.json"))


def test_alias_names_match():
    cfg = sf.canon_config()
    cfg["stocks"]["Cofactor-X"] = cfg["stocks"].pop("CofactorX")
    cfg["stocks"]["Cofactor-X"]["final_mM"] = 0.15
    assert fails(g.check_config(cfg, REG, source="x.json")) == {"cofactorx.final"}


def test_blind_config_without_registry_fields():
    r = g.check_config({"stocks": {"NaCl": {"conc_mM": 1, "type": "buffer"}}, "enzymes": {}}, REG, source="x.json")
    assert r.exit_code() == 2 and r.status == "BLIND"


def test_blind_xlsx_without_params_sheet(tmp_path):
    import openpyxl
    p = tmp_path / "np.xlsx"
    wb = openpyxl.Workbook(); wb.active.title = "Reaction_Mix"; wb.save(p)
    assert g.check_file(p, REG).exit_code() == 2


def test_blind_params_sheet_with_no_known_field(tmp_path):
    import openpyxl
    p = tmp_path / "pp.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Params"
    ws.append(["Component", "Stock conc", "Unit", "Note"]); ws.append(["NaCl", 5, "mM", ""]); wb.save(p)
    assert g.check_file(p, REG).exit_code() == 2


def test_blind_unreadable_and_bad_registry(tmp_path, canon_json):
    bad = tmp_path / "bad.json"; bad.write_text("{not json", encoding="utf-8")
    assert g.check_file(bad, REG).exit_code() == 2
    reg = tmp_path / "r.toml"; reg.write_text("[meta]\nschema=1\n", encoding="utf-8")
    with pytest.raises(g.RegistryError):
        g.load_registry(reg)
    assert g.cli([str(canon_json), "--registry", str(reg)]) == 2


def test_example_registry_entries_have_provenance_and_no_person_names():
    for e in REG["canon"]:
        for k in ("value", "decided", "decided_by", "basis", "source", "field", "names"):
            assert k in e, (e["id"], k)
        assert e["decided_by"] in ("researcher", "n/a (two files agree)"), e["decided_by"]
    ids = {e["id"] for e in REG["canon"]}
    assert {"cofactorx.stock", "cofactorx.final", "enza.batch"} <= ids
    # working stocks must never be canonical
    assert not any(i.startswith(("saltm", "saltc")) for i in ids)


def test_cli_exit_codes_via_subprocess(params_xlsx, canon_json):
    py = sys.executable
    run = lambda p: subprocess.run([py, str(HERE.parent / "canon_gate.py"), str(p), "--registry",
                                    str(g.EXAMPLE_REGISTRY)], capture_output=True, text=True,
                                   encoding="utf-8").returncode
    assert (run(params_xlsx), run(canon_json)) == (1, 0)


def test_validate_config_blocks_registry_contradiction():
    import reaction_matrix as rm
    cfg = sf.canon_config()
    # this neat-stock design is blocked by the pipetting rules (EnzF 0.357 uL added per tube < the 0.5 uL
    # floor; EnzF/EnzG at 3 volumes) -- assert that, then waive them with a reason so the canon-gate path below
    # is tested on its own.
    ok, rep = rm.validate_config(cfg)
    assert not ok and any("< pipette floor 0.5 uL" in l for l in rep) and any("different volumes" in l for l in rep)
    for e in cfg["enzymes"].values():
        e.update(allow_small_volume=True, allow_small_volume_reason="fixture, canon-gate test",
                 allow_variable_volume=True, allow_variable_volume_reason="fixture, canon-gate test")
    # the 0.1 uL rounding check would flag this design's small draws; that rule has its own tests -- use a
    # fine resolution here so only the canon-gate path is exercised
    cfg["pipettes"] = {"resolution_uL": 0.01}
    ok, rep = rm.validate_config(cfg)
    assert ok, rep
    cfg["enzymes"]["EnzA"]["batch"] = "100009"
    ok, rep = rm.validate_config(cfg)
    assert not ok and any(l.startswith("FAIL canon-gate") for l in rep)


def test_config_diluted_working_stock_checks_source_stock(tmp_path):
    """A diluted working stock is pipetted at stock_gL but the canonical check uses source_stock_gL."""
    cfg = {"enzymes": {"EnzF": {"batch": "100006", "stock_gL": 17.5, "source_stock_gL": 70.0}}, "stocks": {},
           "conditions": []}
    p = sf.dump(cfg, tmp_path / "diluted.json")
    assert not fails(g.check_file(p, REG))
    cfg["enzymes"]["EnzF"]["source_stock_gL"] = 60.0
    sf.dump(cfg, p)
    assert fails(g.check_file(p, REG)) == {"enzf.stock_gL"}


def test_equal_volume_expansion_and_canon_check():
    """equal_volume.py: same volume per tube, diluted working stock per level, canon check on source stock."""
    import equal_volume as ev
    spec = {"total_volume_uL": 50, "stocks": {}, "enzymes": {"EnzF": {"batch": "100006", "stock_gL": 70.0,
            "equal_volume_uL": 2, "rxn_gL": {"1x": 0.5, "4x": 2.0}}},
            "conditions": [{"num": 1, "substrates": {}, "enzymes": {"EnzF": "1x"}},
                           {"num": 2, "substrates": {}, "enzymes": {"EnzF": "4x"}}]}
    out = ev.expand(spec)
    assert out["enzymes"]["EnzF [1x]"]["stock_gL"] == 12.5 and out["enzymes"]["EnzF [4x]"]["stock_gL"] == 50.0
    assert out["conditions"][0]["enzymes"] == {"EnzF [1x]": "1x", "EnzF [4x]": "0x"}
    assert not fails(g.check_config(out, REG))
    spec["enzymes"]["EnzF"]["equal_volume_uL"] = 0.5      # would need 200 g/L > lab stock
    with pytest.raises(ValueError):
        ev.expand(spec)

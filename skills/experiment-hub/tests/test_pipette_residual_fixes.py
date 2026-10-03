"""Residual defects found by a second adversarial review of the pipetting rules, one test group each:

a  control_water_in_place: false -> WARN (the no-enzyme control gets one odd DW volume)
b  allow_small_volume waives the sub-floor volume rule only, never a > 5 % rounding FAIL
c  one physical enzyme under two names: batch normalisation, dilution entries without source_stock_gL
d  no DW top-up between 0 and the pipette floor: the common water in MM-A is lowered, every tube closes
e  premix.extra_rxns makes the MM-A margin configurable
g  canon_gate registry lookup order (env, next to the config, skill file, example; none = BLIND)
"""
import copy
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
SK = HERE.parent
sys.path.insert(0, str(SK))
sys.path.insert(0, str(HERE))
import canon_gate as g  # noqa: E402
import equal_volume as ev  # noqa: E402
import pipette_count as pc  # noqa: E402
import reaction_matrix as rm  # noqa: E402
from test_pipette_rule_fixes import base, ne_toy  # noqa: E402
import synthetic_fixtures as sf  # noqa: E402



def fails(rep):
    return [l for l in rep if l.startswith("FAIL")]


def closes(rp, V):
    return all(abs(sum(r[3] for r in rows) - V) < 1e-9 for rows in rp["tubes"].values())


# ---------------------------------------------------------------- a
def test_a_control_water_in_place_off_warns():
    ok, rep = rm.validate_config(ne_toy())
    assert ok and not any("control_water_in_place" in l for l in rep)
    c = ne_toy()
    c["control_water_in_place"] = False
    ok, rep = rm.validate_config(c)
    hit = [l for l in rep if l.startswith("WARN uniformity") and '"control_water_in_place": false' in l]
    assert ok and len(hit) == 1, rep
    assert "no-enzyme tube #3" in hit[0] and "a volume no other tube uses" in hit[0]
    # no enzyme tube anywhere -> nothing to imitate, no warning
    c = base()
    for cond in c["conditions"]:
        cond["enzymes"] = {"E": "0x"}
    c["enzymes"]["E"]["rxn_gL"]["0x"] = 0
    c["control_water_in_place"] = False
    ok, rep = rm.validate_config(c)
    assert not any("control_water_in_place" in l for l in rep)


# ---------------------------------------------------------------- b
def test_b_small_volume_waiver_never_waives_rounding_fail():
    c = base()
    c["enzymes"]["E"]["rxn_gL"]["1x"] = 0.05                 # 0.25 uL -> 0.3 uL after rounding: +20 %
    c["enzymes"]["E"].update(allow_small_volume=True, allow_small_volume_reason="positive-displacement pipette")
    ok, rep = rm.validate_config(c)
    assert any("WAIVED (allow_small_volume" in l and "pipette floor" in l for l in rep)   # floor rule waived
    assert not ok and any(l.startswith("FAIL rounding: E deviates +20.00 %") and "does not waive" in l
                          for l in rep), rep
    c["enzymes"]["E"]["rxn_gL"]["1x"] = 0.21                 # 1.05 uL -> 1.1 uL: +4.76 %, within 2-5 %
    ok, rep = rm.validate_config(c)
    assert ok and any(l.startswith("WARN rounding: E deviates +4.76 %") and "WAIVED" in l for l in rep), rep


# ---------------------------------------------------------------- c
def two_names(dil_extra=None, batch_a="b1", batch_b="b1"):
    c = base()
    c["enzymes"] = {"EnzX": {"batch": batch_a, "stock_gL": 10.0, "rxn_gL": {"1x": 0.4}},
                    "EnzX_dil": {"batch": batch_b, "stock_gL": 2.5, "rxn_gL": {"1x": 0.4}, **(dil_extra or {})}}
    c["conditions"][0]["enzymes"] = {"EnzX": "1x"}
    c["conditions"][1]["enzymes"] = {"EnzX_dil": "1x"}
    return c


@pytest.mark.parametrize("label,cfg", [
    ("declares source_stock_gL", two_names({"source_stock_gL": 10.0})),
    ("no source_stock_gL (name stem + batch)", two_names()),
    ("batch with trailing space", two_names({"source_stock_gL": 10.0}, batch_b="b1 ")),
    ("batch case differs", two_names(batch_a="B1")),
    ("batch int vs str", two_names({"source_stock_gL": 10.0}, batch_a="100001", batch_b=100001)),
    ("batch float vs str", two_names(batch_a="100001", batch_b=100001.0)),
])
def test_c_one_physical_enzyme_under_two_names(label, cfg):
    ok, rep = rm.validate_config(cfg)
    assert not ok, label
    assert any(l.startswith("FAIL EnzX") and "2 different volumes" in l and "[config entries: EnzX, EnzX_dil]" in l
               for l in rep), (label, rep)


def test_c_dilution_x_reconstructs_the_lab_stock_and_unrelated_enzymes_stay_apart():
    c = two_names({"dilution_x": 4})
    c["enzymes"]["Other"] = c["enzymes"].pop("EnzX_dil")        # unrelated name, but 2.5 x 4 = the 10 g/L stock
    c["conditions"][1]["enzymes"] = {"Other": "1x"}
    ok, rep = rm.validate_config(c)
    assert not ok and any("[config entries: EnzX, Other]" in l for l in fails(rep))
    c = base()                                                    # two DIFFERENT enzymes sharing a batch label
    c["enzymes"] = {"EnzP": {"batch": "shared ref", "stock_gL": 10.0, "rxn_gL": {"1x": 0.4}},
                    "EnzQ": {"batch": "Shared Ref", "stock_gL": 12.0, "rxn_gL": {"1x": 0.6}}}
    for cond in c["conditions"]:
        cond["enzymes"] = {"EnzP": "1x", "EnzQ": "1x"}
    ok, rep = rm.validate_config(c)
    assert ok, fails(rep)
    assert rm._name_stem("EnzX [1x]") == rm._name_stem("EnzX_dil") == rm._name_stem("EnzX working") == "enzx"
    assert rm._norm_batch(100001.0) == rm._norm_batch(" 100001 ") == "100001" and rm._norm_batch(True) is None


def test_c_canon_gate_reads_an_integer_batch():
    reg = g.load_registry(g.EXAMPLE_REGISTRY)
    r = g.check_config({"stocks": {}, "enzymes": {"EnzA": {"batch": 100009, "stock_gL": 12.0}}}, reg, "x.json")
    assert r.exit_code() == 1 and {f.entry for f in r.findings if f.level == "FAIL"} == {"enza.batch"}


# ---------------------------------------------------------------- d
def topup_case():
    c = base()
    c["conditions"][1]["substrates"]["S"] = 106               # 5.3 vs 5.0 uL: #1 would need a 0.3 uL top-up
    return c


def test_d_common_water_lowered_so_no_top_up_is_below_the_floor(tmp_path):
    c = topup_case()
    plan = pc.build_plan(c)
    assert plan["water_common_max_uL"] == pytest.approx(38.7)
    assert plan["water_common_uL"] == pytest.approx(38.2)      # lowered by the 0.5 uL floor
    assert plan["topup"] == pytest.approx({1: 0.8, 2: 0.5})
    assert all(v == 0 or v >= 0.5 for v in plan["topup"].values())
    rp = pc.rounded_plan(plan)
    assert rp["dropped"] == {} and closes(rp, 50.0)
    assert all(v == 0 or v >= 0.5 for v in rp["dw"].values())
    ok, rep = rm.validate_config(c)
    assert ok, fails(rep)
    assert not any("is skipped" in l for l in rep)
    assert any(l.startswith("INFO R1: common water in MM-A lowered from 38.7 to 38.2") for l in rep)
    p, x = tmp_path / "t.json", tmp_path / "t.xlsx"
    p.write_text(json.dumps(c), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SK / "reaction_matrix.py"), str(p), str(x)],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout + r.stderr
    pg = openpyxl.load_workbook(x)["Pipetting Guide"]
    row = [i for i in range(1, 30) if str(pg.cell(i, 1).value).startswith("DW (common water")]
    assert row and "lowered" in pg.cell(row[0], 1).value and pg.cell(row[0], 4).value == pytest.approx(38.2)


def test_d_rounding_residual_below_floor_lowers_mm_a_instead_of_dropping():
    c = base()
    c["stocks"]["S2"] = {"conc_mM": 1000, "type": "substrate"}
    c["conditions"][0]["substrates"] = {"S": 100.8, "S2": 12.8}  # 5.04 + 0.64 uL, both round DOWN
    c["conditions"][1]["substrates"] = {"S": 100.8}
    plan = pc.build_plan(c)
    assert plan["topup"][1] == 0.0
    rp = pc.rounded_plan(plan)
    assert rp["dropped"] == {} and closes(rp, 50.0)
    assert all(v == 0 or v >= 0.5 - 1e-9 for v in rp["dw"].values()), rp["dw"]


def test_d_existing_designs_still_close():
    for cfg in (ev.expand(sf.ev_spec()), ne_toy(), topup_case()):
        rp = pc.rounded_plan(pc.build_plan(cfg))
        assert rp["dropped"] == {} and closes(rp, float(cfg["total_volume_uL"]))


# ---------------------------------------------------------------- e
def test_e_premix_extra_rxns_removes_the_premix_rounding_warnings(tmp_path):
    spec = sf.ev_spec()
    ok, rep = rm.validate_config(ev.expand(spec))
    assert any(l.startswith("WARN rounding: CofactorX") for l in rep) and any(l.startswith("WARN rounding: CofactorY") for l in rep)
    spec["premix"] = {"extra_rxns": 5}
    cfg = ev.expand(spec)
    plan = pc.build_plan(cfg)
    draws = {s.source: s.volume_uL for s in plan["steps"] if s.target == "MM-A"}
    assert plan["mm_a_rxns"] == 16
    assert pc._rnd(draws["CofactorX"], 0.1) == pytest.approx(1.2) and draws["CofactorX"] == pytest.approx(1.2)
    assert pc._rnd(draws["CofactorY"], 0.1) == pytest.approx(2.0) and draws["CofactorY"] == pytest.approx(2.0)
    ok, rep = rm.validate_config(cfg)
    assert ok, fails(rep)
    assert not any(l.startswith("WARN rounding: CofactorX") or l.startswith("WARN rounding: CofactorY") for l in rep), rep
    p, x = tmp_path / "e.json", tmp_path / "e.xlsx"
    p.write_text(json.dumps(cfg), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SK / "reaction_matrix.py"), str(p), str(x)],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout
    wb = openpyxl.load_workbook(x)
    pg = wb["Pipetting Guide"]
    assert any(pg.cell(i, 1).value == "# rxns (with margin)" and pg.cell(i, 2).value == 16 for i in range(1, 12))
    assert rm.surplus_note(5) in [c.value for row in wb["Enzyme Additions"].iter_rows() for c in row]
    for bad in (-1, 2.5, "5", True):
        spec["premix"] = {"extra_rxns": bad}
        ok, rep = rm.validate_config(ev.expand(spec))
        assert not ok and any("premix.extra_rxns" in l for l in fails(rep)), bad


# ---------------------------------------------------------------- g
REG_FAIL = """
[meta]
schema = 1
[[canon]]
id = "enza.batch"
names = ["EnzA"]
field = "batch"
value = "100002"
decided = "t"
decided_by = "researcher"
basis = "researcher-decision"
source = "test"
"""


def enz_cfg():
    return {"stocks": {}, "enzymes": {"EnzA": {"batch": "100001", "stock_gL": 12.0}}, "conditions": []}


def test_g_registry_lookup_order(tmp_path, monkeypatch):
    monkeypatch.delenv(g.ENV_REGISTRY, raising=False)
    cfgp = tmp_path / "cfg.json"
    cfgp.write_text(json.dumps(enz_cfg()), encoding="utf-8")
    other = tmp_path / "elsewhere"
    other.mkdir()
    # 4. no lab registry at all -> the example (EnzA batch 100001 = OK there)
    monkeypatch.setattr(g, "DEFAULT_REGISTRY", tmp_path / "missing.toml")
    assert g.resolve_registry(cfgp) == g.EXAMPLE_REGISTRY and g.check_file(cfgp).exit_code() == 0
    # 3. the skill's own registry wins over the example
    own = other / "own.toml"
    own.write_text(REG_FAIL, encoding="utf-8")
    monkeypatch.setattr(g, "DEFAULT_REGISTRY", own)
    assert g.resolve_registry(cfgp) == own and g.check_file(cfgp).exit_code() == 1
    # 2. canonical_constants.toml next to the config wins over the skill's own
    near = tmp_path / g.REGISTRY_NAME
    near.write_text(REG_FAIL.replace("100002", "100001"), encoding="utf-8")
    assert g.resolve_registry(cfgp) == near and g.check_file(cfgp).exit_code() == 0
    assert g.check_config(enz_cfg(), source=str(cfgp)).exit_code() == 0
    # 1. the environment variable wins over everything
    monkeypatch.setenv(g.ENV_REGISTRY, str(own))
    assert g.resolve_registry(cfgp) == own and g.check_file(cfgp).exit_code() == 1
    monkeypatch.setenv(g.ENV_REGISTRY, str(tmp_path / "nope.toml"))   # set but missing = BLIND, no fall-through
    r = g.check_file(cfgp)
    assert r.exit_code() == 2 and g.ENV_REGISTRY in r.blind_reason


def test_g_no_registry_anywhere_is_blind(tmp_path, monkeypatch):
    monkeypatch.delenv(g.ENV_REGISTRY, raising=False)
    monkeypatch.setattr(g, "DEFAULT_REGISTRY", tmp_path / "a.toml")
    monkeypatch.setattr(g, "EXAMPLE_REGISTRY", tmp_path / "b.toml")
    cfgp = tmp_path / "cfg.json"
    cfgp.write_text(json.dumps(enz_cfg()), encoding="utf-8")
    r = g.check_file(cfgp)
    assert r.exit_code() == 2 and r.status == "BLIND" and "no canonical-constants registry" in r.blind_reason
    assert g.cli([str(cfgp)]) == 2
    assert g.check_config(enz_cfg()).exit_code() == 2
    ok, rep = rm.validate_config(base())                  # the pipetting validator shows it, never crashes
    assert any(l.startswith("WARN canon-gate BLIND") for l in rep)


def test_g_cli_finds_the_registry_next_to_the_file(tmp_path):
    cfgp = tmp_path / "cfg.json"
    cfgp.write_text(json.dumps(enz_cfg()), encoding="utf-8")
    (tmp_path / g.REGISTRY_NAME).write_text(REG_FAIL, encoding="utf-8")
    env = {k: v for k, v in __import__("os").environ.items() if k != g.ENV_REGISTRY}
    r = subprocess.run([sys.executable, str(SK / "canon_gate.py"), str(cfgp)], capture_output=True, text=True,
                       encoding="utf-8", env=env)
    assert r.returncode == 1 and "enza.batch" in r.stdout

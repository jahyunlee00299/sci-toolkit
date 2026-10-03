"""R1/R2 enforcement in reaction_matrix.validate_config and the pipette_count.py model.

The toy config is small enough to count by hand; the expected numbers in test_toy_* are that hand count.
"""
import copy
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

HERE = Path(__file__).resolve().parent
SK = HERE.parent
sys.path.insert(0, str(SK))
sys.path.insert(0, str(HERE))
import equal_volume as ev  # noqa: E402
import pipette_count as pc  # noqa: E402
import reaction_matrix as rm  # noqa: E402
import synthetic_fixtures as sf  # noqa: E402


TOY = {
    "title": "toy", "date": "000000", "total_volume_uL": 50,
    "stocks": {"Tris": {"conc_mM": 1000, "type": "buffer", "final_mM": 200},
               "MgCl2": {"conc_mM": 1000, "type": "buffer", "final_mM": 5},
               "S": {"conc_mM": 1000, "type": "substrate"}},
    "enzymes": {"E1": {"batch": "x", "stock_gL": 10.0, "rxn_gL": {"1x": 0.4, "2x": 0.8}},
                "E2": {"batch": "x", "stock_gL": 5.0, "rxn_gL": {"1x": 0.2}}},
    "conditions": [{"num": 1, "substrates": {"S": 100}, "enzymes": {"E1": "1x", "E2": "1x"}},
                   {"num": 2, "substrates": {"S": 100}, "enzymes": {"E1": "2x", "E2": "1x"}},
                   {"num": 3, "substrates": {"S": 50}, "enzymes": {"E1": "1x", "E2": "1x"}}],
}


def toy_b():
    c = copy.deepcopy(TOY)
    c["enzymes"]["E1"]["equal_volume_uL"] = 4
    return c


def fails(rep):
    return [l for l in rep if l.startswith("FAIL")]


# ---------------------------------------------------------------- pipette_count: hand-counted
def test_toy_design_a_hand_count():
    s = pc.count(TOY)
    assert (s["prep_steps"], s["prep_mixes"], s["run_steps"], s["total_steps"], s["total_ops"]) == (4, 2, 14, 18, 20)
    assert (s["dw_steps_with_r1"], s["dw_steps_without_r1"], s["water_common_uL"]) == (2, 3, 28.75)
    assert (s["below_1uL"], s["below_2uL"], s["distinct_volumes_run"]) == (0, 1, 6)
    assert s["per_tube_enzyme_steps_max"] == 2 and s["prep_dilutions"] == 0


def test_toy_design_b_hand_count():
    plan = pc.build_plan(toy_b())
    s = pc.summarize(plan)
    # E1 [1x] working stock: 2 tubes x 4 uL x 1.2 + 10 dead -> 20 uL = 10 uL lab stock + 10 uL water; E1 [2x] is neat
    assert [(w["working_stock"], w["make_uL"], w["lab_stock_uL"], w["water_uL"]) for w in plan["working"]] == \
        [("E1 [1x]", 20, 10.0, 10.0)]
    assert (s["prep_steps"], s["prep_mixes"], s["run_steps"], s["total_steps"], s["total_ops"]) == (6, 3, 13, 19, 22)
    assert (s["dw_steps_with_r1"], s["dw_steps_without_r1"]) == (1, 3)


def test_cocktail_modes_change_only_enzyme_steps():
    a = pc.count(TOY, "none")
    c = pc.count(toy_b(), "per_condition")
    assert a["per_tube_enzyme_steps_max"] == 2 and a["run_by_kind"]["enzyme"] == 6
    # tubes #1 and #3 share a fingerprint -> one cocktail for them (1 step each); #2 is a single-tube
    # fingerprint, where a cocktail would add a step, so it keeps its 2 direct additions
    assert c["run_by_kind"]["cocktail_add"] == 2 and c["run_by_kind"]["enzyme"] == 2
    assert c["per_tube_enzyme_steps_min"] == 1 and c["per_tube_enzyme_steps_max"] == 2
    with pytest.raises(ValueError):
        pc.count(TOY, "bogus")


def test_weighing_is_counted_but_not_pipetting():
    cfg = sf.ev_spec()
    plan = pc.build_plan(cfg)
    s = pc.summarize(plan)
    assert s["weighings"] == 11 and all(w[1] == "SolidT" and w[2] == 1.121 for w in plan["weighings"])
    assert not any("SolidT" in st.source for st in plan["steps"])


def test_equal_volume_design_numbers_and_r1_effect():
    s = pc.count(sf.ev_spec())
    # no-enzyme control #11 gets water in place of the enzyme additions -> no DW top-up left anywhere
    assert (s["prep_dilutions"], s["below_1uL"], s["dw_steps_with_r1"], s["dw_steps_without_r1"]) == (11, 0, 0, 11)
    assert s["feasible"] and s["water_in_place_tubes"] == [11] and s["identical_sequence"]


def test_infeasible_design_is_flagged_and_cli_exit_1(tmp_path):
    c = copy.deepcopy(TOY)
    c["enzymes"]["E1"]["rxn_gL"]["2x"] = 7.0          # 35 uL of E1 -> tube #2 overfills
    s = pc.count(c)
    assert not s["feasible"] and s["min_residual_uL"] < 0
    p = tmp_path / "over.json"
    p.write_text(json.dumps(c), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SK / "pipette_count.py"), str(p)], capture_output=True, text=True,
                       encoding="utf-8")
    assert r.returncode == 1 and "INFEASIBLE" in r.stdout


def test_cli_compare_and_steps_csv(tmp_path):
    out = tmp_path / "steps.csv"
    M7EV = sf.dump(sf.ev_spec(), tmp_path / "spec.json")
    r = subprocess.run([sys.executable, str(SK / "pipette_count.py"), str(M7EV), "--steps", str(out)],
                       capture_output=True, text=True, encoding="utf-8")
    # 90 = 88 + 2: the no-enzyme control receives 3 water additions (in place of the
    # cocktail, EnzF, EnzG) instead of 1 DW top-up; without water-in-place the model gives 88
    assert r.returncode == 0 and "TOTAL 90 steps" in r.stdout
    assert len(out.read_text(encoding="utf-8").splitlines()) == 1 + 90
    cfg = sf.ev_spec()
    assert pc.summarize(pc.build_plan(cfg, water_in_place=False))["total_steps"] == 88
    r = subprocess.run([sys.executable, str(SK / "pipette_count.py"), str(M7EV), "--compare"],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0 and "Cdirect" in r.stdout and "total_steps" in r.stdout


# ---------------------------------------------------------------- validate_config rules
def test_good_equal_volume_config_passes():
    ok, rep = rm.validate_config(ev.expand(sf.ev_spec()))
    assert ok, fails(rep)
    assert any(l.startswith("PASS R2") for l in rep) and any(l.startswith("PASS R1") for l in rep)


def test_sub_microlitre_enzyme_addition_blocks():
    c = copy.deepcopy(TOY)
    c["enzymes"]["E2"]["rxn_gL"]["1x"] = 0.04          # 0.4 uL < 0.5 uL pipette floor
    ok, rep = rm.validate_config(c)
    assert not ok and any("enzyme E2" in l and "0.4 uL < pipette floor 0.5 uL" in l for l in fails(rep))


def test_small_volume_waiver_needs_reason():
    c = copy.deepcopy(TOY)
    c["enzymes"]["E2"]["rxn_gL"]["1x"] = 0.04
    c["enzymes"]["E2"]["allow_small_volume"] = True
    ok, rep = rm.validate_config(c)
    assert not ok and any("malformed waiver" in l and "needs a non-empty string" in l for l in fails(rep))
    c["enzymes"]["E2"]["allow_small_volume_reason"] = "tested P2 pipette, cocktail-pooled"
    c["enzymes"]["E1"].update(allow_variable_volume=True, allow_variable_volume_reason="toy")
    ok, rep = rm.validate_config(c)
    assert ok, fails(rep)
    assert any("WAIVED (allow_small_volume" in l for l in rep)


def test_between_floor_and_comfort_is_warn_only():
    c = toy_b()
    c["enzymes"]["E2"]["rxn_gL"]["1x"] = 0.08          # 0.8 uL: pipettable (>= 0.5 floor), poor accuracy (< 1.0)
    ok, rep = rm.validate_config(ev.expand(c))
    assert ok, fails(rep)
    assert any(l.startswith("WARN") and "0.8 uL is below 1.0 uL" in l for l in rep)


def test_min_volume_is_configurable():
    c = ev.expand(toy_b())
    c["min_addition_uL"] = 2.5                          # legacy key = floor; E2 = 2.0 uL now below it
    ok, rep = rm.validate_config(c)
    assert not ok and any("E2" in l and "2.0 uL < pipette floor 2.5 uL" in l for l in fails(rep))
    c.pop("min_addition_uL")
    c["pipettes"] = {"floor_uL": 2.5}                   # the config-level form
    ok, rep = rm.validate_config(c)
    assert not ok and any("2.0 uL < pipette floor 2.5 uL" in l for l in fails(rep))


def test_variable_volume_of_one_enzyme_blocks_unless_waived():
    ok, rep = rm.validate_config(copy.deepcopy(TOY))   # E1 at 2.0 and 4.0 uL
    assert not ok and any(l.startswith("FAIL E1: one enzyme pipetted at 2 different volumes") for l in rep)
    ok, rep = rm.validate_config(ev.expand(toy_b()))   # same design, equal volume -> passes
    assert ok, fails(rep)
    c = copy.deepcopy(TOY)
    c["enzymes"]["E1"].update(allow_variable_volume=True, allow_variable_volume_reason="dose series from one tube")
    ok, rep = rm.validate_config(c)
    assert ok and any("WAIVED (allow_variable_volume" in l for l in rep)


def test_waiver_survives_equal_volume_expansion():
    c = toy_b()
    c["enzymes"]["E1"].update(allow_small_volume=True, allow_small_volume_reason="r")
    out = ev.expand(c)
    assert all(e.get("allow_small_volume_reason") == "r" for k, e in out["enzymes"].items() if k.startswith("E1 ["))
    assert all(e.get("diluent") == "water" for k, e in out["enzymes"].items() if k.startswith("E1 ["))


def test_r1_common_water_warns_when_off():
    c = ev.expand(toy_b())
    c["common_water_in_mm_a"] = False
    ok, rep = rm.validate_config(c)
    assert ok and any(l.startswith("WARN R1") and "3 tube DW pipettings" in l for l in rep)


# ---------------------------------------------------------------- end-to-end via the CLI
def test_cli_blocks_and_writes_no_xlsx(tmp_path):
    p, x = tmp_path / "bad.json", tmp_path / "bad.xlsx"
    p.write_text(json.dumps(TOY), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SK / "reaction_matrix.py"), str(p), str(x)],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 1 and not x.exists() and "xlsx NOT generated" in r.stdout


def test_cli_good_config_writes_workbook_with_count_and_tables(tmp_path):
    p, x = tmp_path / "good.json", tmp_path / "good.xlsx"
    p.write_text(json.dumps(ev.expand(toy_b())), encoding="utf-8")
    r = subprocess.run([sys.executable, str(SK / "reaction_matrix.py"), str(p), str(x)],
                       capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout + r.stderr
    assert "Pipette count (common cocktail" in r.stdout and "TOTAL 19 steps" in r.stdout
    wb = openpyxl.load_workbook(x)
    assert "Enzyme Additions" in wb.sheetnames
    ea = [[c.value for c in row] for row in wb["Enzyme Additions"].iter_rows()]
    flat = [str(v) for row in ea for v in row if v is not None]
    assert "water (uL)" in flat and "E1 [1x]" in flat and "PER-TUBE ADDITION TABLE (pipetting order; enzymes LAST)" in flat
    pg = wb["Pipetting Guide"]
    water_rows = [r for r in range(1, 20) if str(pg.cell(r, 1).value).startswith("DW (common water")]
    assert water_rows and pg.cell(water_rows[0], 4).value == 28.75


def test_no_mm_a_means_no_common_water_and_empty_config_is_loud():
    c = {"title": "t", "total_volume_uL": 50, "stocks": {"S": {"conc_mM": 1000, "type": "substrate"}},
         "enzymes": {}, "conditions": [{"num": 1, "substrates": {"S": 10}, "enzymes": {}}]}
    plan = pc.build_plan(c)
    assert plan["water_common_uL"] == 0.0
    assert sum(s.volume_uL for s in plan["steps"] if s.phase == "run") == pytest.approx(50.0)
    c["conditions"] = []
    with pytest.raises(ValueError):
        pc.build_plan(c)

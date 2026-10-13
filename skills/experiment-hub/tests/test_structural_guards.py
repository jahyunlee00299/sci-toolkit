"""Structural guards in pipetting_checks.py, the opt-in Methods-folder publish and the `calc` entry point of
reaction_matrix.py. Every workbook here is synthetic and built in tmp_path."""
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook

SK = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(SK))

import pipetting_checks as pc  # noqa: E402
import reaction_matrix as rm  # noqa: E402


def _save(tmp_path, build, name="wb.xlsx"):
    wb = Workbook()
    build(wb)
    p = tmp_path / name
    wb.save(p)
    return p


def _cells(report):
    return {f.cell for f in report.findings}


# ---- cache_freshness_guard -------------------------------------------------

def test_cache_freshness_flags_openpyxl_save(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"] = 1, 2, "=A1+A2"
    rep = pc.cache_freshness_guard(_save(tmp_path, build))
    msgs = " ".join(f.message for f in rep.findings)
    assert not rep.ok
    assert "calcChain.xml missing" in msgs and "0/1 formula caches filled" in msgs


def test_cache_freshness_passes_recalculated_file(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"], ws["A2"], ws["A3"] = 1, 2, "=A1+A2"
    src = _save(tmp_path, build)
    # emulate what an Excel recalculation leaves behind: a calcChain part and a cached <v> after each <f>
    dst = tmp_path / "recalculated.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                data = data.replace(b"</f><v></v>", b"</f><v>3</v>").replace(b"</f></c>", b"</f><v>3</v></c>")
            zout.writestr(item, data)
        zout.writestr("xl/calcChain.xml", '<calcChain><c r="A3" i="1"/></calcChain>')
    assert pc.cache_freshness_guard(dst).ok


# ---- formula_validity_guard / constant_formula_guard (a pair with leading_equals_guard) ----

def test_formula_validity_flags_note_saved_as_formula(tmp_path):
    def build(wb):
        ws = wb.active
        ws["B2"] = "= 100/1000 x 50e-3 x 112.09"
        ws["B3"] = "=ROUND(SUM(A1:A3)*2,1)"
        ws["B4"] = '="dilute x 10 "&A1'
        ws["B5"] = "='Other Sheet'!$B$39*C4"
    rep = pc.formula_validity_guard(_save(tmp_path, build))
    assert _cells(rep) == {"B2"}


def test_constant_formula_flags_bare_numbers_only(tmp_path):
    def build(wb):
        ws = wb.active
        ws["C1"], ws["C2"], ws["C3"] = "=9", "= 27 ", "=A1*9"
    rep = pc.constant_formula_guard(_save(tmp_path, build))
    assert _cells(rep) == {"C1", "C2"}


def test_leading_equals_and_formula_validity_are_blind_to_each_other(tmp_path):
    def build(wb):
        ws = wb.active
        ws["A1"] = "= 1 x 2"                      # stored as a (broken) formula
        ws["A2"] = "=B1*2"
        ws["A2"].data_type = "s"                    # text that only looks like a formula
    p = _save(tmp_path, build)
    assert _cells(pc.formula_validity_guard(p)) == {"A1"}
    assert "A1" not in _cells(pc.leading_equals_guard(p))


# ---- dead_input_cell_guard ------------------------------------------------

def test_dead_input_cell_is_advisory_and_respects_cross_sheet_refs(tmp_path):
    def build(wb):
        ws = wb.active
        ws.title = "Reaction Matrix"
        ws["A1"], ws["B1"] = "Premix scale", 36          # referenced by no formula -> dead
        ws["A2"], ws["B2"] = "Total volume (uL)", 50     # referenced from another sheet
        ws["A3"], ws["B3"] = "Volume per rxn", "=ROUND(2*13,2)"
        ws2 = wb.create_sheet("Pipetting Guide")
        ws2["B1"] = "='Reaction Matrix'!$B$2/10"
        for col in "ABCDEF":                              # a data table row (> 4 filled cells) is skipped
            ws[f"{col}5"] = 1 if col != "A" else "row label"
    rep = pc.dead_input_cell_guard(_save(tmp_path, build))
    assert _cells(rep) == {"B1"}
    assert rep.ok, "dead-input findings are warnings, never blocking"


# ---- opt-in Methods-folder publish -------------------------------------------

def _xlsx(tmp_path, name="exp.xlsx", payload=b"one"):
    p = tmp_path / "out" / name
    p.parent.mkdir(exist_ok=True)
    p.write_bytes(payload)
    return p


def test_publish_is_skipped_when_nothing_configured(tmp_path, monkeypatch):
    monkeypatch.delenv("LAB_METHODS_DIR", raising=False)
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path))
    ok, msg = rm.publish_to_methods(_xlsx(tmp_path))
    assert ok and "no Methods folder configured" in msg


def test_publish_copies_then_backs_up_a_different_file(tmp_path, monkeypatch):
    methods = tmp_path / "Methods"
    methods.mkdir()
    monkeypatch.setenv("LAB_METHODS_DIR", str(methods))
    src = _xlsx(tmp_path)
    assert rm.publish_to_methods(src)[0]
    assert (methods / "exp.xlsx").read_bytes() == b"one"
    ok, msg = rm.publish_to_methods(src)
    assert ok and "identical" in msg
    src.write_bytes(b"two")
    assert rm.publish_to_methods(src)[0]
    assert (methods / "exp.xlsx").read_bytes() == b"two"
    assert len(list(methods.glob("exp.BACKUP-*.xlsx"))) == 1


def test_publish_fails_loudly_when_configured_folder_is_missing(tmp_path, monkeypatch):
    monkeypatch.setenv("LAB_METHODS_DIR", str(tmp_path / "nope"))
    ok, msg = rm.publish_to_methods(_xlsx(tmp_path))
    assert not ok and "not found" in msg


def test_methods_dir_glob_from_per_computer_config(tmp_path, monkeypatch):
    monkeypatch.delenv("LAB_METHODS_DIR", raising=False)
    monkeypatch.setattr(Path, "home", classmethod(lambda cls: tmp_path))
    (tmp_path / "Cloud - Lab" / "Methods").mkdir(parents=True)
    (tmp_path / ".config").mkdir()
    (tmp_path / ".config" / "experiment-hub.json").write_text('{"methods_dir_glob": "Cloud*/Methods"}', encoding="utf-8")
    assert rm.methods_dir() == tmp_path / "Cloud - Lab" / "Methods"


# ---- `reaction_matrix.py calc` entry point ------------------------------------

def _calc(*args):
    return subprocess.run([sys.executable, str(SK / "reaction_matrix.py"), "calc", *args],
                          capture_output=True, text=True, encoding="utf-8", errors="replace")


def test_calc_topup_and_compound_dilution_factor():
    top = _calc("topup", "--target", "31.03", "--have", "29.53")
    assert top.returncode == 0 and "1.5" in top.stdout
    df = _calc("df", "--stages", "2", "10")
    assert df.returncode == 0 and "20" in df.stdout


def test_calc_impossible_question_exits_2():
    assert _calc("topup", "--target", "10", "--have", "12").returncode == 2

#!/usr/bin/env python3
"""
Order-sheet writers
===================
File/text renderers behind ``PrimerOrderSheet`` (order_sheet.py).

Each writer is a plain function of the primer entries (plus the summary
statistics where the sheet shows them) and an output path, so a layout can be
read and changed without opening the data-model class. ``PrimerOrderSheet``
keeps its public ``to_*`` methods and delegates here; nothing outside the
package needs to import this module.

Layout changes are pinned by tests/test_characterization.py (golden cell
contents, styles and column widths of every sheet).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:  # pragma: no cover - typing only, avoids an import cycle
    from .order_sheet import PrimerEntry


# Mapping to the Macrogen order-sheet format
SCALE_TO_UMOL = {
    "25 nmol": 0.025,
    "50 nmol": 0.05,
    "100 nmol": 0.1,
    "1 umol": 1.0,
}

PURIFICATION_TO_MACROGEN = {
    "Desalting": "MOPC",
    "PAGE": "PAGE",
    "HPLC": "HPLC",
}

MACROGEN_OLIGO_HEADERS = ["No.", "Oligo Name", "5` - Oligo Seq - 3`", "Amount", "Purification"]
MACROGEN_OLIGO_ROWS = 1000  # rows in the Macrogen template, blank rows included


# ── XLSX: Macrogen Order + Summary + QC ───────────────────────────────────

def write_order_xlsx(entries: list["PrimerEntry"], summary_data: dict, output_path: Path) -> Path:
    """Three-sheet workbook: "Macrogen Order", "Summary", "QC"."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws_order = wb.active
    ws_order.title = "Macrogen Order"
    _fill_order_sheet(ws_order, entries)
    _fill_summary_sheet(wb.create_sheet("Summary"), summary_data)
    _fill_qc_sheet(wb.create_sheet("QC"), entries)

    wb.save(str(output_path))
    return output_path


def _fill_order_sheet(ws_order, entries: list["PrimerEntry"]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers = [
        "No.", "Primer Name", "Sequence (5'->3')", "Scale",
        "Purification", "Length (nt)", "Notes",
    ]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col_idx, header in enumerate(headers, 1):
        cell = ws_order.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row_idx, entry in enumerate(entries, 2):
        ws_order.cell(row=row_idx, column=1, value=row_idx - 1)
        ws_order.cell(row=row_idx, column=2, value=entry.name)
        seq_cell = ws_order.cell(row=row_idx, column=3, value=entry.sequence)
        seq_cell.font = Font(name="Consolas", size=10)
        ws_order.cell(row=row_idx, column=4, value=entry.scale.value)
        ws_order.cell(row=row_idx, column=5, value=entry.purification.value)
        ws_order.cell(row=row_idx, column=6, value=entry.length)
        ws_order.cell(row=row_idx, column=7, value=entry.notes)

    col_widths = [6, 25, 60, 12, 12, 12, 20]
    for i, width in enumerate(col_widths, 1):
        ws_order.column_dimensions[get_column_letter(i)].width = width


def _fill_summary_sheet(ws_summary, summary_data: dict) -> None:
    from openpyxl.styles import Font

    summary_rows = [
        ("Total Primers", summary_data["total_primers"]),
        ("Total Length (nt)", summary_data["total_length_nt"]),
        ("Estimated Cost (KRW)", f"{summary_data['estimated_cost_krw']:,}"),
        ("Unique Experiments", ", ".join(summary_data["unique_experiments"]) or "-"),
    ]

    ws_summary.cell(row=1, column=1, value="Metric").font = Font(bold=True)
    ws_summary.cell(row=1, column=2, value="Value").font = Font(bold=True)

    for row_idx, (metric, value) in enumerate(summary_rows, 2):
        ws_summary.cell(row=row_idx, column=1, value=metric)
        ws_summary.cell(row=row_idx, column=2, value=value)

    # Scale counts, then purification counts (each a titled two-column block)
    row_offset = len(summary_rows) + 3
    _write_count_block(ws_summary, row_offset, "Scale", summary_data["scale_counts"])
    row_offset2 = row_offset + len(summary_data["scale_counts"]) + 2
    _write_count_block(ws_summary, row_offset2, "Purification", summary_data["purification_counts"])

    ws_summary.column_dimensions["A"].width = 25
    ws_summary.column_dimensions["B"].width = 30


def _write_count_block(ws, row_offset: int, title: str, counts: dict) -> None:
    from openpyxl.styles import Font

    ws.cell(row=row_offset, column=1, value=title).font = Font(bold=True)
    ws.cell(row=row_offset, column=2, value="Count").font = Font(bold=True)
    for i, (label, count) in enumerate(counts.items(), 1):
        ws.cell(row=row_offset + i, column=1, value=label)
        ws.cell(row=row_offset + i, column=2, value=count)


def _fill_qc_sheet(ws_qc, entries: list["PrimerEntry"]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    qc_headers = ["Primer Name", "Tm", "GC%", "QC Verdict", "Experiment"]
    for col_idx, header in enumerate(qc_headers, 1):
        cell = ws_qc.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    verdict_fills = {
        "PASS": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "WARNING": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
        "FAIL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    }

    for row_idx, entry in enumerate(entries, 2):
        ws_qc.cell(row=row_idx, column=1, value=entry.name)
        ws_qc.cell(row=row_idx, column=2, value=entry.tm)
        ws_qc.cell(row=row_idx, column=3, value=entry.gc)
        verdict_cell = ws_qc.cell(row=row_idx, column=4, value=entry.qc_verdict or "-")
        if entry.qc_verdict in verdict_fills:
            verdict_cell.fill = verdict_fills[entry.qc_verdict]
        ws_qc.cell(row=row_idx, column=5, value=entry.experiment)

    qc_col_widths = [25, 10, 10, 15, 25]
    for i, width in enumerate(qc_col_widths, 1):
        ws_qc.column_dimensions[get_column_letter(i)].width = width


# ── Macrogen Oligo order (.xls via xlwt, .xlsx fallback via openpyxl) ─────

def write_macrogen_oligo_xls(entries: list["PrimerEntry"], output_path: Path) -> Path:
    """Macrogen Oligo order sheet (.xls BIFF8) via xlwt."""
    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Sheet")

    header_style = xlwt.easyxf("font: bold on; align: horiz center")
    seq_style = xlwt.easyxf("font: name Consolas, height 200")

    # Column widths (1/256 character units)
    ws.col(0).width = 256 * 6    # No.
    ws.col(1).width = 256 * 35   # Oligo Name
    ws.col(2).width = 256 * 60   # Sequence
    ws.col(3).width = 256 * 10   # Amount
    ws.col(4).width = 256 * 14   # Purification

    for col, h in enumerate(MACROGEN_OLIGO_HEADERS):
        ws.write(0, col, h, header_style)

    # Data + blank rows (1000 rows total)
    for row_num in range(1, MACROGEN_OLIGO_ROWS + 1):
        ws.write(row_num, 0, row_num)

        if row_num <= len(entries):
            entry = entries[row_num - 1]
            ws.write(row_num, 1, entry.name)
            ws.write(row_num, 2, entry.sequence, seq_style)
            ws.write(row_num, 3, SCALE_TO_UMOL.get(entry.scale.value, 0.05))
            ws.write(row_num, 4, PURIFICATION_TO_MACROGEN.get(entry.purification.value, "MOPC"))

    wb.save(str(output_path))
    return output_path


def write_macrogen_oligo_xlsx(entries: list["PrimerEntry"], output_path: Path) -> Path:
    """Macrogen Oligo order sheet (.xlsx) via openpyxl (fallback)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"

    for col_idx, header in enumerate(MACROGEN_OLIGO_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_num in range(1, MACROGEN_OLIGO_ROWS + 1):
        row_idx = row_num + 1
        ws.cell(row=row_idx, column=1, value=row_num)

        if row_num <= len(entries):
            entry = entries[row_num - 1]
            ws.cell(row=row_idx, column=2, value=entry.name)
            seq_cell = ws.cell(row=row_idx, column=3, value=entry.sequence)
            seq_cell.font = Font(name="Consolas", size=10)
            ws.cell(row=row_idx, column=4, value=SCALE_TO_UMOL.get(entry.scale.value, 0.05))
            ws.cell(row=row_idx, column=5,
                    value=PURIFICATION_TO_MACROGEN.get(entry.purification.value, "MOPC"))

    for letter, width in zip("ABCDE", (6, 35, 60, 10, 14)):
        ws.column_dimensions[letter].width = width

    wb.save(str(output_path))
    return output_path


# ── Macrogen Standard Sequencing order (.xlsx) ────────────────────────────

def write_macrogen_seq_xlsx(sample_primer_pairs: list[dict], output_path: Path) -> Path:
    """Macrogen Standard Sequencing order sheet; see PrimerOrderSheet.to_macrogen_seq."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    _fill_seq_order_sheet(ws, sample_primer_pairs)
    _fill_seq_reference_sheet(wb.create_sheet("Sheet2"))

    wb.save(str(output_path))
    return output_path


def _fill_seq_order_sheet(ws, sample_primer_pairs: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    # Row 1: notice
    ws.cell(
        row=1, column=1,
        value="     ※ Only English Alphabet (either capital small letters), "
              "digit 0~9, a hypen (-) or under bar (_) is allowed "
              "without any blanks.",
    )

    # Row 2: group header
    ws.cell(row=2, column=1, value="#")
    ws.cell(row=2, column=2, value="Reaction Information")
    ws.cell(row=2, column=4, value="Sample Information")
    ws.cell(row=2, column=9, value="Primer Information")
    for col in [1, 2, 4, 9]:
        ws.cell(row=2, column=col).font = Font(bold=True)

    # Row 3: column header
    col_headers = [
        "#", "Sample Name *", "Primer Name *",
        "Sample Concentration (ng/ul)", "Plate Name", "Well Position",
        "Product Size(bp)", "Target Size(bp)",
        "Primer Sequence(5 to 3)", "Primer Concentration (pmol/ul)",
    ]
    for col_idx, header in enumerate(col_headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Data rows (starting at row 4)
    total_rows = max(len(sample_primer_pairs), 1000)
    for row_num in range(1, total_rows + 1):
        row_idx = row_num + 3
        ws.cell(row=row_idx, column=1, value=row_num)

        if row_num <= len(sample_primer_pairs):
            sp = sample_primer_pairs[row_num - 1]
            ws.cell(row=row_idx, column=2, value=sp.get("sample_name", ""))
            ws.cell(row=row_idx, column=3, value=sp.get("primer_name", ""))
            ws.cell(row=row_idx, column=4, value=sp.get("sample_conc"))
            ws.cell(row=row_idx, column=5, value=sp.get("plate_name", ""))
            ws.cell(row=row_idx, column=6, value=sp.get("well_position", ""))
            ws.cell(row=row_idx, column=7, value=sp.get("product_size"))
            ws.cell(row=row_idx, column=8, value=sp.get("target_size"))
            seq_cell = ws.cell(row=row_idx, column=9, value=sp.get("primer_seq", ""))
            seq_cell.font = Font(name="Consolas", size=10)
            ws.cell(row=row_idx, column=10, value=sp.get("primer_conc"))

    widths = [5, 25, 25, 15, 12, 12, 12, 12, 45, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fill_seq_reference_sheet(ws2) -> None:
    from openpyxl.styles import Font

    ws2.cell(row=2, column=1, value="Product Size(bp)").font = Font(bold=True)
    ws2.cell(row=2, column=2, value="Sample Align").font = Font(bold=True)
    ws2.cell(row=3, column=1, value="600bp Over")
    ws2.cell(row=3, column=2, value="Vertical")
    ws2.cell(row=4, column=1, value="600bp Less")
    ws2.cell(row=4, column=2, value="Horizontal")


# ── Markdown ──────────────────────────────────────────────────────────────

def render_markdown(project_name: str, entries: list["PrimerEntry"], summary_data: dict) -> str:
    """Markdown order sheet: order list, QC summary and totals."""
    lines: list[str] = []
    lines.append(f"# Primer Order: {project_name}")
    lines.append("")
    lines.append(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append("")
    lines.extend(_md_order_list(entries))
    lines.extend(_md_qc_summary(entries))
    lines.extend(_md_totals(summary_data))
    return "\n".join(lines)


def _md_order_list(entries: list["PrimerEntry"]) -> list[str]:
    lines = ["## Order List", "",
             "| No. | Primer Name | Sequence (5'->3') | Scale | Purification | Length (nt) | Notes |",
             "|-----|-------------|-------------------|-------|--------------|------------|-------|"]
    for i, entry in enumerate(entries, 1):
        lines.append(
            f"| {i} | {entry.name} | `{entry.sequence}` | "
            f"{entry.scale.value} | {entry.purification.value} | "
            f"{entry.length} | {entry.notes} |"
        )
    lines.append("")
    return lines


def _md_qc_summary(entries: list["PrimerEntry"]) -> list[str]:
    lines = ["## QC Summary", "",
             "| Primer Name | Tm | GC% | QC Verdict | Experiment |",
             "|-------------|----|-----|------------|------------|"]
    for entry in entries:
        tm_str = f"{entry.tm:.1f}" if entry.tm is not None else "-"
        gc_str = f"{entry.gc:.1f}" if entry.gc is not None else "-"
        lines.append(
            f"| {entry.name} | {tm_str} | {gc_str} | "
            f"{entry.qc_verdict or '-'} | {entry.experiment} |"
        )
    lines.append("")
    return lines


def _md_totals(summary_data: dict) -> list[str]:
    lines = ["## Summary", "",
             f"- Total primers: {summary_data['total_primers']}",
             f"- Total length: {summary_data['total_length_nt']} nt",
             f"- Estimated cost: {summary_data['estimated_cost_krw']:,} KRW"]
    if summary_data["unique_experiments"]:
        lines.append(f"- Experiments: {', '.join(summary_data['unique_experiments'])}")
    lines.append("")
    return lines

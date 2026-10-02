"""Characterization snapshots for the primer-design refactor (behaviour-preserving splits).

Three big units were split without changing behaviour:

  * order_sheet.PrimerOrderSheet      -> exported cell contents / styles / widths
  * cloning_report.generate_vector_construct_map -> artists of the matplotlib figure
  * restriction_cloning_mode.RestrictionCloningDesigner.design -> the whole result dict

This module builds the *scenarios* and the *dump* of each output; the committed
golden files in tests/golden/ were produced from the code BEFORE the split by
`python tests/golden/regen_golden.py`. test_characterization.py compares a fresh
dump with the golden file, so a refactor that changes any exported cell, any
drawn artist or any design field fails loudly.

Figures are compared as artists (positions, strings, colours, widths), not as
pixels, so the snapshot does not depend on the matplotlib/freetype build.
Dates ("Generated: ...") are normalised away.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
GOLDEN = HERE / "golden"
SRC = HERE.parent / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from primer_design.order_sheet import (  # noqa: E402
    PrimerOrderSheet, PrimerScale, Purification,
)
from primer_design.restriction_cloning_mode import RestrictionCloningDesigner  # noqa: E402

# --------------------------------------------------------------------------
# Order-sheet scenario
# --------------------------------------------------------------------------

_R1 = {
    "f_full": "ATGCGTAACCTGGCGATCAAGCTG", "r_full": "CAGCTTGATCGCCAGGTTACGCAT",
    "f_tm": 62.5, "r_tm": 62.5, "f_gc": 54.2, "r_gc": 54.2,
    "f_qc": {"verdict": "PASS"}, "r_qc": {"verdict": "PASS"},
}
_R2 = {
    "f_full": "GGATCCATGAAAGCTGCCATTGTTCTG", "r_full": "CTCGAGTTATTCAACATCGGTCGC",
    "f_tm": 63.1, "r_tm": 61.8, "f_gc": 51.9, "r_gc": 50.0,
    "f_qc": {"verdict": "PASS"}, "r_qc": {"verdict": "WARNING"},
    "re_5prime": "BamHI", "re_3prime": "XhoI",
}
_R3 = {
    "f_full": "ggatccatgaaagctgccattgttctgaaa", "r_full": "ctcgagttattcaacatcggtcgctttt",
    "f_tm": None, "r_tm": 55.0, "f_gc": None, "r_gc": 40.0,
    "f_qc": {"verdict": "FAIL"}, "r_qc": "not-a-dict",
}


def build_order_sheet() -> PrimerOrderSheet:
    s = PrimerOrderSheet(project_name="char_order")
    s.add_from_design_result(_R1, experiment="site-directed mutagenesis",
                             parent="GeneX_WT", mutation="A123T")
    s.add_from_design_result(_R2, experiment="restriction cloning")
    s.add_from_design_result(_R3, experiment="restriction cloning", name_prefix="RE")
    s.add_custom_primer(name="T7_promoter_F", sequence="TAATACGACTCACTATAGGG",
                        experiment="sequencing", direction="F",
                        notes="standard sequencing primer")
    s.add_custom_primer(name="big_scale", sequence="acgtacgtacgtacgtacgtacgtacgt",
                        scale=PrimerScale.UMOL_1, purification=Purification.HPLC,
                        tm=60.04, gc=50.0, experiment="misc")
    s.add_custom_primer(name="page_primer", sequence="ACGT" * 5,
                        scale=PrimerScale.NMOL_25, purification=Purification.PAGE)
    s.add_custom_primer(name="dup_of_T7", sequence="taatacgactcactatagggg"[:-1])
    return s


SEQ_PAIRS = [
    {"sample_name": "pET28a_GeneX", "primer_name": "T7", "sample_conc": None,
     "plate_name": "P1", "well_position": "A1", "product_size": 1200,
     "target_size": 800, "primer_seq": "taatacgactcactataggg", "primer_conc": 10.0},
    {"sample_name": "pET28a_GeneY", "primer_name": "T7term"},
]


def _style(cell) -> list:
    fill = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type else None
    return [bool(cell.font.b), fill, cell.font.name, cell.font.sz, cell.alignment.horizontal,
            bool(cell.alignment.wrap_text)]


def dump_xlsx(path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(path))
    out = {"sheetnames": wb.sheetnames, "sheets": {}}
    for ws in wb.worksheets:
        cells, flat = [], []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                flat.append(f"{c.coordinate}={c.value!r}")
                if c.row <= 12 or c.column > 1:
                    cells.append([c.coordinate, c.value, _style(c)])
        out["sheets"][ws.title] = {
            "dims": [ws.max_row, ws.max_column],
            "widths": {k: v.width for k, v in sorted(ws.column_dimensions.items())},
            "n_values": len(flat),
            "sha": hashlib.sha256("\n".join(flat).encode()).hexdigest(),
            "cells": cells,
        }
    return out


def dump_xls(path: Path) -> dict | None:
    try:
        import xlrd
    except ImportError:
        return None
    ws = xlrd.open_workbook(str(path)).sheet_by_name("Sheet")
    rows = [[ws.cell_value(r, c) for c in range(ws.ncols)] for r in range(ws.nrows)]
    flat = json.dumps(rows, ensure_ascii=False)
    return {"nrows": ws.nrows, "ncols": ws.ncols, "head": rows[:9],
            "sha": hashlib.sha256(flat.encode()).hexdigest(),
            "col_widths": [ws.colinfo_map[c].width if c in ws.colinfo_map else None
                           for c in range(ws.ncols)]}


def _norm_md(text: str) -> str:
    return re.sub(r"^Generated: .*$", "Generated: <date>", text, flags=re.M)


def snapshot_order_sheet(tmp: Path) -> dict:
    s = build_order_sheet()
    snap: dict = {"dedup_removed": None}
    snap["summary_before"] = s.summary()
    snap["dataframe"] = json.loads(s.to_dataframe().to_json(orient="records"))
    snap["xlsx"] = dump_xlsx(s.to_xlsx(tmp / "o.xlsx"))
    snap["macrogen_oligo_xlsx"] = dump_xlsx(s.to_macrogen_oligo(tmp / "m.xlsx"))
    snap["macrogen_oligo_xls"] = dump_xls(s.to_macrogen_oligo(tmp / "m.xls")) \
        if _has("xlwt") else None
    snap["macrogen_seq"] = dump_xlsx(s.to_macrogen_seq(SEQ_PAIRS, tmp / "seq.xlsx"))
    snap["csv"] = s.to_csv(tmp / "o.csv").read_bytes().decode("utf-8-sig")
    snap["markdown"] = _norm_md(s.to_markdown(tmp / "o.md").read_text(encoding="utf-8"))
    snap["dedup_removed"] = s.deduplicate()
    snap["summary_after"] = s.summary()
    empty = PrimerOrderSheet(project_name="empty")
    snap["empty_xlsx"] = dump_xlsx(empty.to_xlsx(tmp / "e.xlsx"))
    snap["empty_markdown"] = _norm_md(empty.to_markdown(tmp / "e.md").read_text(encoding="utf-8"))
    return snap


def _has(mod: str) -> bool:
    try:
        __import__(mod)
        return True
    except ImportError:
        return False


# --------------------------------------------------------------------------
# Restriction-cloning design scenarios
# --------------------------------------------------------------------------

TEST_INSERT = (
    "ATGAAAGCTGCCATTGTTCTGAGCGAATCCGGTGTGCACGAATGTCCGAAAGAATCCAAACTG"
    "GTTCCGGCACTGAACGGTCTGGAAATCGAAGATGAAGGCGTTATCCCGGAATTCTTCAAGGGC"
    "AAACTGGATCGCGGCTTCAAAGCGATCCTGAACAATGCGAAAGATTGGAGCCGCGTGGAAAAC"
    "TACCAGCCGGATCTGATCAACGAACTGAAAGCGAAAGCGACCGATGTGGAATAA"
)
WITH_BAMHI = ("ATGAAAGCTGCCATTGTTCTGAGCGAATCCGGATCCGCACGAATGTCCGAAAGAATCCAAAC"
              "TGGTTCCGGCACTGAACGGTCTGGAAATCGAAGATGAAGGCGTTATCCCGGAATTCTTCAAG")
# Found by scanning random inserts: each one drives a different branch of design().
AT_RICH_F_TM = "ATGCTATTAAAACTATAATCAACAAATGATTATATGCCGACAATATATGCAATATAATCCATCTAA"
AT_RICH_R_TM = ("ATGTTTGAATAAAAGTAAGATATACAGCCATAATATAATTTATATTATTGTAGCTAGAGAATTCAAAAAAAC"
                "ACATCTTTTTGTTTAGGAGTACAGACAAGATTCAAAGTTCCTTATTTTAAATTATTTAAGAATATGAATTATTTACACCAGTAA")
R_FAILS = "ATGCCGGCCGGCCTAATTTATTTTAATTTTATATATTTATATAATTTAATTATATTTACTACATAA"
WITH_XHOI = ("ATGAAAGCTGCCATTGTTCTGAGCGAATCCCTCGAGGCACGAATGTCCGAAAGAATCCAAAC"
             "TGGTTCCGGCACTGAACGGTCTGGAAATCGAAGATGAAGGCGTTATCCCGGAATTCTTCAAG")
F_FAILS = "ATGATTTAAAAATTTAGAGTTTTAAATTAAAAGATAATGTATATAATATAATGATTAAGATAAAATATAAATTATTTTATACCGCCCCCCCCCTAA"
F_HAIRPIN_STUCK = "ATGGTATACCGCACTAGCACTTGCAGCGGAAGGTATAACCAACAGTTCCCCTCGTGTATACATTGAGACTGTCAGCATTACCCGCCCACGACCTAA"
R_HAIRPIN_STUCK = "ATGTACTGGTGGTATAATGGTAAATCTATCCCACAGTTGAGATATACTCCAGGGGTTTCCCCCTAA"
F_HAIRPIN_FIXED = "ATGACGCCATTAAGGAAGCGGATAAATTCTTAGGTAATAGTTTTCCGTAGAAGTAAGTTGTATGTTGATAAATCTAGATATCTTCACCATTTTTAA"
R_HAIRPIN_FIXED = "ATGCCGGGGCCACCCCGTATCACTGCCCTATACTTCCGACGTCAAGCGGCTCGGGCGCCCAGCGGGGCCGGGGGCCTCGCCTGGGCCGTAATTTAA"

DESIGN_SCENARIOS = {
    "bamhi_xhoi_pet28a": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="XhoI",
                              vector_name="pET-28a(+)", include_start_codon=True,
                              include_stop_codon=False),
    "nhei_noti_pet21a_stop": dict(insert_seq=TEST_INSERT, re_5prime="NheI", re_3prime="NotI",
                                  vector_name="pET-21a(+)", include_stop_codon=True,
                                  stop_codon="TAA"),
    "internal_site": dict(insert_seq=WITH_BAMHI, re_5prime="BamHI", re_3prime="XhoI"),
    "internal_3prime_site": dict(insert_seq=WITH_XHOI, re_5prime="BamHI", re_3prime="XhoI"),
    "custom_protection_literal": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="XhoI",
                                      protection_bases_5="AATTCC", protection_bases_3=6),
    "auto_protection": dict(insert_seq=TEST_INSERT, re_5prime="NotI", re_3prime="XhoI",
                            protection_bases_5=None, protection_bases_3=None),
    "ndei_atg": dict(insert_seq=TEST_INSERT, re_5prime="NdeI", re_3prime="XhoI",
                     vector_name="pET-28a(+)", include_start_codon=True),
    "ecorv_blunt": dict(insert_seq=TEST_INSERT, re_5prime="EcoRV", re_3prime="XhoI"),
    "compatible_overhang": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="BglII"),
    "same_enzyme": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="BamHI"),
    "spacers_lowercase_spaces": dict(insert_seq=" ".join([TEST_INSERT[:30].lower(), TEST_INSERT[30:]]),
                                     re_5prime="BamHI", re_3prime="XhoI",
                                     spacer_5prime="GC", spacer_3prime="AA",
                                     min_ann_len=20, max_ann_len=28, target_tm=60.0),
    "unknown_vector_frame_check_fails": dict(insert_seq=TEST_INSERT, re_5prime="BamHI",
                                             re_3prime="XhoI", vector_name="no-such-vector"),
    "frame_check_disabled": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="XhoI",
                                 vector_name="pET-28a(+)", auto_frame_check=False),
    "at_rich_forward_relaxed_tm": dict(insert_seq=AT_RICH_F_TM, re_5prime="BamHI",
                                       re_3prime="XhoI", vector_name="pET-28a(+)"),
    "at_rich_reverse_relaxed_tm": dict(insert_seq=AT_RICH_R_TM, re_5prime="BamHI",
                                       re_3prime="XhoI", vector_name="pET-28a(+)"),
    "forward_hairpin_unresolved": dict(insert_seq=F_HAIRPIN_STUCK, re_5prime="BamHI",
                                       re_3prime="XhoI", vector_name="pET-28a(+)"),
    "reverse_hairpin_unresolved": dict(insert_seq=R_HAIRPIN_STUCK, re_5prime="BamHI",
                                       re_3prime="XhoI", vector_name="pET-28a(+)"),
    "forward_hairpin_adjusted": dict(insert_seq=F_HAIRPIN_FIXED, re_5prime="BamHI",
                                     re_3prime="XhoI", vector_name="pET-28a(+)"),
    "reverse_hairpin_adjusted": dict(insert_seq=R_HAIRPIN_FIXED, re_5prime="BamHI",
                                     re_3prime="XhoI", vector_name="pET-28a(+)"),
}

DESIGN_ERRORS = {
    "empty_insert": dict(insert_seq="", re_5prime="BamHI", re_3prime="XhoI"),
    "non_atgc": dict(insert_seq="ATGNNN", re_5prime="BamHI", re_3prime="XhoI"),
    "unknown_5prime_re": dict(insert_seq=TEST_INSERT, re_5prime="Nope", re_3prime="XhoI"),
    "hf_variant_is_not_a_design_enzyme": dict(insert_seq=TEST_INSERT, re_5prime="BamHI-HF",
                                              re_3prime="XhoI"),
    "unknown_3prime_re": dict(insert_seq=TEST_INSERT, re_5prime="BamHI", re_3prime="Nope"),
    "reverse_annealing_fails": dict(insert_seq=R_FAILS, re_5prime="BamHI", re_3prime="XhoI"),
    "long_primer_overflows_primer3": dict(insert_seq=TEST_INSERT, re_5prime="BamHI",
                                          re_3prime="XhoI", spacer_5prime="GC" * 20),
    "forward_annealing_fails": dict(insert_seq=F_FAILS, re_5prime="BamHI", re_3prime="XhoI"),
}


def _jsonable(obj):
    return json.loads(json.dumps(obj, sort_keys=True, default=str))


def snapshot_designs() -> dict:
    d = RestrictionCloningDesigner()
    out = {"results": {}, "errors": {}}
    for name, kw in DESIGN_SCENARIOS.items():
        out["results"][name] = _jsonable(d.design(**kw))
    for name, kw in DESIGN_ERRORS.items():
        try:
            d.design(**kw)
        except Exception as exc:  # noqa: BLE001 - the type and message are the snapshot
            out["errors"][name] = [type(exc).__name__, str(exc)]
        else:
            out["errors"][name] = None
    return out


# --------------------------------------------------------------------------
# Circular construct map: dump the matplotlib artists
# --------------------------------------------------------------------------

def _r(v, n=4):
    try:
        return round(float(v), n)
    except (TypeError, ValueError):
        return v


def dump_figure(fig) -> dict:
    out = []
    for ax in fig.axes:
        entry = {"xlim": [_r(x) for x in ax.get_xlim()], "ylim": [_r(y) for y in ax.get_ylim()],
                 "axis_off": not ax.axison, "lines": [], "patches": [], "texts": []}
        for ln in ax.lines:
            xs, ys = ln.get_xdata(), ln.get_ydata()
            pts = [[_r(x), _r(y)] for x, y in zip(xs, ys)]
            entry["lines"].append({
                "n": len(pts), "first": pts[0], "last": pts[-1],
                "sha": hashlib.sha256(json.dumps(pts).encode()).hexdigest()[:16],
                "color": ln.get_color(), "lw": _r(ln.get_linewidth()),
                "z": ln.get_zorder(), "cap": ln.get_solid_capstyle(),
            })
        for p in ax.patches:
            d = {"type": type(p).__name__, "fill": p.get_fill(), "z": p.get_zorder(),
                 "lw": _r(p.get_linewidth()), "fc": list(map(_r, p.get_facecolor())),
                 "ec": list(map(_r, p.get_edgecolor()))}
            if hasattr(p, "center"):
                d["center"] = [_r(c) for c in p.center]
                d["radius"] = _r(p.radius)
            if hasattr(p, "get_x"):
                d["xy"] = [_r(p.get_x()), _r(p.get_y())]
                d["wh"] = [_r(p.get_width()), _r(p.get_height())]
            entry["patches"].append(d)
        for t in ax.texts:
            s = re.sub(r"Generated: .*", "Generated: <date>", t.get_text())
            d = {"type": type(t).__name__, "text": s, "pos": [_r(v) for v in t.get_position()],
                 "size": _r(t.get_fontsize()), "color": str(t.get_color()),
                 "weight": str(t.get_fontweight()), "ha": t.get_ha(), "va": t.get_va(),
                 "family": list(t.get_fontfamily()), "z": t.get_zorder()}
            if hasattr(t, "xy"):
                d["xy"] = [_r(v) for v in t.xy]
                d["xytext"] = [_r(v) for v in getattr(t, "xyann", ())]
                d["arrowstyle"] = str(getattr(t, "arrowprops", None))
            entry["texts"].append(d)
        out.append(entry)
    return {"size": [_r(v) for v in fig.get_size_inches()], "axes": out}


MAP_SCENARIOS = {
    "pet28a_tags_frame_ok": dict(
        vector_name="pET-28a(+)", gene_name="GeneX", insert_len=900,
        re_5prime="BamHI", re_3prime="XhoI",
        frame_check={"topology": "[N-His6][Thrombin][T7-tag]--RE5--[insert]--RE3--[C-His6]",
                     "in_frame_5prime": True, "in_frame_3prime": True}),
    "pet21a_ctags_out_of_frame": dict(
        vector_name="pET-21a(+)", gene_name="GeneY", insert_len=1500,
        re_5prime="NdeI", re_3prime="XhoI", include_stop=True,
        frame_check={"topology": "[C-His6]:OUT-OF-FRAME [S-tag]", "in_frame_5prime": False,
                     "in_frame_3prime": False}),
    "pmal_signal_peptide": dict(
        vector_name="pMAL-c6T", gene_name="GeneZ", insert_len=600,
        re_5prime="NotI", signal_peptide={"has_signal_peptide": True, "cleavage_site_estimate": 24}),
    "pet28a_c_tags_in_frame": dict(
        vector_name="pET-28a(+)", gene_name="GeneW", insert_len=750, re_5prime="NcoI",
        re_3prime="XhoI", frame_check={"topology": "RE5-[insert]-RE3-[C-His6][S-tag]",
                                       "in_frame_5prime": True, "in_frame_3prime": False}),
    "duet_mcs2_minimal": dict(vector_name="pETDuet-1:MCS2", insert_len=300, re_3prime="XhoI"),
    "unknown_vector_fallback": dict(vector_name="mystery(+)", gene_name="G", insert_len=0),
    "signal_peptide_zero_cleavage": dict(
        vector_name="pET-28a(+)", insert_len=450,
        signal_peptide={"has_signal_peptide": True, "cleavage_site_estimate": 0}),
}


def snapshot_construct_maps(tmp: Path) -> dict:
    import matplotlib.figure as mfig
    from primer_design.cloning_report import generate_vector_construct_map

    captured: list = []
    original = mfig.Figure.savefig

    def spy(self, *a, **k):
        captured.append(dump_figure(self))
        return original(self, *a, **k)

    mfig.Figure.savefig = spy
    try:
        out = {}
        for name, kw in MAP_SCENARIOS.items():
            captured.clear()
            path = generate_vector_construct_map(output_path=tmp / f"{name}.png", **kw)
            assert path.exists() and path.stat().st_size > 1000
            assert len(captured) == 1
            out[name] = captured[0]
    finally:
        mfig.Figure.savefig = original
    return out

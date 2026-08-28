#!/usr/bin/env python3
"""Extract tables, figures, and body text from a paper PDF / Word (.docx) in one pass.

Output lands beside the input file (or in --output-dir) as a
`<filename>_extracted/` folder:

    <stem>_extracted/
    ├── table_p{page}_{idx}.xlsx     individual tables
    ├── all_tables.xlsx              all tables (split into sheets)
    ├── fig_p{page}_{idx}.png        images embedded in the body
    ├── fig_p{page}_fullpage.png     pages with no embedded image get rendered whole
    ├── full_text.md                 the full body text (Markdown)
    └── extraction_report.json       extraction summary (counts, warnings, skipped)

Usage:
    python scripts/extract_paper_assets.py "<paper.pdf>"
    python scripts/extract_paper_assets.py "<manuscript.docx>" --output-dir "<out>"
    python scripts/extract_paper_assets.py "<paper.pdf>" --no-fullpage --dpi 200

Design principles:
- **Report what's missing as missing.** If a library is unavailable or a page
  is empty, don't silently move on — record it in `extraction_report.json`'s
  `warnings` / `skipped`. Zero extractions that still look like success is
  the most dangerous failure mode for this tool.
- Optional dependencies degrade gracefully: **having even one still gets you
  that much** (tables only, or figures only, is fine).
"""
from __future__ import annotations

import argparse
import io
import json
import sys
import traceback
from pathlib import Path

# Windows' default console is cp949 and dies on Korean/symbol output. Force UTF-8.
# Use reconfigure instead of TextIOWrapper — the wrapper takes ownership of the
# underlying stream, so once it's GC'd after import, it closes the caller's
# stdout too (measured).
for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

MIN_IMAGE_PX = 50          # images smaller than this are treated as icons/decoration and dropped
DEFAULT_DPI = 300


class Report:
    """A container that collects extraction results. Records failures as faithfully as successes."""

    def __init__(self, source: Path, outdir: Path):
        self.source = str(source)
        self.outdir = str(outdir)
        self.tables: list[str] = []
        self.figures: list[str] = []
        self.text_file: str | None = None
        self.text_chars = 0
        self.warnings: list[str] = []
        self.skipped: list[str] = []

    def warn(self, msg: str) -> None:
        self.warnings.append(msg)
        print(f"  [warning] {msg}")

    def skip(self, msg: str) -> None:
        self.skipped.append(msg)
        print(f"  [skipped] {msg}")

    def to_dict(self) -> dict:
        return {
            "source": self.source,
            "output_dir": self.outdir,
            "table_count": len(self.tables),
            "figure_count": len(self.figures),
            "tables": self.tables,
            "figures": self.figures,
            "text_file": self.text_file,
            "text_chars": self.text_chars,
            "warnings": self.warnings,
            "skipped": self.skipped,
        }


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------

def extract_pdf_tables(pdf: Path, outdir: Path, rep: Report) -> None:
    try:
        import pdfplumber
    except ImportError:
        rep.skip("table extraction: pdfplumber is not installed (pip install pdfplumber)")
        return
    try:
        from openpyxl import Workbook
    except ImportError:
        rep.skip("table extraction: openpyxl is not installed (pip install openpyxl)")
        return

    combined = Workbook()
    combined.remove(combined.active)
    found = 0

    with pdfplumber.open(str(pdf)) as doc:
        for pno, page in enumerate(doc.pages, 1):
            try:
                tables = page.extract_tables()
            except Exception as exc:                      # noqa: BLE001
                rep.warn(f"p{pno} table extraction failed: {exc}")
                continue
            for idx, table in enumerate(tables, 1):
                rows = [r for r in table if r and any(c not in (None, "") for c in r)]
                if not rows:
                    continue
                found += 1
                name = f"table_p{pno}_{idx}.xlsx"
                wb = Workbook()
                ws = wb.active
                ws.title = f"p{pno}_{idx}"[:31]
                for row in rows:
                    ws.append(["" if c is None else str(c) for c in row])
                wb.save(outdir / name)
                rep.tables.append(name)

                sheet = combined.create_sheet(f"p{pno}_{idx}"[:31])
                for row in rows:
                    sheet.append(["" if c is None else str(c) for c in row])

    if found:
        combined.save(outdir / "all_tables.xlsx")
        print(f"  extracted {found} table(s)")
    else:
        rep.warn("found no tables at all — this may be a scanned PDF or the tables may be images")


def extract_pdf_figures(pdf: Path, outdir: Path, rep: Report,
                        dpi: int, fullpage: bool) -> None:
    try:
        import fitz  # PyMuPDF
    except ImportError:
        rep.skip("figure extraction: PyMuPDF (fitz) is not installed (pip install pymupdf)")
        return

    doc = fitz.open(str(pdf))
    embedded_total = 0
    try:
        for pno in range(len(doc)):
            page = doc[pno]
            page_hits = 0
            for idx, info in enumerate(page.get_images(full=True), 1):
                xref = info[0]
                try:
                    pix = fitz.Pixmap(doc, xref)
                except Exception as exc:                  # noqa: BLE001
                    rep.warn(f"p{pno+1} failed to read image {idx}: {exc}")
                    continue
                if pix.width < MIN_IMAGE_PX or pix.height < MIN_IMAGE_PX:
                    pix = None
                    continue
                if pix.n - pix.alpha >= 4:                # CMYK -> RGB
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                name = f"fig_p{pno+1}_{idx}.png"
                pix.save(str(outdir / name))
                rep.figures.append(name)
                page_hits += 1
                embedded_total += 1
                pix = None

            # A page drawn entirely in vector graphics can only be salvaged by rendering it whole.
            if fullpage and page_hits == 0 and page.get_drawings():
                name = f"fig_p{pno+1}_fullpage.png"
                page.get_pixmap(dpi=dpi).save(str(outdir / name))
                rep.figures.append(name)
    finally:
        doc.close()

    if rep.figures:
        print(f"  extracted {len(rep.figures)} figure(s) ({embedded_total} embedded in body)")
    else:
        rep.warn("found no figures at all")


def extract_pdf_text(pdf: Path, outdir: Path, rep: Report) -> None:
    text = None
    try:
        from markitdown import MarkItDown
        text = MarkItDown().convert(str(pdf)).text_content
    except ImportError:
        rep.warn("markitdown is unavailable, falling back to pdfplumber for body text (table structure will be simpler)")
    except Exception as exc:                              # noqa: BLE001
        rep.warn(f"markitdown conversion failed ({exc}) — falling back to pdfplumber")

    if text is None:
        try:
            import pdfplumber
            with pdfplumber.open(str(pdf)) as doc:
                text = "\n\n".join((p.extract_text() or "") for p in doc.pages)
        except ImportError:
            rep.skip("body-text extraction: neither markitdown nor pdfplumber is available")
            return

    _write_text(text, outdir, rep)


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------

def extract_docx(src: Path, outdir: Path, rep: Report) -> None:
    try:
        import docx  # python-docx
    except ImportError:
        rep.skip("docx processing: python-docx is not installed (pip install python-docx)")
        return

    document = docx.Document(str(src))

    # --- tables ---
    try:
        from openpyxl import Workbook
    except ImportError:
        rep.skip("table extraction: openpyxl is unavailable")
    else:
        combined = Workbook()
        combined.remove(combined.active)
        for idx, table in enumerate(document.tables, 1):
            rows = [[c.text for c in r.cells] for r in table.rows]
            rows = [r for r in rows if any(v.strip() for v in r)]
            if not rows:
                continue
            name = f"table_p0_{idx}.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.title = f"table{idx}"[:31]
            for row in rows:
                ws.append(row)
            wb.save(outdir / name)
            rep.tables.append(name)
            sheet = combined.create_sheet(f"table{idx}"[:31])
            for row in rows:
                sheet.append(row)
        if rep.tables:
            combined.save(outdir / "all_tables.xlsx")
            print(f"  extracted {len(rep.tables)} table(s)")
        else:
            rep.warn("found no tables in the docx")

    # --- figures (pulls image parts embedded in the Word package directly) ---
    count = 0
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        try:
            blob = rel.target_part.blob
            ext = Path(rel.target_part.partname).suffix or ".png"
        except Exception as exc:                          # noqa: BLE001
            rep.warn(f"failed to read an image part: {exc}")
            continue
        count += 1
        name = f"fig_p0_{count}{ext}"
        (outdir / name).write_bytes(blob)
        rep.figures.append(name)
    if count:
        print(f"  extracted {count} figure(s)")
    else:
        rep.warn("found no figures in the docx")

    # --- body text ---
    text = None
    try:
        from markitdown import MarkItDown
        text = MarkItDown().convert(str(src)).text_content
    except Exception:                                     # noqa: BLE001
        text = "\n\n".join(p.text for p in document.paragraphs if p.text.strip())
        rep.warn("extracted body text with python-docx (no markitdown) — "
                 "if tracked changes are present, inserted/deleted content may be missing")
    _write_text(text, outdir, rep)


# ----------------------------------------------------------------------

def _write_text(text: str | None, outdir: Path, rep: Report) -> None:
    if not text or not text.strip():
        rep.warn("body text is empty — a scanned copy would need OCR")
        return
    (outdir / "full_text.md").write_text(text, encoding="utf-8")
    rep.text_file = "full_text.md"
    rep.text_chars = len(text)
    print(f"  extracted {len(text):,} characters of body text")


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Extract tables, figures, and body text from a paper PDF/docx")
    ap.add_argument("source", help="input file (.pdf or .docx)")
    ap.add_argument("--output-dir", help="output folder (default: beside the input file)")
    ap.add_argument("--dpi", type=int, default=DEFAULT_DPI,
                    help=f"full-page render resolution (default {DEFAULT_DPI})")
    ap.add_argument("--no-fullpage", action="store_true",
                    help="skip the whole-page render for pages with no embedded image")
    args = ap.parse_args()

    src = Path(args.source).expanduser().resolve()
    if not src.is_file():
        print(f"Error: file not found — {src}", file=sys.stderr)
        return 2

    suffix = src.suffix.lower()
    if suffix not in (".pdf", ".docx"):
        print(f"Error: unsupported format '{suffix}' (.pdf or .docx only)",
              file=sys.stderr)
        return 2

    outdir = (Path(args.output_dir).expanduser().resolve() if args.output_dir
              else src.parent / f"{src.stem}_extracted")
    outdir.mkdir(parents=True, exist_ok=True)

    print(f"Input: {src}")
    print(f"Output: {outdir}")

    rep = Report(src, outdir)
    try:
        if suffix == ".pdf":
            extract_pdf_tables(src, outdir, rep)
            extract_pdf_figures(src, outdir, rep, args.dpi, not args.no_fullpage)
            extract_pdf_text(src, outdir, rep)
        else:
            extract_docx(src, outdir, rep)
    except Exception:                                     # noqa: BLE001
        rep.warn("unexpected error:\n" + traceback.format_exc())

    (outdir / "extraction_report.json").write_text(
        json.dumps(rep.to_dict(), ensure_ascii=False, indent=2), encoding="utf-8")

    total = len(rep.tables) + len(rep.figures) + (1 if rep.text_file else 0)
    print(f"\nSummary: tables {len(rep.tables)} / figures {len(rep.figures)} / "
          f"body text {'present' if rep.text_file else 'absent'}")
    if rep.skipped:
        print(f"{len(rep.skipped)} step(s) skipped — see 'skipped' in extraction_report.json")
    if total == 0:
        print("Nothing was extracted. Read extraction_report.json's warnings/skipped first.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())

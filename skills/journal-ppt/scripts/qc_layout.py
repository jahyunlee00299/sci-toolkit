#!/usr/bin/env python3
"""
qc_layout.py -- geometric / visual QC for journal-ppt decks.

Catches six defect classes that the style-level QC (qc_deck.py) cannot see:

  L1 TEXT OVERFLOW  wrapped-text height estimated with real Arial metrics
                    (PIL ImageFont) vs the box, the Y=floor line and the shapes
                    below.  Text boxes and table cells.
  L2 SHAPE OVERLAP  ink-rectangle intersection among content shapes.
  L3 LOW CONTRAST   WCAG relative-luminance ratio of each run color against its
                    effective background (z-order walk).
  L4 STRUCTURE      hidden slides, content-less slides, placeholder strings.
  L5 ASPECT         picture shown at another aspect ratio than its pixels (stretched).
  L6 SLIDE ORDER    the references slide must close the talk (only Q&A, thanks, backup after it).

Usage:
    python qc_layout.py deck.pptx [--theme NAME] [--json]      (exit 1 on CRITICAL)
    from qc_layout import check_layout; findings = check_layout(prs)

Deliberately independent of deck_builder / qc_deck (no imports from them).
Font size policy (type scale) is NOT checked here.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, asdict
from functools import lru_cache
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.text.text import _Paragraph

CRITICAL = "CRITICAL"
WARNING = "WARNING"
EMU_IN = 914400.0

# ---------------------------------------------------------------- tunables --
FLOOR_IN = 7.0               # nothing but the footer may sit below this line
FLOOR_TOL_IN = 0.03          # estimator slack before crossing the floor is reported
OWN_BOX_TOL_IN = 0.04        # slack before "exceeds its own box" is reported
INVADE_MIN_XFRAC = 0.25      # x-overlap must cover >=25% of the text width to count as stacked
BLANK_INK_FRAC = 0.02        # picture overlap region with <2% dark pixels is just margin
INVADE_TOL_IN = 0.06         # how far text must reach into another shape
OVERLAP_CRIT_RATIO = 0.05    # picture/table overlap > 5% of smaller area -> CRITICAL
TEXT_OVERLAP_CRIT_RATIO = 0.15  # text-vs-text ink overlap > 15% of the smaller ink area -> CRITICAL
ASPECT_WARN = 0.02           # displayed picture aspect differs from the image by > 2% -> WARNING
ASPECT_CRIT = 0.05           # ... by > 5% -> CRITICAL (visibly stretched)
LINE_FACTOR = 1.2            # PowerPoint single line pitch = 1.2 x font size (Arial)
THIN_IN = 0.08               # shapes thinner than this are decorative lines
CONTRAST_CRIT = 3.0          # below this: always CRITICAL
CONTRAST_NORMAL = 4.5        # below this: CRITICAL for normal text, WARNING for large
LARGE_PT, LARGE_BOLD_PT = 18.0, 14.0
DEFAULT_PT = 18.0

# Theme hook: lets the themes track register its default backgrounds without
# this module importing deck_builder.  Only used where a slide declares none.
THEME_DEFAULT_BG = {"default": "FFFFFF"}

PLACEHOLDER_RE = re.compile(
    r"\blorem\b|\bipsum\b|x{4,}|click to (add|edit)|\bTBD\b|\bTODO\b|"
    r"\[(insert|placeholder)[^\]]*\]|\bFIXME\b", re.I)

REFS_RE = re.compile(r"\s*(references?|reference list|bibliography|literature cited|참고\s*문헌)(\s*\(?\s*(cont\.?|continued|\d+\s*/\s*\d+)\s*\)?)?\s*[.:]?\s*$", re.I)
# slides that may legitimately follow the references slide
AFTER_REFS_RE = re.compile(
    r"\s*(questions?|q\s*&\s*a|discussion|thank\s*you|thanks|acknowledg\w*|backup|appendix|supplement\w*|"
    r"감사|질문|토론|부록)", re.I)
TITLE_ZONE_IN = 1.6
APPENDIX_RE = re.compile(r"(appendix|backup)\b", re.I)          # a title sits above this line

FONT_DIR = Path(r"C:\Windows\Fonts")
FONT_FILES = {  # (family-lower, bold) -> file
    ("arial", False): "arial.ttf", ("arial", True): "arialbd.ttf",
    ("calibri", False): "calibri.ttf", ("calibri", True): "calibrib.ttf",
    ("times new roman", False): "times.ttf", ("times new roman", True): "timesbd.ttf",
    ("malgun gothic", False): "malgun.ttf", ("malgun gothic", True): "malgunbd.ttf",
}


@dataclass
class Finding:
    severity: str
    slide: int          # 1-based slide number
    shape_id: int
    shape_name: str
    rule: str           # L1.. L6.x
    message: str

    def __str__(self):
        return (f"{self.severity:8s} slide {self.slide:>2} [{self.rule}] "
                f"{self.shape_name} (id {self.shape_id}): {self.message}")


# ------------------------------------------------------------------ fonts ---
@lru_cache(maxsize=None)
def _font(family: str, bold: bool):
    from PIL import ImageFont
    fam = (family or "arial").lower()
    fn = FONT_FILES.get((fam, bold)) or FONT_FILES[("arial", bold)]
    p = FONT_DIR / fn
    if not p.exists():  # non-Windows fallback
        for alt in ("DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf",):
            try:
                return ImageFont.truetype(alt, 1000)
            except OSError:
                pass
        return ImageFont.load_default()
    return ImageFont.truetype(str(p), 1000)


def _is_cjk(ch: str) -> bool:
    o = ord(ch)
    return (0x1100 <= o <= 0x11FF or 0x3130 <= o <= 0x318F or 0xAC00 <= o <= 0xD7AF
            or 0x4E00 <= o <= 0x9FFF or 0x3040 <= o <= 0x30FF or 0xFF00 <= o <= 0xFFEF)


def text_width_pt(text: str, family: str, bold: bool, pt: float) -> float:
    """Width in points. Hangul/CJK characters are measured with Malgun Gothic."""
    if not text:
        return 0.0
    total, buf, buf_cjk = 0.0, "", None

    def flush():
        nonlocal total, buf
        if buf:
            f = _font("malgun gothic" if buf_cjk else family, bold)
            total += f.getlength(buf) * pt / 1000.0
            buf = ""

    for ch in text:
        c = _is_cjk(ch)
        if buf_cjk is not None and c != buf_cjk:
            flush()
        buf_cjk = c
        buf += ch
    flush()
    return total


# ---------------------------------------------------------- color / WCAG ---
def _lin(c: float) -> float:
    c /= 255.0
    return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4


def luminance(hexrgb: str) -> float:
    r, g, b = (int(hexrgb[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * _lin(r) + 0.7152 * _lin(g) + 0.0722 * _lin(b)


def contrast_ratio(fg: str, bg: str) -> float:
    a, b = luminance(fg), luminance(bg)
    hi, lo = max(a, b), min(a, b)
    return (hi + 0.05) / (lo + 0.05)


_SCHEME_ALIAS = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


class ThemeColors:
    """Resolves schemeClr via the deck's first theme part (no lumMod support)."""

    def __init__(self, prs):
        self.map = {}
        try:
            master = prs.slide_masters[0]
            for rel in master.part.rels.values():
                if rel.reltype.endswith("/theme"):
                    root = etree.fromstring(rel.target_part.blob)
                    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    cs = root.find(".//a:clrScheme", ns)
                    for el in cs:
                        name = etree.QName(el).localname
                        ch = el[0]
                        val = ch.get("val") if etree.QName(ch).localname == "srgbClr" else ch.get("lastClr")
                        if val:
                            self.map[name] = val.upper()
        except Exception:
            pass

    def resolve(self, clr_holder) -> str | None:
        """clr_holder = element containing srgbClr/schemeClr/sysClr child (e.g. a:solidFill)."""
        if clr_holder is None:
            return None
        for ch in clr_holder:
            tag = etree.QName(ch).localname
            if tag == "srgbClr":
                if len(ch):  # lumMod/alpha/etc -> not exactly resolvable
                    return None
                return ch.get("val").upper()
            if tag == "sysClr":
                return (ch.get("lastClr") or "").upper() or None
            if tag == "schemeClr":
                if len(ch):
                    return None
                name = _SCHEME_ALIAS.get(ch.get("val"), ch.get("val"))
                return self.map.get(name)
        return None


# ------------------------------------------------------------- shape model --
@dataclass
class Item:
    shape: object
    z: int
    left: float
    top: float
    width: float
    height: float
    kind: str                 # text | picture | table | card | line | other
    fill: str | None = None   # resolved solid fill hex, if any
    unresolved_fill: bool = False
    ink: tuple | None = None  # (l, t, r, b) inches for text
    est_bottom: float | None = None
    est_height: float | None = None
    lines: list | None = None

    @property
    def right(self):
        return self.left + self.width

    @property
    def bottom(self):
        return self.top + self.height

    @property
    def name(self):
        return self.shape.name

    @property
    def sid(self):
        return self.shape.shape_id


def _flatten(shapes, ox=0, oy=0, sx=1.0, sy=1.0, out=None):
    out = out if out is not None else []
    for sh in shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if l is None or w is None:
            continue
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = sh._element.find(qn("p:grpSpPr")).find(qn("a:xfrm"))
            chOff, chExt = xfrm.find(qn("a:chOff")), xfrm.find(qn("a:chExt"))
            cx, cy = int(chExt.get("cx")) or 1, int(chExt.get("cy")) or 1
            gsx, gsy = w / cx, h / cy
            _flatten(sh.shapes, ox + l - int(chOff.get("x")) * gsx,
                     oy + t - int(chOff.get("y")) * gsy, gsx, gsy, out)
            continue
        out.append((sh, ox + l * sx, oy + t * sy, w * sx, h * sy))
    return out


def _spfill(sh, theme: ThemeColors):
    """Return (hex or None, unresolved) for a shape's own solid fill."""
    spPr = sh._element.find(qn("p:spPr"))
    if spPr is None:
        return None, False
    sf = spPr.find(qn("a:solidFill"))
    if sf is not None:
        c = theme.resolve(sf)
        return c, c is None
    if spPr.find(qn("a:gradFill")) is not None or spPr.find(qn("a:blipFill")) is not None:
        return None, True
    if spPr.find(qn("a:noFill")) is None and spPr.find(qn("a:pattFill")) is None:
        # no fill of its own: the shape style may fill it (fillRef idx > 0), as PowerPoint's default shapes do
        style = sh._element.find(qn("p:style"))
        ref = style.find(qn("a:fillRef")) if style is not None else None
        if ref is not None and ref.get("idx") not in (None, "0"):
            c = theme.resolve(ref)
            return c, c is None
    return None, False


def _has_text(sh) -> bool:
    return sh.has_text_frame and bool(sh.text_frame.text.strip())


# --------------------------------------------------------- text measurement --
def _pt(length_or_none):
    return None if length_or_none is None else length_or_none / 12700.0


def _run_style(r_el, p_el, tf_defaults):
    """Return (family, bold, size_pt) for an a:r element."""
    rPr = r_el.find(qn("a:rPr")) if r_el is not None else None
    sz = rPr.get("sz") if rPr is not None else None
    if sz is None:
        pPr = p_el.find(qn("a:pPr"))
        d = pPr.find(qn("a:defRPr")) if pPr is not None else None
        sz = d.get("sz") if d is not None else None
    if sz is None:
        end = p_el.find(qn("a:endParaRPr"))
        sz = end.get("sz") if (end is not None and r_el is None) else None
    size = int(sz) / 100.0 if sz else tf_defaults.get("size", DEFAULT_PT)
    b = rPr.get("b") if rPr is not None else None
    bold = (b in ("1", "true")) if b is not None else False
    family = "arial"
    if rPr is not None:
        lat = rPr.find(qn("a:latin"))
        if lat is not None and lat.get("typeface"):
            family = lat.get("typeface").lower()
    return family, bold, size


def layout_text_frame(tf, width_in: float, font_scale: float = 1.0):
    """Wrap a text frame into `width_in` inches of *inner* width.

    Returns (lines, height_pt) where lines = [(para_idx, line_w_pt, line_h_pt, align)].
    Height excludes first-paragraph space_before and last-paragraph space_after.
    """
    wrap = tf.word_wrap is not False
    lines = []
    total = 0.0
    paras = list(tf.paragraphs)
    for pi, para in enumerate(paras):
        p_el = para._p
        pPr = p_el.find(qn("a:pPr"))
        marL = int(pPr.get("marL", 0)) / 12700.0 if pPr is not None else 0.0
        indent = int(pPr.get("indent", 0)) / 12700.0 if pPr is not None else 0.0
        algn = (pPr.get("algn") if pPr is not None else None) or "l"
        # --- spacing
        ls = para.line_spacing
        sb = _pt(para.space_before) or 0.0
        sa = _pt(para.space_after) or 0.0
        # --- segments: list of (kind, text, family, bold, size)
        words, cur = [], []
        max_size = 0.0
        forced = []  # indices in words after which a hard break occurs
        for ch in p_el:
            tag = etree.QName(ch).localname
            if tag in ("r", "fld"):
                t = ch.findtext(qn("a:t")) or ""
                fam, bold, size = _run_style(ch, p_el, {})
                size *= font_scale
                max_size = max(max_size, size)
                for piece in re.split(r"(\s+)", t):
                    if piece == "":
                        continue
                    if piece.isspace():
                        if cur:
                            words.append(cur); cur = []
                        words.append([("sp", piece, fam, bold, size)])
                    else:
                        cur.append(("w", piece, fam, bold, size))
            elif tag == "br":
                if cur:
                    words.append(cur); cur = []
                words.append([("br", "", "arial", False, 0)])
        if cur:
            words.append(cur)
        if max_size == 0:  # empty paragraph
            end = p_el.find(qn("a:endParaRPr"))
            max_size = (int(end.get("sz")) / 100.0 if end is not None and end.get("sz")
                        else DEFAULT_PT) * font_scale
        # --- greedy wrap
        avail_first = width_in * 72.0 - marL - indent
        avail_rest = width_in * 72.0 - marL

        def seg_w(seg):
            return text_width_pt(seg[1], seg[2], seg[3], seg[4])

        para_lines = []  # (width_pt, max_size)
        cw, csz, first = 0.0, 0.0, True
        pend_sp = 0.0

        def push():
            nonlocal cw, csz, first, pend_sp
            para_lines.append((cw, csz or max_size))
            cw, csz, first, pend_sp = 0.0, 0.0, False, 0.0

        for wd in words:
            kind = wd[0][0]
            if kind == "br":
                push(); continue
            if kind == "sp":
                pend_sp += sum(seg_w(s) for s in wd)
                continue
            ww = sum(seg_w(s) for s in wd)
            wsz = max(s[4] for s in wd)
            avail = avail_first if first else avail_rest
            if wrap and cw > 0 and cw + pend_sp + ww > avail:
                push()
                avail = avail_rest
            if wrap and ww > avail:  # single over-long token: hard-break by char
                for s in wd:
                    for chx in s[1]:
                        cwid = text_width_pt(chx, s[2], s[3], s[4])
                        if cw + cwid > avail and cw > 0:
                            push(); avail = avail_rest
                        cw += cwid; csz = max(csz, s[4])
                pend_sp = 0.0
                continue
            cw += (pend_sp if cw > 0 else 0.0) + ww
            csz = max(csz, wsz)
            pend_sp = 0.0
        if cw > 0 or not para_lines or cw == 0 and csz == 0 and not words:
            push()
        elif cw > 0:
            push()
        # --- heights
        for (lw, lsz) in para_lines:
            if ls is None:
                mult_h = lsz * LINE_FACTOR
            elif isinstance(ls, float):
                mult_h = lsz * LINE_FACTOR * ls
            else:  # exact points
                mult_h = ls / 12700.0
            lines.append((pi, lw, mult_h, algn))
            total += mult_h
        if pi > 0:
            total += sb
        if pi < len(paras) - 1:
            total += sa
    return lines, total


def _body_pr(sh):
    el = sh._element.find(qn("p:txBody"))
    return el.find(qn("a:bodyPr")) if el is not None else None


def _insets_in(tf):
    def g(v, d):
        return d if v is None else v / EMU_IN
    return (g(tf.margin_left, 0.1), g(tf.margin_top, 0.05),
            g(tf.margin_right, 0.1), g(tf.margin_bottom, 0.05))


def measure_shape(it: Item):
    """Fill it.ink / est_bottom / est_height for a text-bearing shape."""
    sh = it.shape
    tf = sh.text_frame
    bp = _body_pr(sh)
    scale = 1.0
    autofit = None
    if bp is not None:
        if bp.find(qn("a:normAutofit")) is not None:
            na = bp.find(qn("a:normAutofit"))
            autofit = "norm"
            if na.get("fontScale"):
                scale = int(na.get("fontScale")) / 100000.0
        elif bp.find(qn("a:spAutoFit")) is not None:
            autofit = "sp"
    ml, mt, mr, mb = _insets_in(tf)
    inner_w = max(it.width - ml - mr, 0.05)
    lines, h_pt = layout_text_frame(tf, inner_w, scale)
    h = h_pt / 72.0
    anchor = (bp.get("anchor") if bp is not None else None) or "t"
    avail_h = it.height - mt - mb
    if anchor == "ctr":
        top = it.top + mt + (avail_h - h) / 2.0
    elif anchor == "b":
        top = it.bottom - mb - h
    else:
        top = it.top + mt
    bottom = top + h
    # horizontal ink from per-line widths and alignment
    lmin, rmax = 1e9, -1e9
    x0, x1 = it.left + ml, it.right - mr
    for (_, lw, _, algn) in lines:
        w_in = lw / 72.0
        if algn == "ctr":
            a = (x0 + x1) / 2 - w_in / 2
        elif algn == "r":
            a = x1 - w_in
        else:
            a = x0
        lmin, rmax = min(lmin, a), max(rmax, a + w_in)
    if lmin > rmax:  # no lines with width
        lmin, rmax = x0, x0
    it.ink = (lmin, top, max(rmax, lmin), bottom)
    it.est_bottom = bottom
    it.est_height = h + mt + mb
    it.lines = lines
    it.autofit = autofit
    return it


def measure_table(it: Item):
    """Estimated total table height (rows grow to fit their text)."""
    tbl = it.shape.table
    col_w = [c.width / EMU_IN for c in tbl.columns]
    total = 0.0
    for r in tbl.rows:
        rh = r.height / EMU_IN
        need = 0.0
        for ci, cell in enumerate(r.cells):
            if cell.is_spanned:
                continue
            span = cell.span_width if cell.is_merge_origin else 1
            w = sum(col_w[ci:ci + span])
            ml = (cell.margin_left if cell.margin_left is not None else 91440) / EMU_IN
            mr = (cell.margin_right if cell.margin_right is not None else 91440) / EMU_IN
            mt = (cell.margin_top if cell.margin_top is not None else 45720) / EMU_IN
            mb = (cell.margin_bottom if cell.margin_bottom is not None else 45720) / EMU_IN
            if not cell.text_frame.text.strip():
                continue
            _, h_pt = layout_text_frame(cell.text_frame, max(w - ml - mr, 0.05))
            need = max(need, h_pt / 72.0 + mt + mb)
        total += max(rh, need)
    it.est_height = total
    it.est_bottom = it.top + total


# ------------------------------------------------------------ slide analysis --
def _build_items(slide, theme: ThemeColors, sw: float, sh_: float):
    items = []
    for z, (shp, l, t, w, h) in enumerate(_flatten(slide.shapes)):
        l, t, w, h = l / EMU_IN, t / EMU_IN, w / EMU_IN, h / EMU_IN
        st = shp.shape_type
        el_tag = etree.QName(shp._element).localname
        fill, unres = (None, False)
        if el_tag == "sp":
            fill, unres = _spfill(shp, theme)
        if st == MSO_SHAPE_TYPE.PICTURE:
            kind = "picture"
        elif shp.has_table if hasattr(shp, "has_table") else False:
            kind = "table"
        elif el_tag == "cxnSp" or min(w, h) < THIN_IN:
            kind = "line"
        elif getattr(shp, "has_chart", False) and shp.has_chart:
            kind = "picture"  # chart: opaque graphic
        elif _has_text(shp):
            kind = "text"
        elif fill is not None or unres:
            kind = "card"
        else:
            kind = "other"
        it = Item(shp, z, l, t, w, h, kind, fill, unres)
        if kind == "text":
            rot = getattr(shp, "rotation", 0) or 0
            if rot == 0:
                measure_shape(it)
            else:
                it.ink = (l, t, l + w, t + h)
        elif kind == "table":
            measure_table(it)
            it.ink = (l, t, l + w, t + h)
        items.append(it)
    return items


def _inter(a, b):
    w = min(a[2], b[2]) - max(a[0], b[0])
    h = min(a[3], b[3]) - max(a[1], b[1])
    return (w * h) if (w > 0 and h > 0) else 0.0


def _area(r):
    return max(r[2] - r[0], 0) * max(r[3] - r[1], 0)


def _rect(it: Item):
    return (it.left, it.top, it.right, it.bottom)


def _contains(outer, inner, tol=0.03):
    return (outer[0] - tol <= inner[0] and outer[1] - tol <= inner[1]
            and outer[2] + tol >= inner[2] and outer[3] + tol >= inner[3])


def _is_footer(it: Item, floor: float) -> bool:
    return it.top >= floor - 0.01


def _is_fullbleed(it: Item, sw: float) -> bool:
    return it.left <= 0.05 and it.width >= sw - 0.1


def _ink_rect(it: Item):
    """Rect used for overlap tests: declared box for pictures/tables, the part of
    the text actually occupied for text boxes (clipped to the declared box so that
    overflow is reported by L1 only)."""
    if it.kind == "text" and it.ink:
        l, t, r, b = it.ink
        return (l, t, r, min(b, max(it.bottom, t + 0.01)))
    return _rect(it)


def _fid(it, slide_no, rule, sev, msg):
    return Finding(sev, slide_no, it.sid, it.name, rule, msg)


def _l1(items, slide_no, sw, floor, out):
    reported = set()
    for it in items:
        if it.kind not in ("text", "table") or it.est_bottom is None:
            continue
        if _is_footer(it, floor):
            continue
        if it.kind == "text" and getattr(it, "autofit", None) == "norm":
            continue
        ink = it.ink
        bottom = it.est_bottom
        # (a) floor
        if bottom > floor + FLOOR_TOL_IN:
            out.append(_fid(it, slide_no, "L1.floor", CRITICAL,
                            f"estimated text bottom {bottom:.2f}in crosses the {floor:.1f}in floor "
                            f"(box bottom {it.bottom:.2f}in)"))
            reported.add(it.sid)
            continue
        # (b) invading a shape below
        hit = None
        for ot in items:
            if ot is it or ot.kind in ("line", "other") or _is_footer(ot, floor):
                continue
            if ot.kind == "card" and _is_fullbleed(ot, sw):
                continue
            oi = ot.ink if (ot.kind == "text" and ot.ink) else _rect(ot)
            if ot.kind == "card" or ot.kind in ("picture", "table"):
                oi = _rect(ot)
            if ot.top < it.top - 0.01 or ot.top < ink[1]:
                continue                       # not below
            if _contains(_rect(ot), (ink[0], ink[1], ink[2], ink[3]), 0.05) and ot.kind == "card":
                continue                       # we sit in a card
            if _contains(_rect(it), _rect(ot)):
                # other shape lives inside our own box: only ink matters
                pass
            xo = min(ink[2], oi[2]) - max(ink[0], oi[0])
            if xo < 0.1 or xo < INVADE_MIN_XFRAC * max(ink[2] - ink[0], 0.01):
                continue                       # side-by-side columns, not stacked
            if bottom > oi[1] + INVADE_TOL_IN and ink[1] < oi[1]:
                hit = ot
                break
        if hit is not None:
            out.append(_fid(it, slide_no, "L1.invade", CRITICAL,
                            f"estimated text bottom {bottom:.2f}in runs into '{hit.name}' "
                            f"(top {hit.top:.2f}in)"))
            reported.add(it.sid)
            continue
        # (c) own box
        if it.est_height > it.height + OWN_BOX_TOL_IN and \
                not (it.kind == "text" and getattr(it, "autofit", None) == "sp"):
            what = "table" if it.kind == "table" else "text"
            out.append(_fid(it, slide_no, "L1.box", WARNING,
                            f"estimated {what} height {it.est_height:.2f}in exceeds its box "
                            f"{it.height:.2f}in (overshoot {it.est_height - it.height:.2f}in)"))


def _picture_ink_fraction(pic_item: Item, region):
    """Fraction of non-white opaque pixels of the picture inside `region`
    (l, t, r, b inches), or None when it cannot be read."""
    try:
        import io
        from PIL import Image
        im = Image.open(io.BytesIO(pic_item.shape.image.blob)).convert("RGBA")
        W, H = im.size
        x0 = int((region[0] - pic_item.left) / pic_item.width * W)
        x1 = int((region[2] - pic_item.left) / pic_item.width * W)
        y0 = int((region[1] - pic_item.top) / pic_item.height * H)
        y1 = int((region[3] - pic_item.top) / pic_item.height * H)
        x0, x1 = max(0, x0), min(W, max(x1, x0 + 1))
        y0, y1 = max(0, y0), min(H, max(y1, y0 + 1))
        crop = im.crop((x0, y0, x1, y1)).convert("RGBA")
        px = list(crop.getdata())
        dark = sum(1 for (r, g, b, a) in px if a > 40 and (r + g + b) / 3 < 225)
        return dark / max(len(px), 1)
    except Exception:
        return None


def _l2(items, slide_no, sw, floor, out):
    cands = []
    for it in items:
        if it.kind not in ("text", "picture", "table"):
            continue
        if it.name.startswith("logo:") or _is_footer(it, floor):
            continue
        if it.kind == "text" and _is_fullbleed(it, sw):
            continue                           # header bar carrying its own title
        cands.append(it)
    for i, a in enumerate(cands):
        for b in cands[i + 1:]:
            ra, rb = _ink_rect(a), _ink_rect(b)
            ov = _inter(ra, rb)
            if ov <= 0:
                continue
            # intentional: text fully inside a filled card / filled text box
            def card_of(x, y):
                return (x.fill is not None or x.unresolved_fill) and x.kind == "text" and \
                    y.kind == "text" and _contains(_rect(x), _rect(y), 0.03)
            if card_of(a, b) or card_of(b, a):
                continue
            small = min(_area(ra), _area(rb))
            if small <= 0:
                continue
            ratio = ov / small
            if ratio < 0.005:
                continue
            kinds = {a.kind, b.kind}
            hard = bool((kinds & {"picture", "table"}) and ratio > OVERLAP_CRIT_RATIO)
            if kinds == {"text"} and ratio > TEXT_OVERLAP_CRIT_RATIO:
                hard = True                    # two text blocks printed over each other
            note = ""
            if hard and kinds == {"picture", "text"}:
                pic = a if a.kind == "picture" else b
                reg = (max(ra[0], rb[0]), max(ra[1], rb[1]), min(ra[2], rb[2]), min(ra[3], rb[3]))
                frac = _picture_ink_fraction(pic, reg)
                if frac is not None and frac < BLANK_INK_FRAC:
                    hard, note = False, f" (picture region under the text is blank: {frac * 100:.1f}% ink)"
            sev = CRITICAL if hard else WARNING
            out.append(Finding(sev, slide_no, a.sid, a.name, "L2.overlap",
                               f"overlaps '{b.name}' (id {b.sid}) by {ratio * 100:.0f}% of the "
                               f"smaller shape ({a.kind} vs {b.kind}){note}"))


# ----------------------------------------------------------------- contrast --
def _slide_bg(slide, theme: ThemeColors, theme_name: str | None) -> str:
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        cSld = part._element.find(qn("p:cSld"))
        bg = cSld.find(qn("p:bg")) if cSld is not None else None
        if bg is None:
            continue
        bgPr = bg.find(qn("p:bgPr"))
        if bgPr is not None:
            sf = bgPr.find(qn("a:solidFill"))
            if sf is not None:
                c = theme.resolve(sf)
                if c:
                    return c
            continue
        ref = bg.find(qn("p:bgRef"))
        if ref is not None:
            c = theme.resolve(ref)
            if c:
                return c
    return THEME_DEFAULT_BG.get((theme_name or "default").lower(), "FFFFFF")


def _bg_for(items, idx, probe, slide_bg):
    """Walk z-order downward from item idx; return (hex|None, reason)."""
    for j in range(idx, -1, -1):
        c = items[j]
        if c.kind in ("line", "other"):
            continue
        if j != idx and not (c.left <= probe[0] <= c.right and c.top <= probe[1] <= c.bottom):
            continue
        if j == idx and c.kind not in ("text", "card"):
            continue
        if c.kind == "picture" and j != idx:
            return None, "picture"
        if c.fill is not None:
            return c.fill, None
        if c.unresolved_fill:
            return None, "unresolved fill"
    return slide_bg, None


def _run_color(r_el, theme: ThemeColors):
    rPr = r_el.find(qn("a:rPr"))
    if rPr is None:
        return None
    sf = rPr.find(qn("a:solidFill"))
    return theme.resolve(sf) if sf is not None else None


def _check_runs(tf, bg, owner, slide_no, loc, theme, out, seen, cap_warning=False):
    for para in tf.paragraphs:
        for r in para._p.findall(qn("a:r")):
            txt = r.findtext(qn("a:t")) or ""
            if not txt.strip():
                continue
            fg = _run_color(r, theme)
            if fg is None:
                continue
            fam, bold, size = _run_style(r, para._p, {})
            ratio = contrast_ratio(fg, bg)
            large = size >= LARGE_PT or (bold and size >= LARGE_BOLD_PT)
            if ratio < CONTRAST_CRIT:
                sev = CRITICAL
            elif ratio < CONTRAST_NORMAL:
                sev = WARNING if large else CRITICAL
            else:
                continue
            if cap_warning and sev == CRITICAL:
                sev = WARNING     # incidental footer/page-number text is never a gate failure
            key = (owner.sid, loc, fg, bg, sev)
            if key in seen:
                continue
            seen.add(key)
            out.append(_fid(owner, slide_no, "L3.contrast", sev,
                            f"{loc}text #{fg} on #{bg} = {ratio:.2f}:1 "
                            f"({size:g}pt{' bold' if bold else ''}, {'large' if large else 'normal'}); "
                            f"'{txt.strip()[:30]}'"))


def _l3(slide, items, slide_no, theme, theme_name, out, floor=FLOOR_IN):
    slide_bg = _slide_bg(slide, theme, theme_name)
    seen = set()
    for idx, it in enumerate(items):
        if it.kind == "text":
            probe = ((it.ink[0] + it.ink[2]) / 2, (it.ink[1] + it.ink[3]) / 2) if it.ink else \
                (it.left + it.width / 2, it.top + it.height / 2)
            bg, why = _bg_for(items, idx, probe, slide_bg)
            if bg is None:
                if why == "picture":
                    out.append(_fid(it, slide_no, "L3.unverifiable", WARNING,
                                    "text sits over a picture: contrast unverifiable"))
                continue
            _check_runs(it.shape.text_frame, bg, it, slide_no, "", theme, out, seen,
                        cap_warning=_is_footer(it, floor))
        elif it.kind == "table":
            tbl = it.shape.table
            tblPr = it.shape._element.find(".//" + qn("a:tblPr"))
            styled = tblPr is not None and tblPr.find(qn("a:tableStyleId")) is not None
            bg_under, why = _bg_for(items, idx, (it.left + 0.1, it.top + 0.1), slide_bg) \
                if False else (slide_bg, None)
            for ri, row in enumerate(tbl.rows):
                for ci, cell in enumerate(row.cells):
                    if cell.is_spanned or not cell.text_frame.text.strip():
                        continue
                    tcPr = cell._tc.find(qn("a:tcPr"))
                    sf = tcPr.find(qn("a:solidFill")) if tcPr is not None else None
                    if sf is not None:
                        bg = theme.resolve(sf)
                        if bg is None:
                            continue
                    elif styled:
                        continue                   # table-style fill unknown
                    else:
                        bg = slide_bg
                    _check_runs(cell.text_frame, bg, it, slide_no,
                                f"cell r{ri + 1}c{ci + 1} ", theme, out, seen)


# ---------------------------------------------------------------- structure --
def _all_texts(items):
    for it in items:
        if it.kind in ("text", "card") and it.shape.has_text_frame:
            yield it, it.shape.text_frame.text
        elif it.kind == "table":
            for row in it.shape.table.rows:
                for cell in row.cells:
                    yield it, cell.text_frame.text


def _is_appendix(items) -> bool:
    """A hidden slide titled Appendix/Backup is intentional (journal-club preset); any other hidden slide is flagged."""
    for it in items:
        if it.kind == "text" and it.top < TITLE_ZONE_IN and APPENDIX_RE.match(it.shape.text_frame.text.strip()):
            return True
    return False


def _l4(slide, items, slide_no, floor, out):
    if slide._element.get("show") == "0" and not _is_appendix(items):
        out.append(Finding(WARNING, slide_no, 0, "(slide)", "L4.hidden",
                           "slide is hidden (show=0): it will be skipped in the slideshow"))
    content = [it for it in items
               if it.kind in ("picture", "table") or (it.kind == "text" and not _is_footer(it, floor))]
    if not content:
        out.append(Finding(CRITICAL, slide_no, 0, "(slide)", "L4.blank",
                           "slide has no content shapes besides footer/decoration"))
    for it, txt in _all_texts(items):
        m = PLACEHOLDER_RE.search(txt)
        if m:
            out.append(_fid(it, slide_no, "L4.placeholder", CRITICAL,
                            f"placeholder text '{m.group(0)}' in '{txt.strip()[:40]}'"))


# ------------------------------------------------------------ L5 aspect ---
def _l5(items, slide_no, out):
    """A picture shown at another aspect ratio than its pixels (after cropping) is stretched."""
    for it in items:
        if it.kind != "picture" or etree.QName(it.shape._element).localname != "pic":
            continue
        if (getattr(it.shape, "rotation", 0) or 0) % 180:
            continue
        try:
            iw, ih = it.shape.image.size
        except Exception:
            continue                           # vector or unreadable image: nothing to compare
        if not (iw and ih and it.width > 0 and it.height > 0):
            continue
        vis_w = iw * max(1.0 - it.shape.crop_left - it.shape.crop_right, 1e-6)
        vis_h = ih * max(1.0 - it.shape.crop_top - it.shape.crop_bottom, 1e-6)
        dev = abs((it.width / it.height) / (vis_w / vis_h) - 1.0)
        if dev > ASPECT_WARN:
            sev = CRITICAL if dev > ASPECT_CRIT else WARNING
            out.append(_fid(it, slide_no, "L5.aspect", sev,
                            f"picture is shown at aspect {it.width / it.height:.2f} but its pixels are "
                            f"{vis_w / vis_h:.2f} ({dev * 100:.0f}% distortion)"))


# --------------------------------------------------------- L6 slide order ---
def _slide_texts(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            yield sh, sh.text_frame.text.strip()


def _l6(prs, out):
    """The references slide closes the talk: only Q&A / thanks / backup slides may follow it."""
    slides = list(prs.slides)
    for n, slide in enumerate(slides, 1):
        title = next((txt for sh, txt in _slide_texts(slide)
                      if sh.top / EMU_IN < TITLE_ZONE_IN and REFS_RE.match(txt)), None)
        if title is None:
            continue
        for m, later in enumerate(slides[n:], n + 1):
            texts = [txt for _, txt in _slide_texts(later)]
            if any(REFS_RE.match(x) for x in texts) or any(AFTER_REFS_RE.match(x) for x in texts):
                continue                       # continued references or an allowed closing slide
            out.append(Finding(CRITICAL, n, 0, "(slide)", "L6.refs_order",
                               f"references slide {n} is followed by content slide {m}: "
                               f"references must come last (before Q&A / thanks / backup)"))
            break


# --------------------------------------------------------------------- API ---
REF_H_IN = 7.5               # slide height the 7.0in floor was defined for (13.333 x 7.5 in)


def scaled_floor(sh_in: float) -> float:
    """Y floor for a slide of height `sh_in`: 7.0 in on the reference 7.5 in height, proportional
    elsewhere (a 20 x 11.25 in template gets 10.5 in). Reference-size slides return FLOOR_IN exactly."""
    if abs(sh_in - REF_H_IN) < 0.01:
        return FLOOR_IN
    return FLOOR_IN * sh_in / REF_H_IN


def check_layout(prs, theme: str | None = None, floor_in: float | None = None,
                 rules=("L1", "L2", "L3", "L4", "L5", "L6"), external_template: bool = False) -> list[Finding]:
    """Run the geometric QC on an open Presentation; returns Findings.

    floor_in=None derives the Y floor from the slide height (scaled_floor). external_template
    is accepted for symmetry with qc_deck: every layout rule (overflow, overlap, contrast,
    structure) applies unchanged to a template deck, so it switches nothing off."""
    sw = prs.slide_width / EMU_IN
    sh_ = prs.slide_height / EMU_IN
    if floor_in is None:
        floor_in = scaled_floor(sh_)
    tcol = ThemeColors(prs)
    out: list[Finding] = []
    for n, slide in enumerate(prs.slides, 1):
        items = _build_items(slide, tcol, sw, sh_)
        if "L1" in rules:
            _l1(items, n, sw, floor_in, out)
        if "L2" in rules:
            _l2(items, n, sw, floor_in, out)
        if "L3" in rules:
            _l3(slide, items, n, tcol, theme, out, floor_in)
        if "L4" in rules:
            _l4(slide, items, n, floor_in, out)
        if "L5" in rules:
            _l5(items, n, out)
    if "L6" in rules:
        _l6(prs, out)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="Layout QC (overflow / overlap / contrast / structure)")
    ap.add_argument("pptx")
    ap.add_argument("--theme", default=None, help="theme name (default background lookup)")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--floor", type=float, default=None,
                    help="Y floor in inches (default: 7.0 scaled by slide height / 7.5)")
    ap.add_argument("--external-template", action="store_true",
                    help="deck built on a third-party template (all layout rules stay on)")
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    prs = Presentation(a.pptx)
    fs = check_layout(prs, theme=a.theme, floor_in=a.floor, external_template=a.external_template)
    nc = sum(f.severity == CRITICAL for f in fs)
    if a.json:
        print(json.dumps([asdict(f) for f in fs], ensure_ascii=False, indent=1))
    else:
        for f in fs:
            print(f)
        print(f"\nlayout QC: CRITICAL={nc} WARNING={len(fs) - nc}")
    return 1 if nc else 0


if __name__ == "__main__":
    sys.exit(main())

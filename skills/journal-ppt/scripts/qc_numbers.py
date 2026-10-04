#!/usr/bin/env python3
"""
qc_numbers.py -- number-verification gate for journal-ppt decks.

Compares every numeric claim on the slides (NOT the speaker notes) against the
numbers recorded in content_analysis.json, the "number-verification checkpoint"
of references/pipeline.md (Phase 2 Step 5 / Step 7) and SKILL.md
"Content-integrity rules".

Usage:
    python qc_numbers.py deck.pptx --source content_analysis.json [--ledger out.json] [--strict]

Status per slide number
    MATCH             in source with same value+unit (unit conversion and honest
                      rounding allowed)                                   -> OK
    DRIFT             not in source, but a near-miss of a source number (digit
                      transposition, one-digit change, decimal shift, unit
                      mismatch, false precision, or within 10%)           -> CRITICAL
    SOURCE-CONFLICT  matches a number the source itself flags as conflicting
                      (validated_numbers match:false or a disagreement in
                      validation_warnings) and no visible caveat on the slide -> CRITICAL
    DERIVED-UNLABELED reproducible from two source numbers (ratio, difference,
                      percent change, sum) but no estimate marker (~, approx,
                      est., calculated ...) on the slide paragraph or in the
                      slide's notes                                       -> WARNING
    DERIVED-LABELED   same, with a marker present                         -> OK
    UNSOURCED         not in source and no near-miss                      -> WARNING

Exit codes (shared verification-gates contract)
    0 = clean (no CRITICAL; warnings are printed, --strict makes them fail)
    1 = at least one CRITICAL finding
    2 = BLIND: no source numbers, no slide numbers, unreadable/invalid inputs.
        Never read a 2 as a pass.

The ledger JSON lists every slide number with slide, value, unit, context,
status, source ref and nearest source candidate, plus the numbers ignored as
trivial (with the rule that ignored them), so a verifier can audit the scope.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

EXIT_CLEAN, EXIT_CRITICAL, EXIT_BLIND = 0, 1, 2

# ---------------------------------------------------------------------------
# Units
# ---------------------------------------------------------------------------

_SUPER = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺", "0123456789-+")
_PREFIX = {"p": 1e-12, "n": 1e-9, "u": 1e-6, "m": 1e-3, "c": 1e-2, "k": 1e3}
# base atom -> (dimension dict, scale to base)
_BASE = {
    "M": ({"amount": 1, "vol": -1}, 1.0), "mol": ({"amount": 1}, 1.0),
    "g": ({"mass": 1}, 1.0), "L": ({"vol": 1}, 1.0), "l": ({"vol": 1}, 1.0),
    "m": ({"len": 1}, 1.0), "s": ({"time": 1}, 1.0), "Da": ({"da": 1}, 1.0),
    "Hz": ({"freq": 1}, 1.0), "V": ({"volt": 1}, 1.0), "A": ({"amp": 1}, 1.0),
    "W": ({"watt": 1}, 1.0), "Pa": ({"pressure": 1}, 1.0), "U": ({"U": 1}, 1.0),
    "bp": ({"bp": 1}, 1.0), "J": ({"energy": 1}, 1.0), "cal": ({"energy": 1}, 4.184),
}
_NO_PREFIX = {
    "min": ({"time": 1}, 60.0), "h": ({"time": 1}, 3600.0), "hr": ({"time": 1}, 3600.0),
    "d": ({"time": 1}, 86400.0), "K": ({"kelvin": 1}, 1.0), "rpm": ({"rpm": 1}, 1.0),
    "kb": ({"bp": 1}, 1e3), "Mb": ({"bp": 1}, 1e6), "bar": ({"pressure": 1}, 1e5),
    "psi": ({"pressure": 1}, 6894.76), "eq": ({"eq": 1}, 1.0), "equiv": ({"eq": 1}, 1.0),
    "OD": ({"od": 1}, 1.0), "CFU": ({"cfu": 1}, 1.0), "ppm": ({"ppm": 1}, 1.0),
    "ppb": ({"ppb": 1}, 1.0), "%": ({"pct": 1}, 1.0), "‰": ({"permil": 1}, 1.0),
    "mol%": ({"molpct": 1}, 1.0), "wt%": ({"wtpct": 1}, 1.0), "vol%": ({"volpct": 1}, 1.0),
    "°C": ({"celsius": 1}, 1.0), "°F": ({"fahrenheit": 1}, 1.0), "°": ({"deg": 1}, 1.0),
    "fold": ({"fold": 1}, 1.0), "p": ({"pvalue": 1}, 1.0), "n": ({"count_n": 1}, 1.0),
    "pH": ({"ph": 1}, 1.0),
}


def _atom_info(atom: str):
    """(dim dict, scale) for a unit atom such as 'mM', 'min', 'kDa', else None."""
    if atom in _NO_PREFIX:
        return _NO_PREFIX[atom]
    if atom in _BASE:
        return _BASE[atom]
    if len(atom) > 1 and atom[0] in _PREFIX and atom[1:] in _BASE:
        dims, sc = _BASE[atom[1:]]
        return dims, sc * _PREFIX[atom[0]]
    return None


_ATOM_RE = re.compile(r"^([A-Za-z°%‰]+)(?:\^?(-?\d+))?$")


@dataclass(frozen=True)
class Unit:
    canon: str          # normalized text, '' = no unit
    dim: frozenset      # frozenset of (dimension, exponent)
    scale: float        # multiply a value by this to reach the dimension's base

    @property
    def empty(self) -> bool:
        return self.canon == ""


NO_UNIT = Unit("", frozenset(), 1.0)


def parse_unit(text: str | None) -> Unit | None:
    """Parse '(µ)M', 'mg/mL', 's⁻¹·mM⁻¹', '%' ... Returns None when text is not
    a recognized unit (so 'salts' or 'ligands' are never mistaken for units)."""
    if text is None:
        return None
    t = text.strip().replace("μ", "u").replace("µ", "u").replace("℃", "°C")
    if not t:
        return NO_UNIT
    # superscript exponents -> ^n
    t = re.sub(r"([⁻⁺]?[⁰¹²³⁴⁵⁶⁷⁸⁹]+)", lambda m: "^" + m.group(1).translate(_SUPER), t)
    t = re.sub(r"(?<=[A-Za-z0-9])[ \u2009\u00a0\u202f]+(?=[A-Za-z])", "·", t)
    t = re.sub(r"\s+", "", t)
    dims: dict[str, int] = {}
    scale = 1.0
    sign = 1
    pieces = re.split(r"([/·⋅*])", t)
    for piece in pieces:
        if piece == "/":
            sign = -1
            continue
        if piece in ("·", "⋅", "*"):
            continue
        m = _ATOM_RE.match(piece)
        if not m:
            return None
        info = _atom_info(m.group(1))
        if info is None:
            return None
        exp = int(m.group(2)) if m.group(2) else 1
        for d, e in info[0].items():
            dims[d] = dims.get(d, 0) + sign * exp * e
        scale *= info[1] ** (sign * exp)
    dimset = frozenset((d, e) for d, e in dims.items() if e)
    return Unit(t, dimset, scale)


# ---------------------------------------------------------------------------
# Numeric claim model
# ---------------------------------------------------------------------------

@dataclass
class Num:
    value: float
    unit: Unit
    raw: str
    kind: str = "value"         # value | pct | p | n | fold | range_lo | range_hi | pm_err | sci | ph
    operator: str = ""          # for p-values: '<', '=', ...
    decimals: int | None = 0    # None for scientific notation
    digits: str = ""            # digit string of the mantissa as typed
    start: int = 0
    end: int = 0
    # location (slide side)
    slide: int = 0
    shape: str = ""
    context: str = ""
    paragraph: str = ""
    # location (source side)
    ref: str = ""
    conflict: bool = False
    conflict_note: str = ""

    @property
    def base(self) -> float:
        return self.value * self.unit.scale

    def to_json(self) -> dict:
        return {"value": self.value, "unit": self.unit.canon, "raw": self.raw, "kind": self.kind,
                "operator": self.operator or None}


# ---------------------------------------------------------------------------
# Text extraction (one string -> numeric claims + ignored trivial numbers)
# ---------------------------------------------------------------------------

_SP = r"[    ]"
_NUM = r"(?:\d{1,3}(?:,\d{3})+|\d{1,3}(?:[   ]\d{3})+|\d+)(?:\.\d+)?|\.\d+"
_SUPDIG = "⁰¹²³⁴⁵⁶⁷⁸⁹"
_TOKEN_RE = re.compile(
    rf"(?P<sci>(?P<m1>{_NUM}){_SP}?[×x·⋅*]{_SP}?10{_SP}?(?:\^{_SP}?)?(?P<e1>[-−–⁻+⁺]?[0-9{_SUPDIG}]+)"
    rf"|(?P<m2>\d+(?:\.\d+)?)[eE](?P<e2>[-−+]?\d+))(?![0-9])"
    rf"|(?P<num>{_NUM})"
)
_UNIT_RE = re.compile(
    rf"(?P<sp>{_SP}?)(?P<u>[A-Za-zµμ°%‰]+(?:\^?-?[0-9]+|[⁻⁺]?[{_SUPDIG}]+)?"
    rf"(?:(?:[/·⋅*]|{_SP}(?=[A-Za-zµμ]+(?:\^-|⁻)))[A-Za-zµμ]+(?:\^?-?[0-9]+|[⁻⁺]?[{_SUPDIG}]+)?)*)"
)
_PCT_RE = re.compile(rf"{_SP}?(?P<u>[%‰])")
_DEG_RE = re.compile(rf"{_SP}?(?P<u>°{_SP}?[CF]|℃)(?![A-Za-z])")
_FOLD_RE = re.compile(r"(?:-|‑|\s)?fold\b|[×x](?![A-Za-z0-9])")
_FOLD_SPACED_TIMES = re.compile(rf"{_SP}[×](?!{_SP}?[A-Za-z0-9])")
_PM_RE = re.compile(rf"{_SP}?(?:±|\+/-|\+/−|\+-)\s?")
_RANGE_RE = re.compile(rf"{_SP}?(?:[–—-]|\bto\b){_SP}?(?=[0-9.])")
_PRE_P = re.compile(r"(?<![A-Za-z])[pP]\s?(<=|>=|[<>=≤≥≈])\s?$")
_PRE_N = re.compile(r"(?<![A-Za-z])[nN]\s?=\s?$")
_PRE_PH = re.compile(r"(?<![A-Za-z])pH\s?[=:~≈]?\s?$")

# Trivial-number masks: (rule name, regex). Applied before claim extraction.
_FIGREF = r"(?:Figures?|Figs?\.?|Tables?|Tab\.?|Schemes?|Equations?|Eqs?\.?|Supplementary|Suppl\.?|Slides?|Steps?|Panels?|Sections?|Chapters?|Appendix|Movie|Video)"
_MASKS: list[tuple[str, re.Pattern]] = [
    ("doi", re.compile(r"\b10\.\d{4,9}/[^\s,;)\]]+")),
    ("url", re.compile(r"(?:https?://|www\.)[^\s,;)\]]+")),
    ("iso_date", re.compile(r"\b(?:19|20)\d{2}[-/.]\d{1,2}[-/.]\d{1,2}\b")),
    ("date", re.compile(r"\b\d{1,2}[-/.]\d{1,2}[-/.](?:19|20)?\d{2}\b")),
    ("figure_table_ref", re.compile(
        rf"\b{_FIGREF}{_SP}?\(?S?\d+[A-Za-z]?\)?(?:{_SP}?(?:,|and|&|[–-]){_SP}?S?\d+[A-Za-z]?)*(?:\([A-Za-z]\))?")),
    ("journal_volume_pages", re.compile(
        r"\b\d{1,4},\s\d{3,6}(?:\s?[–-]\s?\d{1,6})?|\b\d{1,4}:\d{3,6}(?:\s?[–-]\s?\d{1,6})?")),
    ("reference_index", re.compile(r"\[\s*\d+(?:\s*[,–-]\s*\d+)*\s*\]")),
    ("journal_citation", re.compile(r"\b\d{1,4}\s?\(\d{1,3}\)\s?:\s?[A-Za-z]?\d+(?:\s?[–-]\s?\d+)?")),
    ("page_range", re.compile(r"\bpp?\.\s?\d+(?:\s?[–-]\s?\d+)?")),
    ("volume_issue", re.compile(r"\b(?:Vol|No|Issue|Nr)\.?\s?\d+", re.I)),
    ("time_or_ratio", re.compile(r"(?<![\d.])\d{1,3}(?::\d{1,3})+(?![\d])")),
    ("version", re.compile(r"\bv\d+(?:\.\d+)*\b")),
    ("citation_year_paren", re.compile(r"\([^()]*?\b(?:19|20)\d{2}[a-z]?[^()]*?\)")),
    ("et_al_year", re.compile(r"\bet al\.?,?\s?(?:19|20)\d{2}[a-z]?")),
]
_REFLINE_RE = re.compile(r"\bet al\b|\bdoi\b|\b10\.\d{4,9}/", re.I)
_BARE_YEAR = re.compile(r"^(?:19|20)\d{2}$")


def _to_float(txt: str) -> float:
    return float(re.sub(r"[,   ]", "", txt))


def _decimals(txt: str) -> int:
    t = re.sub(r"[,   ]", "", txt)
    return len(t.split(".")[1]) if "." in t else 0


def _digits(txt: str) -> str:
    return re.sub(r"[^0-9]", "", txt)


def _sci_value(mant: str, exp: str) -> float:
    e = exp.translate(_SUPER).replace("−", "-").replace("–", "-").replace(" ", "")
    return _to_float(mant) * 10 ** int(e)


def _clean_context(text: str, s: int, e: int, width: int = 40) -> str:
    return re.sub(r"\s+", " ", text[max(0, s - width): e + width]).strip()


def _find_unit(text: str, pos: int):
    """Unit immediately after position pos. Returns (Unit, end, is_fold) or None."""
    m = _PCT_RE.match(text, pos)
    if m:
        # '85 %' ok; 'mol%' / 'wt%' handled by the generic matcher below first
        return parse_unit(m.group("u")), m.end(), False
    m = _DEG_RE.match(text, pos)
    if m:
        return parse_unit(m.group("u").replace(" ", "")), m.end(), False
    m = _FOLD_RE.match(text, pos) or _FOLD_SPACED_TIMES.match(text, pos)
    if m:
        return parse_unit("fold"), m.end(), True
    m = _UNIT_RE.match(text, pos)
    if m:
        u_text = m.group("u")
        end = m.end()
        # an unattached noun that merely starts like a unit is not a unit
        if end < len(text) and (text[end].isalnum()):
            return None
        unit = parse_unit(u_text)
        if unit is not None and not unit.empty:
            return unit, end, False
        # 'mg/mL/x' style: retry with the part before the last separator
        return None
    return None


def extract_numbers(text: str, *, whole_text_bare_ok: bool = True, small_bare_ok: bool = False,
                    header_unit: Unit | None = None):
    """Extract numeric claims from one string.

    Returns (claims, ignored) where ignored is a list of dicts {raw, rule, context}.
    `small_bare_ok` keeps single-digit unit-less integers (used for none by default).
    """
    ignored: list[dict] = []
    masked: list[tuple[int, int, str]] = []
    for rule, rx in _MASKS:
        for m in rx.finditer(text):
            masked.append((m.start(), m.end(), rule))
    for s, e, rule in masked:
        if any(ch.isdigit() for ch in text[s:e]):
            ignored.append({"raw": text[s:e], "rule": rule, "context": _clean_context(text, s, e)})
    refline = bool(_REFLINE_RE.search(text)) or any(
        r in ("citation_year_paren", "et_al_year", "journal_volume_pages", "journal_citation")
        for _, _, r in masked)

    def is_masked(s: int, e: int) -> bool:
        return any(ms <= s and e <= me for ms, me, _ in masked)

    claims: list[Num] = []
    consumed_until = 0
    for m in _TOKEN_RE.finditer(text):
        s, e = m.start(), m.end()
        if s < consumed_until or is_masked(s, e):
            continue
        before = text[:s]
        is_sci = m.group("sci") is not None
        if is_sci:
            mant, ex = (m.group("m1"), m.group("e1")) if m.group("m1") else (m.group("m2"), m.group("e2"))
            value, decimals, digits, raw_kind = _sci_value(mant, ex), None, _digits(mant), "sci"
        else:
            tok = m.group("num")
            value, decimals, digits, raw_kind = _to_float(tok), _decimals(tok), _digits(tok), "value"
        kind, operator, forced_unit = raw_kind, "", None
        if not is_sci and re.fullmatch(r"0\d+", m.group("num")):
            ignored.append({"raw": m.group(0), "rule": "leading_zeros", "context": _clean_context(text, s, e)})
            continue

        mp = _PRE_P.search(before)
        if mp:
            kind, operator, forced_unit = "p", mp.group(1).replace("<=", "≤").replace(">=", "≥"), "p"
        elif _PRE_N.search(before):
            kind, forced_unit = "n", "n"
        elif _PRE_PH.search(before):
            kind, forced_unit = "ph", "pH"
        else:
            # identifier guard: 'T4', 'Cas9', 'ZIF-8', 'BL21', '_3'
            if s > 0 and (text[s - 1].isalpha() or text[s - 1] == "_" or
                          (text[s - 1] in "-‑" and s > 1 and text[s - 2].isalpha())):
                ignored.append({"raw": m.group(0), "rule": "identifier", "context": _clean_context(text, s, e)})
                continue
        # sign
        sign = 1
        if s > 0 and text[s - 1] in "-−–" and (s == 1 or text[s - 2] in " \t(=[,;:~≈±" + "  "):
            sign = -1
        value *= sign
        start = s - (1 if sign == -1 else 0)

        unit: Unit = NO_UNIT
        end = e
        if forced_unit:
            unit = parse_unit(forced_unit)
        else:
            found = _find_unit(text, e)
            if found:
                unit, end, is_fold = found
                if is_fold:
                    kind = "fold"
                elif unit.canon == "%":
                    kind = "pct"
        own_end = end
        # '5.2 ± 0.3 mM' : error term shares the unit
        extras: list[tuple[Num, str]] = []
        pm = _PM_RE.match(text, end) if not forced_unit else None
        if pm:
            m2 = _TOKEN_RE.match(text, pm.end())
            if m2 and not is_masked(m2.start(), m2.end()) and m2.group("num"):
                tok2 = m2.group("num")
                end2 = m2.end()
                u2 = _find_unit(text, end2)
                if unit.empty and u2:
                    unit = u2[0]
                    kind = "pct" if unit.canon == "%" else kind
                if u2:
                    end2 = u2[1]
                err = Num(_to_float(tok2), unit, text[pm.end():end2], "pm_err", "", _decimals(tok2),
                          _digits(tok2), pm.end(), end2)
                extras.append((err, "pm"))
                end = end2
        # '10-20 %' / '3 to 5 mM' ranges share the unit
        elif not forced_unit and not is_sci:
            rg = _RANGE_RE.match(text, end)
            if rg:
                m2 = _TOKEN_RE.match(text, rg.end())
                if m2 and not is_masked(m2.start(), m2.end()) and m2.group("num"):
                    tok2 = m2.group("num")
                    u2 = _find_unit(text, m2.end())
                    end2 = u2[1] if u2 else m2.end()
                    if unit.empty and u2:
                        unit = u2[0]
                        kind = "pct" if unit.canon == "%" else kind
                    hi = Num(_to_float(tok2), unit, text[rg.end():end2], "range_hi", "", _decimals(tok2),
                             _digits(tok2), rg.end(), end2)
                    extras.append((hi, "range"))
                    kind = "range_lo"
                    end = end2
        # unit-less number glued to letters or hyphen-modifier: '3D', '96-well', '2nd'
        if unit.empty and not forced_unit and not extras:
            if end < len(text) and text[end].isalpha():
                ignored.append({"raw": m.group(0), "rule": "attached_letters",
                                "context": _clean_context(text, s, e)})
                continue
            if end + 1 < len(text) and text[end] in "-‑" and text[end + 1].isalpha():
                ignored.append({"raw": m.group(0), "rule": "hyphenated_modifier",
                                "context": _clean_context(text, s, e)})
                continue
        if unit.empty and header_unit is not None and kind in ("value", "range_lo"):
            unit = header_unit
        primary = Num(value, unit, text[start:own_end].strip(), kind, operator, decimals, digits, start, own_end)
        group = [primary] + [x for x, _ in extras]
        for n in group:
            n.context = _clean_context(text, start, end)
            n.paragraph = text
            if n.unit.empty and n.kind in ("value", "range_lo", "range_hi") and n.decimals is not None:
                reason = _trivial_bare(n, text, refline)
                if reason and not (small_bare_ok and reason == "small_bare_integer"):
                    ignored.append({"raw": n.raw, "rule": reason, "context": n.context})
                    continue
            claims.append(n)
        consumed_until = end
    return claims, ignored


def _trivial_bare(n: Num, text: str, refline: bool) -> str | None:
    """Rule name when a unit-less number is trivial, else None."""
    if n.decimals == 0 and abs(n.value) == int(abs(n.value)):
        iv = int(abs(n.value))
        if _BARE_YEAR.match(str(iv)) and n.raw.strip().isdigit():
            return "year"
        if iv <= 9 and n.raw.strip().lstrip("-−").isdigit():
            return "small_bare_integer"
    if refline:
        return "citation_line"
    return None


# ---------------------------------------------------------------------------
# Deck reading
# ---------------------------------------------------------------------------

_BARE_INT_SHAPE = re.compile(r"^\s*(?:‹#›)?\s*(\d{1,3})\s*(?:(?:/|of)\s*\d{1,3})?\s*$")


def _iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # group
            yield from _iter_shapes(sh.shapes)


def _is_footer_placeholder(sh) -> bool:
    try:
        if sh.is_placeholder:
            return sh.placeholder_format.type in (13, 15, 16)  # SLIDE_NUMBER, FOOTER, DATE
    except Exception:
        pass
    return False


def _header_units(tbl) -> dict[int, Unit]:
    out: dict[int, Unit] = {}
    try:
        for ci, cell in enumerate(tbl.rows[0].cells):
            m = re.search(r"[\(\[]\s*([^()\[\]]+?)\s*[\)\]]\s*$", cell.text.strip())
            if m:
                u = parse_unit(m.group(1))
                if u and not u.empty:
                    out[ci] = u
    except Exception:
        pass
    return out


def read_deck(path: str | Path):
    """Return (slide_units, notes, n_slides).

    slide_units: list of (slide_index, shape_name, text, header_unit_or_None, whole_shape)
    notes: {slide_index: notes text}
    """
    from pptx import Presentation
    prs = Presentation(str(path))
    units = []
    notes: dict[int, str] = {}
    n_slides = len(prs.slides)
    for si, slide in enumerate(prs.slides, 1):
        for sh in _iter_shapes(slide.shapes):
            name = sh.name
            if _is_footer_placeholder(sh):
                continue
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                whole = sh.text_frame.text
                for pi, para in enumerate(sh.text_frame.paragraphs):
                    units.append((si, f"{name}/p{pi}", "".join(r.text for r in para.runs) or para.text,
                                  None, whole))
            if getattr(sh, "has_table", False) and sh.has_table:
                hu = _header_units(sh.table)
                for ri, row in enumerate(sh.table.rows):
                    for ci, cell in enumerate(row.cells):
                        units.append((si, f"{name}/r{ri}c{ci}", cell.text, hu.get(ci) if ri > 0 else None,
                                      cell.text))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes[si] = slide.notes_slide.notes_text_frame.text
    return units, notes, n_slides


def extract_slide_numbers(path: str | Path):
    """Extract numeric claims from slide text (never notes).

    Returns (claims, ignored, notes, n_slides, slide_text) where slide_text maps
    slide index -> concatenated visible text (used for caveat detection).
    """
    units, notes, n_slides = read_deck(path)
    claims: list[Num] = []
    ignored: list[dict] = []
    slide_text: dict[int, str] = {}
    for si, shape, text, hunit, whole in units:
        slide_text.setdefault(si, "")
        slide_text[si] += text + "\n"
        mb = _BARE_INT_SHAPE.match(whole or "")
        if mb and int(mb.group(1)) <= max(n_slides, 1) + 0 and text.strip() == (whole or "").strip():
            ignored.append({"slide": si, "shape": shape, "raw": text.strip(), "rule": "slide_number",
                            "context": text.strip()})
            continue
        cs, ig = extract_numbers(text, header_unit=hunit)
        for c in cs:
            c.slide, c.shape = si, shape
            claims.append(c)
        for i in ig:
            i.update({"slide": si, "shape": shape})
            ignored.append(i)
    return claims, ignored, notes, n_slides, slide_text


# ---------------------------------------------------------------------------
# Source (content_analysis.json): schema, validation, number set
# ---------------------------------------------------------------------------

class SchemaError(ValueError):
    """content_analysis.json does not have the shape qc_numbers needs."""


# Text fields scanned for stated numbers (strings or lists of strings).
SOURCE_TEXT_FIELDS = ("abstract_summary", "introduction_bullets", "methods_summary",
                      "novel_contributions", "strengths", "limitations")
_CONFLICT_RE = re.compile(r"disagree|mismatch|conflict|discrepan|inconsisten|contradict|differs?\b|"
                          r"does not match|doesn't match|not match", re.I)


def _is_num(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def validate_source(data) -> list[str]:
    """Return a list of human-readable schema problems (empty = valid)."""
    errs: list[str] = []
    if not isinstance(data, dict):
        return ["top level must be a JSON object, got " + type(data).__name__]
    if not any(k in data for k in ("validated_numbers", "results", "tables", "abstract_summary")):
        errs.append("none of validated_numbers / results / tables / abstract_summary present; "
                    "see references/pipeline.md Step 7 (number schema)")

    def check_vn(lst, where):
        if not isinstance(lst, list):
            errs.append(f"{where} must be a list of objects")
            return
        for i, e in enumerate(lst):
            w = f"{where}[{i}]"
            if not isinstance(e, dict):
                errs.append(f"{w} must be an object with at least 'claim' and 'match'")
                continue
            if not isinstance(e.get("claim"), str) or not e.get("claim", "").strip():
                errs.append(f"{w}.claim must be a non-empty string")
            if "match" not in e or not isinstance(e["match"], bool):
                errs.append(f"{w}.match must be true or false (got {e.get('match', 'missing')!r})")
            for k in ("value", "alt_value"):
                if k in e and not _is_num(e[k]):
                    errs.append(f"{w}.{k} must be a number (got {e[k]!r}); put the unit in 'unit'")
            for k in ("unit", "source", "operator", "alt_source"):
                if k in e and not isinstance(e[k], str):
                    errs.append(f"{w}.{k} must be a string")
            if isinstance(e.get("unit"), str) and parse_unit(e["unit"]) is None:
                errs.append(f"{w}.unit {e['unit']!r} is not a recognised unit (use e.g. 'mM', '%', 'fold', 'p', 'n')")
            if "value" in e and _is_num(e["value"]) and "unit" not in e:
                errs.append(f"{w} has 'value' but no 'unit' (use \"\" for a unit-less number)")

    if "validated_numbers" in data:
        check_vn(data["validated_numbers"], "validated_numbers")
    if "results" in data:
        if not isinstance(data["results"], list):
            errs.append("results must be a list of objects")
        else:
            for i, r in enumerate(data["results"]):
                if not isinstance(r, dict):
                    errs.append(f"results[{i}] must be an object")
                    continue
                for k in ("finding", "explanation"):
                    if k in r and not isinstance(r[k], str):
                        errs.append(f"results[{i}].{k} must be a string")
                if "validated_numbers" in r:
                    check_vn(r["validated_numbers"], f"results[{i}].validated_numbers")
    if "tables" in data:
        if not isinstance(data["tables"], list):
            errs.append("tables must be a list of {caption, headers, rows}")
        else:
            for i, t in enumerate(data["tables"]):
                if not isinstance(t, dict) or not isinstance(t.get("rows"), list):
                    errs.append(f"tables[{i}] must be an object with a 'rows' list (optional 'headers', 'caption')")
                    continue
                if "headers" in t and not (isinstance(t["headers"], list) and all(isinstance(h, str) for h in t["headers"])):
                    errs.append(f"tables[{i}].headers must be a list of strings")
                for ri, row in enumerate(t["rows"]):
                    if not isinstance(row, list):
                        errs.append(f"tables[{i}].rows[{ri}] must be a list of cells")
    if "abstract_summary" in data and not isinstance(data["abstract_summary"], str):
        errs.append("abstract_summary must be a string")
    if "validation_warnings" in data and not (isinstance(data["validation_warnings"], list)
                                              and all(isinstance(w, str) for w in data["validation_warnings"])):
        errs.append("validation_warnings must be a list of strings")
    if "extra_numbers" in data:
        lst = data["extra_numbers"]
        if not isinstance(lst, list):
            errs.append("extra_numbers must be a list of {value, unit, source}")
        else:
            for i, e in enumerate(lst):
                if not isinstance(e, dict) or not _is_num(e.get("value")) or not isinstance(e.get("unit", ""), str):
                    errs.append(f"extra_numbers[{i}] must be an object with numeric 'value' and string 'unit'")
    return errs


def _structured(value, unit_text, operator, ref, text="", conflict=False, note=""):
    unit = parse_unit(unit_text or "") or NO_UNIT
    sv = str(value)
    dec = None if ("e" in sv.lower()) else (len(sv.split(".")[1]) if "." in sv else 0)
    kind = "p" if unit.canon == "p" else "n" if unit.canon == "n" else "pct" if unit.canon == "%" else "value"
    return Num(float(value), unit, text or f"{value} {unit_text or ''}".strip(), kind, operator or "",
               dec, _digits(sv.split("e")[0]), ref=ref, context=text, conflict=conflict, conflict_note=note)


def build_source(data: dict) -> list[Num]:
    """Flatten content_analysis.json into the source number set (after validate_source)."""
    out: list[Num] = []

    def add_text(text, ref, conflict=False, note=""):
        cs, _ = extract_numbers(text)
        for c in cs:
            c.ref, c.conflict, c.conflict_note = ref, conflict, note
            c.context = _clean_context(text, c.start, c.end, 60)
            out.append(c)

    def add_vn(lst, where):
        for i, e in enumerate(lst):
            ref = f"{where}[{i}]"
            conflict = e.get("match") is False
            note = f"validated_numbers match:false ({e.get('claim', '')[:60]})" if conflict else ""
            text = e.get("claim", "")
            if _is_num(e.get("value")):
                out.append(_structured(e["value"], e.get("unit"), e.get("operator"), ref + ".value",
                                       text, conflict, note))
            else:
                add_text(text, ref + ".claim", conflict, note)
            if _is_num(e.get("alt_value")):
                out.append(_structured(e["alt_value"], e.get("unit"), e.get("operator"), ref + ".alt_value",
                                       e.get("alt_source", text), conflict, note))
            if e.get("source"):
                add_text(e["source"], ref + ".source", conflict, note)
            if e.get("alt_source"):
                add_text(e["alt_source"], ref + ".alt_source", conflict, note)
            if conflict and _is_num(e.get("value")):
                add_text(text, ref + ".claim", True, note)

    if isinstance(data.get("validated_numbers"), list):
        add_vn(data["validated_numbers"], "validated_numbers")
    for i, r in enumerate(data.get("results", []) or []):
        for k in ("finding", "explanation"):
            if isinstance(r.get(k), str):
                add_text(r[k], f"results[{i}].{k}")
        if isinstance(r.get("validated_numbers"), list):
            add_vn(r["validated_numbers"], f"results[{i}].validated_numbers")
    for i, t in enumerate(data.get("tables", []) or []):
        heads = t.get("headers") or []
        hunits = {}
        for ci, h in enumerate(heads):
            m = re.search(r"[\(\[]\s*([^()\[\]]+?)\s*[\)\]]\s*$", h.strip())
            u = parse_unit(m.group(1)) if m else None
            if u and not u.empty:
                hunits[ci] = u
        if t.get("caption"):
            add_text(t["caption"], f"tables[{i}].caption")
        for ri, row in enumerate(t["rows"]):
            for ci, cell in enumerate(row):
                ref = f"tables[{i}].rows[{ri}][{ci}]"
                if _is_num(cell):
                    n = _structured(cell, "", "", ref, str(cell))
                    if ci in hunits:
                        n.unit = hunits[ci]
                    out.append(n)
                elif isinstance(cell, str):
                    cs, _ = extract_numbers(cell, header_unit=hunits.get(ci), small_bare_ok=True)
                    for c in cs:
                        c.ref, c.context = ref, cell
                        out.append(c)
    for k in SOURCE_TEXT_FIELDS:
        v = data.get(k)
        if isinstance(v, str):
            add_text(v, k)
        elif isinstance(v, list):
            for j, s in enumerate(v):
                if isinstance(s, str):
                    add_text(s, f"{k}[{j}]")
    for j, e in enumerate(data.get("extra_numbers", []) or []):
        out.append(_structured(e["value"], e.get("unit"), e.get("operator"), f"extra_numbers[{j}]",
                               e.get("source", "")))
    # abstract-vs-results disagreement flagged in validation_warnings
    for j, w in enumerate(data.get("validation_warnings", []) or []):
        if _CONFLICT_RE.search(w):
            add_text(w, f"validation_warnings[{j}]", True, "validation_warnings disagreement: " + w[:80])
    return out


def load_source(path: str | Path) -> list[Num]:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SchemaError(f"cannot read {path}: {exc}") from exc
    errs = validate_source(data)
    if errs:
        raise SchemaError("content_analysis.json schema problems:\n  - " + "\n  - ".join(errs))
    return build_source(data)


# ---------------------------------------------------------------------------
# Classification
# ---------------------------------------------------------------------------

MARKER_RE = re.compile(r"~|≈|\bapprox(?:imately|\.)?|\best\.|\bestimat\w*|\bcalculat\w*|\bcomputed?\b|\bderived\b|"
                       r"\broughly\b|\bour (?:calculation|estimate)\b|\bc\.\s?(?=\d)", re.I)
CAVEAT_RE = re.compile(r"caveat|discrepan|inconsisten|mismatch|disagree|conflict|differs?\b|contradict|"
                       r"abstract\s+(?:states|says|reports|claims|gives)|vs\.?\s+abstract|unresolved|⚠|"
                       r"note:|flagged|does not match|doesn't match", re.I)

STATUS_SEVERITY = {
    "MATCH": "OK", "DERIVED-LABELED": "OK", "DRIFT": "CRITICAL", "SOURCE-CONFLICT": "CRITICAL",
    "DERIVED-UNLABELED": "WARNING", "UNSOURCED": "WARNING",
}
REL_TOL = 1e-9


def _same_dim(a: Num, b: Num) -> bool:
    return a.unit.dim == b.unit.dim


def _close(x: float, y: float) -> bool:
    return math.isclose(x, y, rel_tol=1e-9, abs_tol=1e-12)


def _unit_compatible(c: Num, s: Num) -> bool:
    """Same unit dimension, or a missing unit on either side (reported as a note)."""
    if c.unit.dim == s.unit.dim:
        return True
    return (c.unit.empty or s.unit.empty) and c.kind not in ("p", "n") and s.kind not in ("p", "n")


def _op_ok(c: Num, s: Num) -> bool:
    return not (c.kind == "p" and s.kind == "p" and c.operator and s.operator and c.operator != s.operator)


def _match_exact(c: Num, s: Num) -> str | None:
    """Return a note ('' exact, 'rounded', 'unit_converted', ...) when s explains c, else None."""
    if c.kind == "p" or s.kind == "p":
        if not (c.kind == "p" and s.kind == "p"):
            return None
    if c.kind == "n" or s.kind == "n":
        if not (c.kind == "n" and s.kind == "n"):
            return None
    if not _unit_compatible(c, s):
        return None
    if not _op_ok(c, s):
        return None
    note = ""
    if c.unit.dim != s.unit.dim:           # a unit is missing on one side
        cv, sv = c.value, s.value
        note = "unit_missing_on_slide" if c.unit.empty else "unit_not_in_source"
    else:
        cv, sv = c.base, s.base
        if c.unit.scale != s.unit.scale and not _close(c.unit.scale, s.unit.scale):
            note = "unit_converted"
    if _close(cv, sv):
        if c.decimals is not None and s.decimals is not None and c.decimals > s.decimals:
            return None                       # false precision: handled as DRIFT (sig_fig_change)
        return note
    # honest rounding: the slide shows fewer decimals than the source and equals source rounded
    if c.decimals is not None and s.decimals is not None and c.decimals < s.decimals:
        sv_in_c = (sv / c.unit.scale) if c.unit.dim == s.unit.dim else sv
        cv_in_c = c.value
        if _close(_round_half_up(sv_in_c, c.decimals), cv_in_c):
            return "rounded"
    return None


def _round_half_up(x: float, nd: int) -> float:
    from decimal import Decimal, ROUND_HALF_UP
    q = Decimal(1).scaleb(-nd)
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


def _swap_adjacent(a: str, b: str) -> bool:
    if len(a) != len(b) or a == b:
        return False
    for i in range(len(a) - 1):
        t = list(a)
        t[i], t[i + 1] = t[i + 1], t[i]
        if "".join(t) == b:
            return True
    return False


def _near_miss(c: Num, s: Num) -> tuple[str, str] | None:
    """(strength 'strong'|'weak', reason) when slide number c looks like a corrupted s, else None."""
    if not _op_ok(c, s) and _close(c.value, s.value):
        return "strong", "p_operator_changed"
    if (c.kind == "p") != (s.kind == "p") or (c.kind == "n") != (s.kind == "n"):
        return None
    cu, su = c.unit, s.unit
    if cu.empty and c.kind == "value" and c.decimals == 0:
        return None            # a bare integer is a count/label: exact match or UNSOURCED, never a guess
    both_units = not cu.empty and not su.empty
    # same numeric value, different unit (mM vs uM, % vs mM, ...)
    if both_units and cu.canon != su.canon and cu.dim == su.dim and _close(c.value, s.value) \
            and not _close(c.base, s.base) and c.kind not in ("p", "n"):
        return "strong", "unit_mismatch"
    # percent vs fraction
    if (cu.empty and su.canon == "%" and 0 < abs(c.value) <= 1 and _close(c.value * 100, s.value)) or \
       (su.empty and cu.canon == "%" and 0 < abs(s.value) <= 1 and _close(s.value * 100, c.value)):
        return "strong", "percent_vs_fraction"
    if cu.dim != su.dim:
        return None
    # same dimension from here on
    cb, sb = c.base, s.base
    if c.decimals is not None and s.decimals is not None and cu.canon == su.canon:
        # false precision: more decimals than the source states, but inside its rounding band
        if c.decimals > s.decimals and abs(c.value - s.value) <= 0.5 * 10 ** (-s.decimals) + 1e-12:
            return "strong", "sig_fig_change"
        sig_c, sig_s = c.digits.lstrip("0"), s.digits.lstrip("0")
        if len(sig_c) >= 2 and len(sig_s) >= 2 and c.digits != s.digits:
            if c.decimals == s.decimals and len(c.digits) == len(s.digits):
                if _swap_adjacent(c.digits, s.digits):
                    return "strong", "digit_transposition"
                if sum(x != y for x, y in zip(c.digits, s.digits)) == 1:
                    return "strong", "one_digit_change"
        if c.digits.strip("0") == s.digits.strip("0") and c.digits.strip("0") and not _close(cb, sb):
            k = math.log10(abs(cb / sb)) if cb and sb else 0.5
            if abs(k - round(k)) < 1e-9 and 1 <= abs(round(k)) <= 3:
                return "strong", "decimal_shift"
    if sb != 0 and abs(cb - sb) / abs(sb) <= 0.10 and not _close(cb, sb):
        return "weak", "within_10pct"
    return None


def _distance(c: Num, s: Num) -> float:
    if c.unit.dim != s.unit.dim:
        return float("inf")
    cb, sb = c.base, s.base
    if cb == 0 or sb == 0:
        return abs(cb - sb)
    return abs(cb - sb) / abs(sb)


def _candidate_json(s: Num | None, reason: str = "", dist: float | None = None) -> dict | None:
    if s is None:
        return None
    d = {"value": s.value, "unit": s.unit.canon, "ref": s.ref, "text": s.context[:120]}
    if reason:
        d["reason"] = reason
    if dist is not None and math.isfinite(dist):
        d["relative_distance"] = round(dist, 4)
    return d


def _derivations(c: Num, src: list[Num]):
    """Yield (op_text, a, b) when slide number c is reproducible from two source numbers."""
    if c.kind in ("p", "n", "ph", "sci") or c.decimals is None:
        return
    sig = len(c.digits.lstrip("0"))
    if sig < 2 or (c.unit.empty and c.decimals == 0):
        return                  # one significant figure or a bare integer (count/label): too coincidence-prone
    tol = 0.5 * 10 ** (-c.decimals) + 1e-12
    pool = [s for s in src if s.kind not in ("p", "n", "ph") and s.value != 0]
    seen = set()
    cu = c.unit
    for i, a in enumerate(pool):
        for b in pool[i + 1:] if False else pool:
            if a is b:
                continue
            key = (a.base, a.unit.canon, b.base, b.unit.canon)
            if key in seen:
                continue
            seen.add(key)
            same = a.unit.dim == b.unit.dim
            # ratio / fold / percent-of  (any dims -> dimensionless)
            if same and b.base != 0:
                r = a.base / b.base
                if cu.empty or cu.canon in ("fold", "x", "×"):
                    if abs(r - c.value) <= tol:
                        yield f"{a.value:g}/{b.value:g} ratio", a, b
                elif cu.canon == "%":
                    if abs(r * 100 - c.value) <= tol:
                        yield f"{a.value:g}/{b.value:g} as percent", a, b
                    if abs((r - 1) * 100 - c.value) <= tol or abs(abs(r - 1) * 100 - abs(c.value)) <= tol:
                        yield f"percent change {b.value:g} -> {a.value:g}", b, a
            # difference and sum in the claim's own unit
            if same and cu.dim == a.unit.dim and not cu.empty:
                if abs(abs(a.base - b.base) / cu.scale - abs(c.value)) <= tol:
                    yield f"|{a.value:g} - {b.value:g}| difference", a, b
                if abs((a.base + b.base) / cu.scale - c.value) <= tol:
                    yield f"{a.value:g} + {b.value:g} sum", a, b
            elif same and cu.empty and a.unit.empty:
                if abs(abs(a.value - b.value) - abs(c.value)) <= tol:
                    yield f"|{a.value:g} - {b.value:g}| difference", a, b
                if abs(a.value + b.value - c.value) <= tol:
                    yield f"{a.value:g} + {b.value:g} sum", a, b


def _has_marker(c: Num, notes_text: str) -> bool:
    if MARKER_RE.search(c.paragraph):
        return True
    if notes_text:
        raw = c.raw.strip()
        for m in re.finditer(re.escape(raw), notes_text):
            if MARKER_RE.search(notes_text[max(0, m.start() - 80): m.end() + 80]):
                return True
    return False


def classify(claims: list[Num], source: list[Num], notes: dict[int, str] | None = None,
             slide_text: dict[int, str] | None = None) -> list[dict]:
    """Classify every slide claim; returns ledger rows (dicts)."""
    notes = notes or {}
    slide_text = slide_text or {}
    rows: list[dict] = []
    for c in claims:
        row = {"slide": c.slide, "shape": c.shape, "raw": c.raw, "value": c.value, "unit": c.unit.canon,
               "kind": c.kind, "operator": c.operator or None, "context": c.context}
        status, reason, ref, cand = None, "", None, None

        # 1. MATCH (possibly onto a conflicting source number)
        matches = [(s, n) for s in source if (n := _match_exact(c, s)) is not None]
        if matches:
            conflicts = [(s, n) for s, n in matches if s.conflict]
            if conflicts and not CAVEAT_RE.search(slide_text.get(c.slide, "")):
                s, n = conflicts[0]
                status, reason, ref, cand = "SOURCE-CONFLICT", s.conflict_note or "source flags this number", s.ref, s
            else:
                s, n = sorted(matches, key=lambda x: (x[0].conflict, x[1] != ""))[0]
                status, reason, ref, cand = "MATCH", (n or "exact") + (
                    "; conflict caveated on slide" if s.conflict else ""), s.ref, s
        else:
            # 2. strong near-miss -> DRIFT
            strong, weak = [], []
            for s in source:
                nm = _near_miss(c, s)
                if nm:
                    (strong if nm[0] == "strong" else weak).append((s, nm[1]))
            if strong:
                s, r = min(strong, key=lambda x: _distance(c, x[0]))
                status, reason, ref, cand = "DRIFT", r, s.ref, s
            else:
                # 3. derived
                der = next(iter(_derivations(c, source)), None)
                if der:
                    op, a, b = der
                    marked = _has_marker(c, notes.get(c.slide, ""))
                    status = "DERIVED-LABELED" if marked else "DERIVED-UNLABELED"
                    reason = f"{op} of source numbers {a.ref} and {b.ref}"
                    ref, cand = a.ref + " & " + b.ref, a
                elif weak:
                    s, r = min(weak, key=lambda x: _distance(c, x[0]))
                    status, reason, ref, cand = "DRIFT", r, s.ref, s
                else:
                    status, reason = "UNSOURCED", "no source number or near-miss"
                    comp = [s for s in source if _distance(c, s) != float("inf")]
                    cand = min(comp, key=lambda s: _distance(c, s)) if comp else None
        row.update({"status": status, "severity": STATUS_SEVERITY[status], "reason": reason,
                    "source_ref": ref,
                    "nearest_source": _candidate_json(cand, reason if status == "DRIFT" else "",
                                                      _distance(c, cand) if cand else None)})
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# Orchestration / CLI
# ---------------------------------------------------------------------------

def run(deck: str | Path, source_path: str | Path, ledger_path: str | Path | None = None,
        strict: bool = False) -> tuple[int, dict]:
    ledger: dict = {"tool": "qc_numbers", "deck": str(deck), "source": str(source_path)}
    try:
        source = load_source(source_path)
    except SchemaError as exc:
        ledger.update({"exit_code": EXIT_BLIND, "verdict": "BLIND", "error": str(exc)})
        _write_ledger(ledger, ledger_path)
        return EXIT_BLIND, ledger
    try:
        claims, ignored, notes, n_slides, slide_text = extract_slide_numbers(deck)
    except Exception as exc:  # unreadable pptx
        ledger.update({"exit_code": EXIT_BLIND, "verdict": "BLIND", "error": f"cannot read deck: {exc}"})
        _write_ledger(ledger, ledger_path)
        return EXIT_BLIND, ledger
    ledger["n_slides"], ledger["n_source_numbers"] = n_slides, len(source)
    if not source or not claims:
        why = "no numbers in source" if not source else "no numeric claims found on slides"
        ledger.update({"exit_code": EXIT_BLIND, "verdict": "BLIND", "error": why,
                       "claims": [], "ignored": ignored})
        _write_ledger(ledger, ledger_path)
        return EXIT_BLIND, ledger
    rows = classify(claims, source, notes, slide_text)
    counts: dict[str, int] = {}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    critical = [r for r in rows if r["severity"] == "CRITICAL"]
    warnings = [r for r in rows if r["severity"] == "WARNING"]
    code = EXIT_CRITICAL if critical or (strict and warnings) else EXIT_CLEAN
    ledger.update({"exit_code": code, "verdict": "CRITICAL" if critical else "CLEAN",
                   "counts": counts, "claims": rows, "ignored": ignored,
                   "ignored_by_rule": _count_rules(ignored)})
    _write_ledger(ledger, ledger_path)
    return code, ledger


def _count_rules(ignored: list[dict]) -> dict[str, int]:
    out: dict[str, int] = {}
    for i in ignored:
        out[i["rule"]] = out.get(i["rule"], 0) + 1
    return out


def _write_ledger(ledger: dict, path) -> None:
    if path:
        Path(path).write_text(json.dumps(ledger, ensure_ascii=False, indent=2), encoding="utf-8")


def format_report(ledger: dict) -> str:
    lines = []
    if ledger.get("verdict") == "BLIND":
        return f"BLIND (exit 2): {ledger.get('error')}\nThis is NOT a pass: no verification happened."
    lines.append(f"qc_numbers: {ledger['deck']}  vs  {ledger['source']}")
    lines.append(f"slides={ledger['n_slides']} source_numbers={ledger['n_source_numbers']} "
                 f"claims={len(ledger['claims'])} ignored={len(ledger['ignored'])}")
    lines.append("counts: " + ", ".join(f"{k}={v}" for k, v in sorted(ledger["counts"].items())))
    for sev in ("CRITICAL", "WARNING"):
        rows = [r for r in ledger["claims"] if r["severity"] == sev]
        if rows:
            lines.append(f"\n{sev} ({len(rows)})")
            for r in rows:
                nc = r.get("nearest_source")
                near = f"  nearest source: {nc['value']:g} {nc['unit']} @ {nc['ref']}" if nc else ""
                lines.append(f"  slide {r['slide']:>2} [{r['status']}] {r['raw']!r} ({r['reason']}){near}\n"
                             f"           ...{r['context']}...")
    lines.append(f"\nverdict: {ledger['verdict']} (exit {ledger['exit_code']})")
    return "\n".join(lines)


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass
    ap = argparse.ArgumentParser(description="Compare slide numbers with content_analysis.json")
    ap.add_argument("deck")
    ap.add_argument("--source", required=True, help="content_analysis.json")
    ap.add_argument("--ledger", help="write the full ledger JSON here")
    ap.add_argument("--strict", action="store_true", help="WARNING findings also exit 1")
    args = ap.parse_args(argv)
    code, ledger = run(args.deck, args.source, args.ledger, args.strict)
    print(format_report(ledger))
    return code


if __name__ == "__main__":
    sys.exit(main())

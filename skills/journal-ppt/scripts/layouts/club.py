"""Journal-club layouts (presenter preset, 261004).

  thank_you   closing slide: one huge 'Thank you' (54 pt) and one line, nothing else.

The 54 pt size is declared on the layout (extra_sizes), not on any theme, so every theme can use it and
qc_deck allows it only in decks that actually built this layout.
"""
from __future__ import annotations

from pptx.util import Inches

from deck_builder import CLUB_MARGIN, SLIDE_W, STAT_BIG_PT
from layouts import register


@register("thank_you", kind="closing", extra_sizes=(STAT_BIG_PT,),
          summary="huge 'Thank you' plus one line, no kicker / paper line / presenter block",
          source="journal-club presenter feedback 261004")
def thank_you(deck, headline: str = "Thank you", subline: str = "for listening. Please ask me anything."):
    slide = deck.new_slide()
    deck._next_num()
    w = SLIDE_W - 2 * CLUB_MARGIN
    h = deck.add_textbox(slide, CLUB_MARGIN, Inches(2.6), w, Inches(1.5), headline, role="stat_big")
    for r in h.text_frame.paragraphs[0].runs:
        r.font.color.rgb = deck.theme.palette["body"]
    s = deck.add_textbox(slide, CLUB_MARGIN, Inches(4.15), w, Inches(0.8), subline, role="title_main")
    for r in s.text_frame.paragraphs[0].runs:
        r.font.color.rgb = deck.theme.palette["subtitle"]
    return slide

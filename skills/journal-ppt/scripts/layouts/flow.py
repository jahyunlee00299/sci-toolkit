"""Flow family: pipeline rail (G11), walkthrough highlight, numbered agenda grid (G10), geometric divider (G7).

Registered layouts
  pipeline_rail         numbered nodes + duration pills on a dashed rail, one card per step (enzyme cascade)
  walkthrough_highlight the same rail with ONE step enlarged and the rest dimmed (Morph-linkable)
  numbered_grid_agenda  chapter-numbered grid, optional highlighted current item (recurring section slide)
  geometric_divider     section divider built from circles / bars / quarter-circles

All colors come from the active Theme palette, all text sizes from ROLES (the only extra size is stat_big 54 pt in
geometric_divider). Text fit is checked by construction with real Arial metrics; content that cannot fit raises
StyleError instead of overflowing.
"""
from __future__ import annotations

import math
from dataclasses import dataclass

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from deck_builder import (LINE_SPACING, MARGIN, SLIDE_W, STAT_BIG_PT, StyleError, contrast_ratio,
                          disable_autofit, required_contrast, set_line_spacing)
from qc_layout import text_width_pt

from . import register

KINDS = ("enzyme", "substrate", "product", "cofactor")
SLIDE_W_IN = SLIDE_W / 914400.0
MARGIN_IN = MARGIN / 914400.0
LINE_H = 1.2                 # Arial single-line pitch / pt, same factor qc_layout uses
WRAP_SLACK_IN = 0.05         # keep our wrap strictly narrower than the QC estimator's


# ============================================================================ color policy
def _hex(c) -> str:
    return str(c)


def _uniq(colors):
    seen, out = set(), []
    for c in colors:
        if _hex(c) not in seen:
            seen.add(_hex(c))
            out.append(c)
    return out


class _Pal:
    """Theme-derived color decisions: which fills, which ink on a fill. Only palette colors are returned."""

    def __init__(self, theme):
        self.theme = theme
        p = theme.palette
        self.p = p
        self.bg = theme.bg
        self.inks = _uniq([theme.bg, p["body"], p["title_main"], p["title_bar_text"],
                           p["table_hdr_text"], p["caption"]])
        self.body = p["body"]
        self.title = p["title_main"]
        self.primary = self._fill_for(("subtitle", "stat", "accent", "title_main", "body"), ())
        self.secondary = self._fill_for(("accent", "stat", "subtitle", "title_main", "body", "caption"),
                                        (_hex(self.primary),), min_ink=3.0)
        self.rail = self._muted()
        self.vivid = self._vivid()

    def best_ink(self, fill):
        return max(self.inks, key=lambda c: contrast_ratio(c, fill))

    def _fill_for(self, keys, avoid, min_ink=4.5):
        for k in keys:
            c = self.p[k]
            if _hex(c) in avoid:
                continue
            if contrast_ratio(c, self.bg) >= 3.0 and contrast_ratio(self.best_ink(c), c) >= min_ink:
                return c
        for k in ("body", "title_main", "caption"):
            if _hex(self.p[k]) not in avoid:
                return self.p[k]
        return self.p["body"]

    def _muted(self):
        for k in ("caption", "body"):
            if contrast_ratio(self.p[k], self.bg) >= 4.5:
                return self.p[k]
        return self.p["body"]

    def _vivid(self):
        keys = ("accent", "subtitle", "stat", "title_main", "body", "caption", "highlight_bg", "zebra")
        cols = [self.p[k] for k in keys if _hex(self.p[k]) != _hex(self.bg)
                and contrast_ratio(self.p[k], self.bg) >= 1.25]
        return _uniq(cols) or [self.p["body"]]

    def text_on(self, fill, candidates, pt, bold):
        """First candidate reaching the WCAG ratio on `fill`; else the best palette ink."""
        need = required_contrast(pt, bold)
        for c in candidates:
            if contrast_ratio(c, fill) >= need:
                return c
        return self.best_ink(fill)

    def ink_for_graphic(self, color):
        """Color to draw text in when the graphic color itself is used as text (18 pt bold numerals)."""
        return color if contrast_ratio(color, self.bg) >= 4.5 else self.body


# ============================================================================ low-level shape helpers
def _iv(x):
    return int(round(x * 914400))


def _style_shape(shp, fill=None, line=None, line_w=1.0, dash=None, alpha=None, name=None, rot=0):
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if alpha is not None:
            clr = shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 1000)))
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash is not None:
            shp.line.dash_style = dash
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = True
    disable_autofit(shp.text_frame)
    if rot:
        shp.rotation = rot
    if name:
        shp.name = name
    return shp


def _shape(slide, mso, l, t, w, h, **style):
    shp = slide.shapes.add_shape(mso, _iv(l), _iv(t), _iv(w), _iv(h))
    return _style_shape(shp, **style)


def _write_run(deck, run, spec, color, bold=None, italic=None):
    run.font.name = deck.theme.font
    run.font.size = Pt(spec["pt"])
    run.font.bold = spec["bold"] if bold is None else bold
    run.font.italic = spec["italic"] if italic is None else italic
    run.font.color.rgb = color


@dataclass
class _Para:
    text: str
    role: str
    color: RGBColor
    bold: bool | None = None
    italic: bool | None = None
    after: float = 0.0
    line_spacing: float | None = None


def _para_metrics(deck, para: _Para):
    spec = deck._spec(para.role)
    bold = spec["bold"] if para.bold is None else para.bold
    return spec["pt"], bold


def _wrap(deck, text, pt, bold, width_in):
    """Greedy word wrap with real Arial metrics; a word wider than the box is a hard error."""
    fam = deck.theme.font
    limit = width_in - WRAP_SLACK_IN
    lines, cur = [], ""
    for w in text.split():
        if text_width_pt(w, fam, bold, pt) / 72.0 > limit:
            raise StyleError(f"word {w!r} is wider than its {width_in:.2f} in box at {pt:g} pt; shorten it")
        trial = (cur + " " + w).strip()
        if cur and text_width_pt(trial, fam, bold, pt) / 72.0 > limit:
            lines.append(cur)
            cur = w
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return lines or [""]


def _paras_height(deck, paras, width_in):
    """Text height in inches of stacked paragraphs (line pitch x role spacing + space_after between)."""
    total_pt = 0.0
    for i, pr in enumerate(paras):
        pt, bold = _para_metrics(deck, pr)
        ls = LINE_SPACING[pr.role] if pr.line_spacing is None else pr.line_spacing
        total_pt += len(_wrap(deck, pr.text, pt, bold, width_in)) * pt * LINE_H * ls
        if i < len(paras) - 1:
            total_pt += pr.after
    return total_pt / 72.0


def _rich_box(deck, slide, l, t, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
              inset_lr=0.08, inset_tb=0.04, name=None):
    """One text box, several paragraphs in different roles. Raises StyleError if the text cannot fit."""
    need = _paras_height(deck, paras, w - 2 * inset_lr) + 2 * inset_tb
    if need > h + 0.02:
        raise StyleError(f"text needs {need:.2f} in but its box is {h:.2f} in tall: {paras[0].text[:40]!r}")
    tb = slide.shapes.add_textbox(_iv(l), _iv(t), _iv(w), _iv(h))
    tf = tb.text_frame
    tf.word_wrap = True
    disable_autofit(tf)
    tf.margin_left = tf.margin_right = _iv(inset_lr)
    tf.margin_top = tf.margin_bottom = _iv(inset_tb)
    tf.vertical_anchor = anchor
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        set_line_spacing(p, LINE_SPACING[pr.role] if pr.line_spacing is None else pr.line_spacing)
        p.space_after = Pt(pr.after if i < len(paras) - 1 else 0)
        _write_run(deck, p.add_run(), deck._spec(pr.role), pr.color, pr.bold, pr.italic)
        p.runs[0].text = pr.text
    if name:
        tb.name = name
    return tb


def _shape_text(deck, shp, text, role, color, align=PP_ALIGN.CENTER, italic=False):
    """Text inside an autoshape (numerals, pills): zero insets, centered both ways."""
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    set_line_spacing(p, LINE_SPACING[role])
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    _write_run(deck, run, deck._spec(role), color, italic=italic)


def _arrow(slide, p1, p2, color, width_pt, dash, src, dst, src_idx=6, dst_idx=2, name=None):
    """Straight connector with a real glue to both shapes (ellipse sites: 6 = right, 2 = left)."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, _iv(p1[0]), _iv(p1[1]), _iv(p2[0]), _iv(p2[1]))
    c.line.color.rgb = color
    c.line.width = Pt(width_pt)
    if dash is not None:
        c.line.dash_style = dash
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    cnv = c._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    for tag, shp, idx in (("a:stCxn", src, src_idx), ("a:endCxn", dst, dst_idx)):
        el = etree.SubElement(cnv, qn(tag))
        el.set("id", str(shp.shape_id))
        el.set("idx", str(idx))
    if name:
        c.name = name
    return c


def _bezier_loop(slide, a, b, apex_h, color, width_pt, name=None):
    """Free-form cubic arc from point a up over to point b, arrow head at b (pointing down onto the node)."""
    (xa, ya), (xb, yb) = a, b
    y0 = min(ya, yb)
    yc = y0 - apex_h / 0.75                     # cubic with both controls at yc peaks exactly apex_h above y0
    left, top = min(xa, xb), yc
    w, h = abs(xb - xa), max(ya, yb) - yc
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _iv(left), _iv(top), _iv(w), _iv(h))
    _style_shape(shp, fill=None, line=color, line_w=width_pt, name=name)
    spPr = shp._element.spPr
    spPr.remove(spPr.find(qn("a:prstGeom")))
    W, H = _iv(w), _iv(h)

    def pt(x, y):
        return f'<a:pt x="{_iv(x - left)}" y="{_iv(y - top)}"/>'

    xml = (f'<a:custGeom xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:avLst/><a:gdLst/>'
           f'<a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
           f'<a:path w="{W}" h="{H}" fill="none"><a:moveTo>{pt(xa, ya)}</a:moveTo>'
           f'<a:cubicBezTo>{pt(xa, yc)}{pt(xb, yc)}{pt(xb, yb)}</a:cubicBezTo></a:path></a:pathLst></a:custGeom>')
    spPr.find(qn("a:xfrm")).addnext(etree.fromstring(xml))
    ln = shp.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return shp


# ============================================================================ pipeline rail (G11)
@dataclass
class _Step:
    label: str
    sublabel: str
    duration: str
    kind: str
    note: str


def _parse_steps(steps):
    if not isinstance(steps, (list, tuple)) or not steps:
        raise StyleError("steps must be a non-empty list of {label, sublabel, duration, kind, note} dicts")
    if not 3 <= len(steps) <= 7:
        raise StyleError(f"a rail takes 3-7 steps (got {len(steps)}); split the cascade over two slides")
    out = []
    for i, s in enumerate(steps):
        if not isinstance(s, dict) or not str(s.get("label", "")).strip():
            raise StyleError(f"step {i}: needs a dict with a non-empty 'label'")
        kind = s.get("kind", "enzyme")
        if kind not in KINDS:
            raise StyleError(f"step {i}: kind {kind!r} not in {KINDS}")
        out.append(_Step(str(s["label"]).strip(), str(s.get("sublabel") or "").strip(),
                         str(s.get("duration") or "").strip(), kind, str(s.get("note") or "").strip()))
    return out


def _parse_loop(loop, n):
    if loop is None:
        return None
    try:
        a, b, label = loop
    except (TypeError, ValueError):
        raise StyleError("cofactor_loop must be (from_idx, to_idx, label) with 0-based step indices") from None
    if not (isinstance(a, int) and isinstance(b, int) and 0 <= a < n and 0 <= b < n and a != b):
        raise StyleError(f"cofactor_loop indices must be two different 0-based step indices < {n}")
    if not str(label).strip():
        raise StyleError("cofactor_loop needs a label")
    return a, b, str(label).strip()


def _check_active(active, n):
    if not (isinstance(active, int) and not isinstance(active, bool) and 0 <= active < n):
        raise StyleError(f"active must be a 0-based step index < {n}")
    return active


_TOP = 1.30          # first usable y below any header
_BAND_TOP, _BAND_H = 6.20, 0.70
_LOOP_APEX, _LOOP_LABEL_H, _LOOP_ZONE = 0.50, 0.32, 0.95
_CARD_PAD_LR, _CARD_PAD_TB = 0.10, 0.08


@dataclass(frozen=True)
class _Tier:
    """Type/size tier of a rail: fewer steps -> bigger roles (closed type scale only)."""
    node_d: float
    active_scale: float
    node_role: str
    pill_role: str
    pill_h: float
    pill_gap: float
    card_gap: float
    tag_role: str
    label_role: str
    sub_role: str
    act_label_role: str
    act_sub_role: str
    dim_label_role: str
    card_max_w: float
    card_active_max_w: float


_TIERS = {
    "big": _Tier(0.95, 1.3, "title_bar", "fig_subtitle", 0.38, 0.16, 0.22, "table_cell", "subtitle",
                 "fig_subtitle", "title_bar", "body", "fig_subtitle", 3.6, 5.4),
    "mid": _Tier(0.70, 1.5, "subtitle", "table_header", 0.30, 0.10, 0.14, "caption", "fig_subtitle",
                 "table_cell", "subtitle", "fig_subtitle", "table_header", 2.9, 4.8),
    "small": _Tier(0.62, 1.5, "subtitle", "table_header", 0.30, 0.10, 0.14, "caption", "table_header",
                   "table_cell", "subtitle", "fig_subtitle", "table_header", 2.9, 4.8),
}


class _Rail:
    """Geometry + drawing of one rail slide. ``active=None`` -> pipeline_rail, else walkthrough_highlight.

    The rail center line and the card top are computed from BOTH states (normal and every possible active step),
    so a pipeline_rail slide and a walkthrough_highlight slide of the same content put them at the same y."""

    def __init__(self, deck, steps, loop, takeaway, active=None):
        self.deck, self.pal = deck, _Pal(deck.theme)
        self.steps = _parse_steps(steps)
        self.n = len(self.steps)
        self.loop = _parse_loop(loop, self.n)
        self.takeaway = (takeaway or "").strip()
        self.active = None if active is None else _check_active(active, self.n)
        self.t = _TIERS["big" if self.n <= 4 else ("mid" if self.n <= 6 else "small")]
        self.D = self.t.node_d
        self.has_pills = any(s.duration for s in self.steps)
        self._plan_vertical()

    # ---- horizontal slots
    def slots(self, active):
        W = SLIDE_W_IN - 2 * MARGIN_IN
        if active is None:
            weights = [1.0] * self.n
        else:
            a = 1.9 if self.n >= 6 else 2.2
            weights = [a if i == active else 1.0 for i in range(self.n)]
        tot = sum(weights)
        x, out = MARGIN_IN, []
        for wgt in weights:
            sw = W * wgt / tot
            out.append((x, sw))
            x += sw
        return out

    def card_w(self, i, active):
        _, sw = self.slots(active)[i]
        cap = self.t.card_active_max_w if (active is not None and i == active) else self.t.card_max_w
        return min(sw - 0.12, cap)

    # ---- paragraphs per card state
    def _tag_color(self, st, fill):
        kc = self.pal.primary if st.kind in ("enzyme", "substrate") else self.pal.secondary
        return self.pal.text_on(fill, [kc, self.pal.body], 9.5, False)

    def card_paras(self, i, state):
        """state: 'normal' | 'active' | 'dim'."""
        st, pal = self.steps[i], self.pal
        fill = pal.p["highlight_bg"] if state != "dim" else pal.bg
        body = pal.text_on(fill, [pal.body], 12, False)
        mute = pal.text_on(fill, [pal.p["caption"], pal.body], 12, False)
        title = pal.text_on(fill, [pal.title, pal.body], 14, True)
        t = self.t
        if state == "dim":
            return [_Para(st.label, t.dim_label_role, mute, bold=True, italic=False)]
        paras = [_Para(st.kind.upper(), t.tag_role, self._tag_color(st, fill), bold=True, italic=False, after=2)]
        if state == "active":
            role, ls = t.act_sub_role, 1.0
            paras.append(_Para(st.label, t.act_label_role, title, bold=True, italic=False, after=2))
        else:
            role, ls = t.sub_role, None
            paras.append(_Para(st.label, t.label_role, title, bold=True, italic=False, after=2))
        if st.sublabel:
            paras.append(_Para(st.sublabel, role, body, bold=False, italic=False, after=2, line_spacing=ls))
        if st.note:
            paras.append(_Para(st.note, role, mute, bold=False, italic=True, line_spacing=ls))
        return paras

    def _raw_card_h(self, i, state, active):
        w = self.card_w(i, active) - 2 * _CARD_PAD_LR
        return _paras_height(self.deck, self.card_paras(i, state), w) + 2 * _CARD_PAD_TB + 0.08

    def card_h(self, i, state, active):
        """Normal cards share one height (the tallest) so the row reads as a set; active/dim size to content."""
        if state == "normal":
            return self.normal_h
        return self._raw_card_h(i, state, active)

    def _plan_vertical(self):
        self.normal_h = max(self._raw_card_h(i, "normal", None) for i in range(self.n))
        normal = self.normal_h
        act = max(self._raw_card_h(i, "active", i) for i in range(self.n))
        self.reserve_h = max(normal, act, 1.9 if self.t is _TIERS["big"] else 1.3)
        rmax = self.D * self.t.active_scale / 2
        above = _LOOP_ZONE if self.loop else 0.0
        pill_zone = (self.t.pill_gap + self.t.pill_h) if self.has_pills else 0.0
        below = rmax + pill_zone + self.t.card_gap + self.reserve_h
        bottom_limit = (_BAND_TOP - 0.12) if self.takeaway else 6.90
        avail = bottom_limit - _TOP
        block = above + rmax + below
        if block > avail:
            raise StyleError(f"rail content needs {block:.2f} in of height but only {avail:.2f} in are free; "
                             f"shorten label/sublabel/note text")
        dy = min((avail - block) / 2, 0.9)
        self.c = _TOP + dy + above + rmax
        self.card_top = self.c + rmax + pill_zone + self.t.card_gap
        self.pill_top = self.c + rmax + self.t.pill_gap
        self.loop_top = _TOP + dy
        self.rmax = rmax

    # ---- drawing
    def build(self, slide):
        slots = self.slots(self.active)
        nodes, cards = [], []
        for i, st in enumerate(self.steps):
            nodes.append(self._node(slide, i, slots[i], st))
        for i, st in enumerate(self.steps):
            if self.has_pills and st.duration:
                self._pill(slide, i, slots[i])
            cards.append(self._card(slide, i, slots[i]))
        self._arrows(slide, nodes, slots)
        if self.loop:
            self._loop(slide, nodes, slots)
        if self.takeaway:
            self._takeaway(slide)

    def _node_diam(self, i):
        return self.D * self.t.active_scale if i == self.active else self.D

    def _node_style(self, st, dim):
        pal = self.pal
        if dim:
            return dict(fill=pal.bg, line=pal.rail, line_w=1.5, dash=None, ink=pal.rail)
        main = pal.primary if st.kind in ("enzyme", "substrate") else pal.secondary
        if st.kind == "enzyme":
            return dict(fill=main, line=None, line_w=1, dash=None, ink=pal.best_ink(main))
        if st.kind == "product":
            return dict(fill=main, line=None, line_w=1, dash=None, ink=pal.best_ink(main))
        dash = MSO_LINE.DASH if st.kind == "cofactor" else None
        return dict(fill=pal.bg, line=main, line_w=3.0, dash=dash, ink=pal.ink_for_graphic(main))

    def _node(self, slide, i, slot, st):
        dim = self.active is not None and i != self.active
        d = self._node_diam(i)
        cx = slot[0] + slot[1] / 2
        s = self._node_style(st, dim)
        shp = _shape(slide, MSO_SHAPE.OVAL, cx - d / 2, self.c - d / 2, d, d, fill=s["fill"], line=s["line"],
                     line_w=s["line_w"], dash=s["dash"])
        _shape_text(self.deck, shp, str(i + 1), self.t.node_role if self.t.node_role == "title_bar" or i != self.active else "title_bar", s["ink"])
        self.deck.name_shape(shp, f"pr_node_{i}")
        return shp

    def _pill(self, slide, i, slot):
        st, pal = self.steps[i], self.pal
        dim = self.active is not None and i != self.active
        tw = text_width_pt(st.duration, self.deck.theme.font, True, self.deck._spec(self.t.pill_role)["pt"]) / 72.0
        w = max(0.85, tw + 0.34)
        if w > slot[1] - 0.08:
            raise StyleError(f"duration {st.duration!r} is too long for its pill column ({slot[1]:.2f} in)")
        cx = slot[0] + slot[1] / 2
        if dim:
            fill, line, ink = None, pal.rail, pal.rail
        else:
            fill, line, ink = pal.primary, None, pal.best_ink(pal.primary)
        shp = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx - w / 2, self.pill_top, w, self.t.pill_h, fill=fill,
                     line=line, line_w=1.0)
        shp.adjustments[0] = 0.5
        _shape_text(self.deck, shp, st.duration, self.t.pill_role, ink)
        self.deck.name_shape(shp, f"pr_pill_{i}")
        return shp

    def _card(self, slide, i, slot):
        pal = self.pal
        state = "normal" if self.active is None else ("active" if i == self.active else "dim")
        w = self.card_w(i, self.active)
        h = self.card_h(i, state, self.active)
        left = slot[0] + (slot[1] - w) / 2
        if state == "dim":
            rect = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, self.card_top, w, h, fill=None,
                          line=pal.rail, line_w=0.75, dash=MSO_LINE.DASH)
        else:
            lc = pal.primary if state == "active" else pal.rail
            rect = _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, self.card_top, w, h, fill=pal.p["highlight_bg"],
                          line=lc, line_w=2.25 if state == "active" else 0.75)
        rect.adjustments[0] = 0.08
        self.deck.name_shape(rect, f"pr_card_{i}")
        paras = self.card_paras(i, state)
        tb = _rich_box(self.deck, slide, left + _CARD_PAD_LR / 2, self.card_top + _CARD_PAD_TB / 2,
                       w - _CARD_PAD_LR, h - _CARD_PAD_TB, paras, inset_lr=_CARD_PAD_LR / 2,
                       inset_tb=_CARD_PAD_TB / 2)
        self.deck.name_shape(tb, f"pr_text_{i}")
        return rect

    def _arrows(self, slide, nodes, slots):
        dim_color = self.pal.rail
        for i in range(self.n - 1):
            xa = slots[i][0] + slots[i][1] / 2 + self._node_diam(i) / 2
            xb = slots[i + 1][0] + slots[i + 1][1] / 2 - self._node_diam(i + 1) / 2
            if xb - xa < 0.25:
                raise StyleError("nodes are too close for an arrow; use fewer steps")
            hot = self.active is not None and self.active in (i, i + 1)
            color = self.pal.primary if (self.active is None or hot) else dim_color
            _arrow(slide, (xa, self.c), (xb, self.c), color, 2.25, MSO_LINE.DASH, nodes[i], nodes[i + 1],
                   name="!!pr_arrow_%d" % i)

    def _loop(self, slide, nodes, slots):
        a, b, label = self.loop
        pal = self.pal
        top_of = lambda i: (slots[i][0] + slots[i][1] / 2, self.c - self._node_diam(i) / 2)   # noqa: E731
        pa, pb = top_of(a), top_of(b)
        color = pal.secondary if contrast_ratio(pal.secondary, pal.bg) >= 3.0 else pal.body
        shp = _bezier_loop(slide, pa, pb, _LOOP_APEX, color, 2.25, name="!!pr_loop")
        text_col = pal.text_on(pal.bg, [pal.secondary, pal.body], 14, True)
        lw = text_width_pt(label, self.deck.theme.font, True, 14) / 72.0 + 0.3
        mid = (pa[0] + pb[0]) / 2
        left = min(max(mid - lw / 2, MARGIN_IN), SLIDE_W_IN - MARGIN_IN - lw)
        apex_y = min(pa[1], pb[1]) - _LOOP_APEX
        tb = _rich_box(self.deck, slide, left, apex_y - _LOOP_LABEL_H - 0.04, lw, _LOOP_LABEL_H,
                       [_Para(label, "fig_subtitle", text_col, bold=True, italic=False)],
                       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, inset_lr=0.05, inset_tb=0.02)
        self.deck.name_shape(tb, "pr_looplabel")
        return shp

    def _takeaway(self, slide):
        pal, deck = self.pal, self.deck
        left, width = MARGIN_IN, SLIDE_W_IN - 2 * MARGIN_IN
        fill = pal.p["highlight_bg"]
        _shape(slide, MSO_SHAPE.RECTANGLE, left, _BAND_TOP, width, _BAND_H, fill=fill,
               line=pal.rail, line_w=0.75, name="deco:takeaway_band")
        _shape(slide, MSO_SHAPE.RECTANGLE, left, _BAND_TOP, 0.12, _BAND_H, fill=pal.secondary,
               name="deco:takeaway_bar")
        col = pal.text_on(fill, [pal.title, pal.body], 18, True)
        lines = _wrap(deck, self.takeaway, 18, True, width - 0.42 - 0.2)
        if len(lines) > 1:
            raise StyleError("takeaway must fit on one line at 18 pt bold (about 90 characters)")
        _rich_box(deck, slide, left + 0.3, _BAND_TOP + 0.05, width - 0.42, _BAND_H - 0.1,
                  [_Para(self.takeaway, "subtitle", col)], anchor=MSO_ANCHOR.MIDDLE, inset_lr=0.1, inset_tb=0.02,
                  name="pr_takeaway")


def _rail_slide(deck, title, rail: _Rail):
    slide = deck.content_slide(title or "Cascade overview")
    rail.build(slide)
    return slide


@register("pipeline_rail", kind="content",
          summary="Numbered nodes + duration pills on a dashed rail, one card per step, optional cofactor loop",
          source="G11 Cell Press graphical-abstract pipeline rail")
def pipeline_rail(deck, steps, title="Cascade overview", cofactor_loop=None, takeaway=None):
    """Enzyme-cascade flagship layout.

    steps: 3-7 dicts {label, sublabel, duration, kind, note}; kind in enzyme|substrate|product|cofactor
      (node style: filled disc | solid ring | filled disc in the second color | dashed ring, plus a text tag).
    cofactor_loop: (from_idx, to_idx, label), 0-based indices into steps, drawn as a curved arrow above the rail.
    takeaway: one line (<= ~90 chars at 18 pt bold) in a band at the bottom.
    """
    return _rail_slide(deck, title, _Rail(deck, steps, cofactor_loop, takeaway))


@register("walkthrough_highlight", kind="content",
          summary="The pipeline rail with one step enlarged and the rest dimmed; shape names match pipeline_rail for Morph",
          source="G11 rail, walkthrough variant")
def walkthrough_highlight(deck, steps, active, title="Cascade overview", cofactor_loop=None, takeaway=None,
                          morph=False):
    """Same content as pipeline_rail plus ``active`` (0-based). Shapes carry the stable names !!pr_node_<i>,
    !!pr_card_<i>, !!pr_text_<i>, !!pr_pill_<i>, !!pr_arrow_<i>, !!pr_loop so Deck.morph() animates between a
    pipeline_rail slide and any walkthrough_highlight slide. ``morph=True`` sets the Morph transition on this slide."""
    if active is None:
        raise StyleError("walkthrough_highlight needs active (0-based index of the highlighted step)")
    slide = _rail_slide(deck, title, _Rail(deck, steps, cofactor_loop, takeaway, active=active))
    if morph:
        deck.morph(slide)
    return slide


# ============================================================================ numbered grid agenda (G10)
def _parse_items(items):
    if not isinstance(items, (list, tuple)) or not items:
        raise StyleError("items must be a non-empty list")
    if not 4 <= len(items) <= 8:
        raise StyleError(f"agenda grid takes 4-8 items (got {len(items)})")
    out = []
    for i, it in enumerate(items):
        if isinstance(it, str):
            label, desc = it, ""
        elif isinstance(it, dict):
            label, desc = it.get("label", ""), it.get("desc", "")
        else:
            label, desc = (tuple(it) + ("",))[:2]
        if not str(label).strip():
            raise StyleError(f"item {i}: empty label")
        out.append((str(label).strip(), str(desc or "").strip()))
    return out


@register("numbered_grid_agenda", kind="agenda",
          summary="Chapter-numbered grid of outlined number squares; optional highlighted current item",
          source="G10 Slidesgo chapter-numbered grid")
def numbered_grid_agenda(deck, items, title="Agenda", current=None):
    """items: 4-8 strings, (label, desc) pairs or {label, desc} dicts. current: 0-based index to highlight
    (filled square + tinted panel) so the slide can recur as a section marker."""
    its = _parse_items(items)
    n = len(its)
    if current is not None and not (isinstance(current, int) and 0 <= current < n):
        raise StyleError(f"current must be a 0-based index < {n}")
    pal = _Pal(deck.theme)
    slide = deck.content_slide(title)
    cols = 2 if n == 4 else (3 if n <= 6 else 4)
    rows = math.ceil(n / cols)
    area_top, area_bot = 1.45, 6.90
    col_w = (SLIDE_W_IN - 2 * MARGIN_IN) / cols
    sq_w, sq_h, gap = 1.05, 0.76, 0.16
    text_w = col_w - 0.5
    panel_fill = pal.p["highlight_bg"]

    def paras(i, label, desc):
        cur = i == current
        lab = pal.text_on(panel_fill if cur else pal.bg, [pal.title, pal.body], 18, True)
        out = [_Para(label, "subtitle", lab, after=6)]
        if desc:
            fill = panel_fill if cur else pal.bg
            out.append(_Para(desc, "table_cell", pal.text_on(fill, [pal.p["caption"], pal.body], 12, False)))
        return out

    text_h = max(_paras_height(deck, paras(i, *its[i]), text_w - 0.16) for i in range(n)) + 0.12
    cell_h = sq_h + gap + text_h
    pitch = (area_bot - area_top) / rows
    if cell_h + 0.4 > pitch:
        raise StyleError(f"agenda text needs {cell_h + 0.4:.2f} in per row but only {pitch:.2f} in are free; "
                         f"shorten labels/descriptions")
    for i, (label, desc) in enumerate(its):
        r, c = divmod(i, cols)
        in_row = cols if r < rows - 1 else n - cols * (rows - 1)
        x0 = MARGIN_IN + (SLIDE_W_IN - 2 * MARGIN_IN - in_row * col_w) / 2 + c * col_w
        cx, y0 = x0 + col_w / 2, area_top + r * pitch + (pitch - cell_h) / 2
        cur = i == current
        if cur:
            _shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x0 + 0.1, y0 - 0.14, col_w - 0.2, cell_h + 0.28,
                   fill=panel_fill, line=pal.primary, line_w=1.5, name="deco:current_panel").adjustments[0] = 0.06
        sq = _shape(slide, MSO_SHAPE.RECTANGLE, cx - sq_w / 2, y0, sq_w, sq_h,
                    fill=pal.primary if cur else pal.bg, line=pal.primary, line_w=1.5)
        ink = pal.best_ink(pal.primary) if cur else pal.ink_for_graphic(pal.primary)
        _shape_text(deck, sq, f"{i + 1:02d}", "title_bar", ink)
        _rich_box(deck, slide, cx - text_w / 2, y0 + sq_h + gap, text_w, text_h, paras(i, label, desc),
                  align=PP_ALIGN.CENTER, inset_lr=0.08, inset_tb=0.04)
    return slide


# ============================================================================ geometric divider (G7)
def _wrap_n(deck, text, role, width_in, max_lines):
    spec = deck._spec(role)
    lines = _wrap(deck, text, spec["pt"], spec["bold"], width_in)
    if len(lines) > max_lines:
        raise StyleError(f"{role} text wraps to {len(lines)} lines (max {max_lines}) in {width_in:.1f} in: {text[:40]!r}")
    return len(lines)


def _compose_circles(slide, col):
    _shape(slide, MSO_SHAPE.OVAL, 7.1, 0.45, 6.4, 6.4, fill=col(0), name="deco:circle_big")
    _shape(slide, MSO_SHAPE.OVAL, 6.75, 3.75, 3.5, 3.5, fill=col(1), alpha=72, name="deco:circle_overlap")
    _shape(slide, MSO_SHAPE.DONUT, 7.3, 0.35, 2.1, 2.1, fill=col(2), name="deco:ring").adjustments[0] = 0.28
    _shape(slide, MSO_SHAPE.OVAL, 11.7, 5.3, 1.6, 1.6, fill=col(3), name="deco:circle_small")


def _compose_bars(slide, col):
    heights = [4.2, 5.9, 3.4, 6.6, 5.0, 6.2, 4.4]
    x0, w, gap = 6.55, 0.88, 0.10
    for i, h in enumerate(heights):
        shp = _shape(slide, MSO_SHAPE.ROUND_2_SAME_RECTANGLE, x0 + i * (w + gap), 7.5 - h, w, h + 0.4,
                     fill=col(i), name=f"deco:bar_{i}")
        shp.adjustments[0] = 0.5
        shp.adjustments[1] = 0.0
    _shape(slide, MSO_SHAPE.OVAL, x0 + 2 * (w + gap) + 0.07, 7.5 - 3.4 - 1.05, 0.74, 0.74, fill=col(1),
           name="deco:bar_dot")


def _compose_quarter(slide, col):
    s, x0, y0 = 2.95, 0.55, 0.85
    _shape(slide, MSO_SHAPE.PIE_WEDGE, x0, y0, s, s, fill=col(0), name="deco:quarter_a")
    _shape(slide, MSO_SHAPE.RECTANGLE, x0 + s, y0, s, s, fill=col(2), name="deco:tile_b")
    _shape(slide, MSO_SHAPE.PIE_WEDGE, x0 + s, y0, s, s, fill=col(1), rot=270, name="deco:quarter_b")
    _shape(slide, MSO_SHAPE.PIE_WEDGE, x0, y0 + s, s, s, fill=col(3), rot=90, name="deco:quarter_c")
    _shape(slide, MSO_SHAPE.OVAL, x0 + s + 0.45, y0 + s + 0.45, s - 0.9, s - 0.9, fill=col(0), alpha=80,
           name="deco:circle_d")


_COMPOSE = {"circles": (_compose_circles, "left"), "bars": (_compose_bars, "left"),
            "quarter": (_compose_quarter, "right")}


@register("geometric_divider", kind="divider", extra_sizes=(STAT_BIG_PT,),
          summary="Section divider from circles / bars / quarter-circles with a big section number",
          source="G7 geometric primitives (Mike Kus conference-talk slides)")
def geometric_divider(deck, title, subtitle="", number=None, variant="circles"):
    """variant: 'circles' | 'bars' | 'quarter'. number: int or str shown at 54 pt in the theme stat color."""
    if variant not in _COMPOSE:
        raise StyleError(f"variant {variant!r} not in {sorted(_COMPOSE)}")
    if not str(title).strip():
        raise StyleError("title is empty")
    compose, side = _COMPOSE[variant]
    pal = _Pal(deck.theme)
    slide = deck.new_slide()
    vivid = pal.vivid
    compose(slide, lambda i: vivid[i % len(vivid)])

    text_x = 0.8 if side == "left" else 7.1
    text_w = 5.5
    num_txt = "" if number is None else (f"{number:02d}" if isinstance(number, int) else str(number))
    n_title = _wrap_n(deck, str(title), "title_main", text_w - 0.16, 3)
    n_sub = _wrap_n(deck, subtitle, "subtitle", text_w - 0.16, 3) if subtitle else 0
    parts = []
    if num_txt:
        parts.append(("num", 54 * LINE_H / 72 + 0.1))
        parts.append(("rule", 0.3))
    parts.append(("title", n_title * 30 * LINE_H / 72 + 0.1))
    if subtitle:
        parts.append(("sub", n_sub * 18 * LINE_H / 72 + 0.1 + 0.12))
    total = sum(h for _, h in parts)
    y = max(1.2, (7.0 - total) / 2)
    for kind, h in parts:
        if kind == "num":
            deck.add_textbox(slide, _iv(text_x), _iv(y), _iv(text_w), _iv(h), num_txt, role="stat_big")
        elif kind == "rule":
            _shape(slide, MSO_SHAPE.RECTANGLE, text_x + 0.08, y + 0.1, 1.1, 0.07, fill=pal.secondary,
                   name="deco:rule")
        elif kind == "title":
            deck.add_textbox(slide, _iv(text_x), _iv(y), _iv(text_w), _iv(h), str(title), role="title_main")
        else:
            deck.add_textbox(slide, _iv(text_x), _iv(y + 0.12), _iv(text_w), _iv(h - 0.12), subtitle,
                             role="subtitle")
        y += h
    num = deck._next_num()
    deck.add_textbox(slide, MARGIN, Inches(7.1), Inches(1.0), Inches(0.3), str(num), role="footer")
    return slide

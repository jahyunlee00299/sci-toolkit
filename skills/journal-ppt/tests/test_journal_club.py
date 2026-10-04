# lang-gate: skip-file  (test decks carry seeded notes problems on purpose)
"""Journal-club preset (Deck(prefs='journal_club')) and the QC-20 prose gate (qc_prose.py)."""
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import deck_builder  # noqa: E402
import qc_deck  # noqa: E402
import qc_prose  # noqa: E402
from deck_builder import Deck  # noqa: E402

NOTES = "【테스트】 이 슬라이드는 발표자가 읽는 한국어 노트이며 백 자 이상이어야 하므로 문장을 충분히 길게 이어서 쓰고, 같은 문장을 반복하지 않도록 서로 다른 내용을 한 번씩만 적습니다."


def club_deck(tmp_path, theme="nyt_hairline"):
    d = Deck(footer_right="Test et al., J. Test 2026", theme=theme, prefs="journal_club")
    s = d.title_slide("A Test Title", "Test et al. (2026)", "J. Test, 1(1):1-2", "Presenter", "2026-10-04")
    d.notes(s, NOTES)
    s = d.content_slide("Background")
    d.add_textbox(s, Inches(0.4), Inches(1.2), Inches(12), Inches(0.5), "A direct claim without a period.", role="subtitle")
    d.notes(s, NOTES)
    s = d.content_slide("Results")
    d.bullets(s, ["First point ends with a period.", "Second point as well."], role="body", left=Inches(0.4), top=Inches(1.5),
              width=Inches(12), height=Inches(3), bullet_char="")
    d.notes(s, NOTES)
    s = d.conclusions_slide("One takeaway line", [("Shows", "Text one.", False), ("Uncertain", "Text two.", True)])
    d.notes(s, NOTES)
    s = d.references_slide(["Test A (2026) A paper. J. Test."], highlight_index=0)
    d.notes(s, NOTES)
    s = d.closing_slide()
    d.notes(s, NOTES)
    s = d.appendix_slide("strengths and limitations")
    d.notes(s, NOTES)
    out = tmp_path / "club.pptx"
    d.save(str(out))
    return out


def test_prefs_loader():
    assert deck_builder.load_prefs("journal_club")["caption"]["pt"] == 12
    assert deck_builder.load_prefs(None) == {}
    with pytest.raises(deck_builder.StyleError):
        deck_builder.load_prefs("no_such_preset")


def test_structure_and_hidden_appendix(tmp_path):
    prs = Presentation(str(club_deck(tmp_path)))
    titles = [qc_prose.slide_title(s) for s in prs.slides]
    assert titles[-1].startswith("Appendix")
    assert prs.slides[-1]._element.get("show") == "0"          # hidden
    assert prs.slides[3]._element.get("show") is None            # conclusions stays visible
    assert any("Conclusions" in t for t in titles)


def test_trailing_periods_stripped_but_references_kept(tmp_path):
    prs = Presentation(str(club_deck(tmp_path)))
    for s in prs.slides:
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]
        if any(t.strip() == "References" for t in texts):
            assert any(t.rstrip().endswith("J. Test.") for t in texts)
            continue
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    assert not p.text.rstrip().endswith("."), p.text


def test_closing_slide_is_just_thank_you(tmp_path):
    prs = Presentation(str(club_deck(tmp_path)))
    closing = [s for s in prs.slides if any(sh.has_text_frame and sh.text_frame.text == "Thank you" for sh in s.shapes)][0]
    texts = [sh.text_frame.text for sh in closing.shapes if sh.has_text_frame]
    assert texts == ["Thank you", "for listening. Please ask me anything"]   # no kicker, no paper, no presenter
    big = [r.font.size.pt for sh in closing.shapes if sh.has_text_frame and sh.text_frame.text == "Thank you"
           for p in sh.text_frame.paragraphs for r in p.runs]
    assert big == [54]


def test_closing_54pt_is_declared_by_the_layout_in_any_theme(tmp_path):
    for theme in ("navy_lab", "nyt_hairline", "dark_seminar"):
        d = Deck(theme=theme, prefs="journal_club")
        s = d.content_slide("Background")
        d.notes(s, NOTES)
        s = d.closing_slide()
        d.notes(s, NOTES)
        out = tmp_path / f"{theme}.pptx"
        d.save(str(out))
        rep = qc_deck.qc_deck(str(out), language=False)
        assert not [c for c in rep.critical if "54" in c or "type scale" in c], (theme, rep.critical)


def test_plain_centered_caption_with_prefs(tmp_path):
    d = Deck(theme="nyt_hairline", prefs="journal_club")
    s = d.content_slide("Figure")
    box = d.add_figure(s, str(Path(__file__).parent / "_fig.png"), top=Inches(1.6), max_height=Inches(4)) \
        if (Path(__file__).parent / "_fig.png").exists() else (Inches(1), Inches(1.6), Inches(6), Inches(4))
    cap = d.caption(s, "Figure 1. Effect of a thing\n(a) time, (b) ratio", box)
    runs = [r for p in cap.text_frame.paragraphs for r in p.runs]
    assert all(r.font.size == Pt(12) and not r.font.italic and not r.font.bold for r in runs)
    assert all(p.alignment is not None for p in cap.text_frame.paragraphs)
    assert len(cap.text_frame.paragraphs) == 2


def test_default_caption_unchanged_without_prefs():
    d = Deck(theme="navy_lab")
    s = d.content_slide("Figure")
    cap = d.caption(s, "Figure 1. Effect of et al.", (Inches(1), Inches(1.6), Inches(6), Inches(4)))
    r = cap.text_frame.paragraphs[0].runs[0]
    assert r.font.size == Pt(9.5) and r.font.italic is True and r.font.bold is True


def test_club_deck_passes_qc_without_critical(tmp_path):
    rep = qc_deck.qc_deck(str(club_deck(tmp_path)), language=False)
    assert not rep.critical, rep.critical
    assert not [w for w in rep.warning if "QC-20" in w], [w for w in rep.warning if "QC-20" in w]


# ---- qc_prose rules on seeded problems -------------------------------------------------------

def seeded(tmp_path, notes, body="Plain body text", title="Results", extra_numbered=0):
    d = Deck(theme="nyt_hairline", prefs="journal_club")
    for t in ("Background", "Methods", "Conclusions", "Discussion", "Extra"):
        s = d.content_slide(t)
        d.notes(s, NOTES)
    s = d.content_slide(title)
    d.add_textbox(s, Inches(0.4), Inches(1.2), Inches(12), Inches(0.5), body, role="subtitle")
    if extra_numbered:
        d.bullets(s, [f"{i + 1}   Question number {i + 1}" for i in range(extra_numbered)], role="body",
                  left=Inches(0.4), top=Inches(2), width=Inches(12), height=Inches(4), bullet_char="")
    d.notes(s, notes)
    out = tmp_path / "seed.pptx"
    d.save(str(out))
    return Presentation(str(out))


def rules(prs):
    return {f.rule for f in qc_prose.check_prose(prs, prefs=deck_builder.load_prefs("journal_club"))}


def test_p1_notes_residue(tmp_path):
    prs = seeded(tmp_path, NOTES + " 앞선 초안에는 27로 적었지만 이번에 바로잡았습니다. 발표 전에 보완하겠습니다.")
    assert "P1.notes-residue" in rules(prs)


def test_p2_count_mismatch_and_match(tmp_path):
    bad = seeded(tmp_path, NOTES + " 다섯 가지 질문을 준비했습니다.", extra_numbered=6)
    assert "P2.notes-count" in rules(bad)
    ok = seeded(tmp_path, NOTES + " 여섯 가지 질문을 준비했습니다.", extra_numbered=6)
    assert "P2.notes-count" not in rules(ok)


def test_p3_duplicate_sentence(tmp_path):
    s = "이 문장은 같은 노트 안에서 두 번 반복되는 충분히 긴 한국어 문장입니다."
    prs = seeded(tmp_path, NOTES + " " + s + " 중간 문장입니다 그리고 다른 내용이 들어갑니다 충분히 길게. " + s)
    assert "P3.notes-dup" in rules(prs)


def test_p4_aiism_and_semicolon(tmp_path):
    prs = seeded(tmp_path, NOTES, body="We leverage a robust method; the result is seamless today")
    assert "P4.slide-aiism" in rules(prs)
    clean = seeded(tmp_path, NOTES, body="The method gives a 64% yield at 25 C")
    assert "P4.slide-aiism" not in rules(clean)


def test_p5_trailing_period_only_with_prefs(tmp_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tb.text_frame.paragraphs[0].add_run().text = "A sentence that ends with a period."
    assert "P5.trailing-period" in {f.rule for f in qc_prose.check_prose(prs, prefs=deck_builder.load_prefs("journal_club"))}
    assert "P5.trailing-period" not in {f.rule for f in qc_prose.check_prose(prs, prefs=None)}


def test_p6_no_conclusions_and_p7_visible_appendix(tmp_path):
    d = Deck(theme="nyt_hairline", prefs="journal_club")
    for t in ("Background", "Methods", "Results", "Discussion", "References"):
        s = d.content_slide(t)
        d.notes(s, NOTES)
    s = d.appendix_slide("extra", hidden=False)
    d.notes(s, NOTES)
    out = tmp_path / "x.pptx"
    d.save(str(out))
    r = rules(Presentation(str(out)))
    assert "P6.no-conclusions" in r and "P7.appendix-shown" in r

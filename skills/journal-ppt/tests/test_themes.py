"""Theme / contrast / closed-list tests for journal-ppt.

Run:  "$PY" -m pytest scientific-skills/active/journal-ppt/tests/test_themes.py -q
(or `"$PY" tests/test_themes.py [outdir]` to just write one deck per theme and print QC verdicts).
"""
from __future__ import annotations

import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))

import pytest                                   # noqa: E402
from PIL import Image, ImageDraw                # noqa: E402
from pptx import Presentation                   # noqa: E402

import deck_builder as db                       # noqa: E402
import qc_deck as qc                            # noqa: E402

KO = "발표자 노트입니다. 이 슬라이드의 핵심 메시지를 설명하고 질문에 답합니다. " * 3
# The registry IS the test matrix: every theme added to db.THEMES is picked up automatically
THEME_NAMES = sorted(db.THEMES)
BUILTIN = {"navy_lab", "assertion_evidence", "journal_print", "dark_seminar", "mono_one"}
STAT_BIG_THEMES = {"mono_one", "swiss_grid", "metric_first"}   # the only themes that list 54 pt in extra_sizes


def make_figure(path: Path, size=(1600, 900)) -> Path:
    """White-background chart-like PNG (the glare case dark_seminar's plate exists for)."""
    im = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(im)
    w, h = size
    d.line([(100, h - 100), (100, 80), (w - 80, h - 100)], fill="black", width=5)
    for i, col in enumerate(["#0072B2", "#D55E00", "#009E73"]):
        x = 220 + i * 420
        d.rectangle([x, h - 100 - (300 + 120 * i), x + 220, h - 100], fill=col)
    path.parent.mkdir(parents=True, exist_ok=True)
    im.save(path)
    return path


def build_deck(theme: str, out: Path, fig: Path) -> Path:
    d = db.Deck(footer_right="Test et al., J Test 2026", theme=theme)
    t = d.theme
    d.notes(d.title_slide("Cascade Design Enables High-Yield Synthesis",
                          "Test et al. (2026)", "J Test 1:1-2 -- DOI: 10.0000/test",
                          "J. Lee", "2026-10-04"), KO)
    if t.allow_bullets:
        s = d.content_slide("Background & Motivation")
        d.bullets(s, ["Enzyme cascades avoid intermediate isolation.",
                      "Cofactor regeneration limits productivity.",
                      "Single-pot yields remain below 50%.",
                      "Kinetic models guide enzyme selection.",
                      "Bayesian optimization narrows conditions."], role="body")
        d.notes(s, KO)
    d.notes(d.reactions_slide("Reaction scheme", [str(fig), str(fig)],
                              "Scheme 1. Two-step cascade. (Figure 1 from [1])"), KO)
    if t.allow_bullets:
        pts = ["Yield reached 85% in 3 h.", "Selectivity above 99%.", "Titer doubled versus control."]
        if t.caption_h > db.CAPTION_H:     # long-caption themes leave room for <= 2 key points
            pts = pts[:2]
        s = d.figure_slide("Key Result", "What this figure shows", pts,
                           str(fig), "Figure 2. Time course. (Figure 2 from [1])")
    else:
        s = d.figure_slide("Yield reached 85 percent within three hours", "", [],
                           str(fig), "Figure 2. Time course. (Figure 2 from [1])")
    d.notes(s, KO)
    d.notes(d.prior_work_slide("Comparison with prior work", str(fig),
                               "Figure 5. Prior-work comparison. (Figure 5 from [2])",
                               note="Same substrate, different enzymes"), KO)
    if 54 in t.extra_sizes:
        d.notes(d.stat_slide("Headline number", "85%", "isolated yield after three hours",
                             "versus 41% for the best single-enzyme route"), KO)
    d.notes(d.table_slide("Kinetic Parameters", ["Enzyme", "kcat (s-1)", "Km (mM)"],
                          [["A", "12", "0.4"], ["B", "7", "1.1"]]), KO)
    d.notes(d.references_slide(["Test et al. (2026) J Test 1, 1-2.",
                                "Other et al. (2020) J Other 2, 3-4."]), KO)
    return Path(d.save(str(out)))


@pytest.fixture(scope="module")
def fig(tmp_path_factory):
    return make_figure(tmp_path_factory.mktemp("figs") / "fig.png")


# --------------------------------------------------------------------------
# 1. every theme builds, passes qc with 0 CRITICAL, theme round-trips via keywords
# --------------------------------------------------------------------------
@pytest.mark.parametrize("theme", THEME_NAMES)
def test_theme_deck_zero_critical(theme, fig, tmp_path):
    p = build_deck(theme, tmp_path / f"{theme}.pptx", fig)
    prs = Presentation(str(p))
    assert prs.core_properties.keywords == f"journal-ppt-theme:{theme}"
    assert qc.detect_theme(prs) == theme
    rep = qc.qc_deck(str(p))
    assert rep.critical == [], rep.render(str(p))
    assert not [w for w in rep.warning if "contrast" in w or "not in the" in w], rep.warning
    assert any(m.startswith(f"Theme: {theme} (file keywords)") for m in rep.passed)
    assert any("S3: reactions slide present" in m for m in rep.passed)


def test_fallback_and_override(fig, tmp_path):
    p = build_deck("dark_seminar", tmp_path / "d.pptx", fig)
    # forcing the wrong palette must light up warnings (the palette check really uses the theme)
    rep = qc.qc_deck(str(p), theme_override="navy_lab")
    assert any("not in the navy_lab palette" in w for w in rep.warning)
    assert any("overrides the theme recorded" in w for w in rep.warning)
    # a file without the keyword falls back to navy_lab
    prs = Presentation(str(p))
    prs.core_properties.keywords = ""
    q = tmp_path / "nokw.pptx"
    prs.save(str(q))
    assert qc.detect_theme(Presentation(str(q))) is None
    assert any("default fallback" in m for m in qc.qc_deck(str(q)).passed)


def test_stat_big_only_in_declared_themes(fig, tmp_path):
    for name in THEME_NAMES:
        d = db.Deck(theme=name)
        if name in STAT_BIG_THEMES:
            d.stat_slide("t", "85%", "meaning")
        else:
            with pytest.raises(db.StyleError):
                d.stat_slide("t", "85%", "meaning")
    # a 54 pt run in a navy_lab deck is a CRITICAL
    p = build_deck("navy_lab", tmp_path / "n.pptx", fig)
    prs = Presentation(str(p))
    run = [r for sh in prs.slides[1].shapes if sh.has_text_frame
           for pa in sh.text_frame.paragraphs for r in pa.runs][0]
    run.font.size = db.Pt(54)
    prs.save(str(p))
    assert any("54" in c for c in qc.qc_deck(str(p)).critical)


@pytest.mark.parametrize("style", db.HEADER_STYLES)
def test_every_header_style_generic(style, fig, tmp_path, monkeypatch):
    """A registry entry with each supported header style builds and passes QC with no new
    code path (data-only extension)."""
    import dataclasses
    base = db.THEMES["navy_lab"]
    # titles sit on the plain background unless the style is 'bar' -> they need a dark color
    pal = dict(base.palette, title_bar_text=db._rgb("1A355E")) if style != "bar" else base.palette
    th = dataclasses.replace(base, name="probe_" + style, header_style=style, palette=pal,
                             rule_color=db._rgb("C8102E") if style == "rule" else None)
    monkeypatch.setitem(db.THEMES, th.name, th)
    p = build_deck(th.name, tmp_path / f"{style}.pptx", fig)
    rep = qc.qc_deck(str(p))
    assert rep.critical == [], rep.render(str(p))
    assert not [w for w in rep.warning if "contrast" in w], rep.warning
    assert all(ok for *_, ok in db.contrast_pairs(th.name))
    titles = [sh.text_frame.text for sh in Presentation(str(p)).slides[2].shapes if sh.has_text_frame]
    assert ("Reaction scheme" in titles) == (style != "none")


def test_font_override_is_checked(fig, tmp_path, monkeypatch):
    import dataclasses
    th = dataclasses.replace(db.THEMES["navy_lab"], name="probe_font", font="Calibri")
    monkeypatch.setitem(db.THEMES, th.name, th)
    p = build_deck(th.name, tmp_path / "f.pptx", fig)
    assert qc.qc_deck(str(p)).critical == []                    # theme font accepted
    assert any("Calibri" in c for c in qc.qc_deck(str(p), theme_override="navy_lab").critical)


def test_assertion_evidence_rules(fig):
    d = db.Deck(theme="assertion_evidence")
    with pytest.raises(db.StyleError):
        d.bullets(d.content_slide("A full sentence headline here"), ["x"])
    with pytest.raises(db.StyleError):
        d.figure_slide("A full sentence headline here", "", ["bullet"], str(fig), "c")
    with pytest.raises(db.StyleError):
        d.figure_slide("Too short", "", [], str(fig), "c")


# --------------------------------------------------------------------------
# 2. closed lists were not widened beyond what is documented
# --------------------------------------------------------------------------
def test_closed_lists_documented_only():
    assert db.ALLOWED_LINE_SPACING == {1.0, 1.15, 2.0} == qc.ALLOWED_LINE_SPACING
    assert db.ALLOWED_SPACE_AFTER_PT == {0, 2, 6, 12, 16} == qc.ALLOWED_SPACE_AFTER_PT
    base = {9, 9.5, 11, 12, 14, 16, 18, 26, 30}
    assert qc.ALLOWED_FONT_SIZES == base
    assert qc.ALLOWED_FONT_SIZES_RESEARCH == base | {20}
    assert db.ALLOWED_FONT_SIZES == base | {20, 54}
    assert {r["pt"] for r in db.ROLES.values()} <= db.ALLOWED_FONT_SIZES
    # the only theme-specific size is stat_big=54, only in the documented themes
    assert db.EXTRA_SIZES_DOC == {54}
    for name, th in db.THEMES.items():
        assert (name in STAT_BIG_THEMES) == (54 in th.extra_sizes), name
    # line-spacing / space_after of every role come from the closed lists
    assert set(db.LINE_SPACING.values()) <= db.ALLOWED_LINE_SPACING
    assert set(db.SPACE_AFTER.values()) <= db.ALLOWED_SPACE_AFTER_PT
    assert db.LINE_SPACING["body"] == db.LINE_SPACING["sidebar_bullet"] == db.LINE_SPACING["presenter"] == 2.0
    assert db.LINE_SPACING["reference"] == 1.15
    for r in ("title_bar", "title_main", "footer", "caption", "table_header", "table_cell"):
        assert db.LINE_SPACING[r] == 1.0
    assert db.ROLES["body"]["pt"] == 16 and db.ROLES["sidebar_bullet"]["pt"] == 14
    assert BUILTIN <= set(db.THEMES)
    for th in db.THEMES.values():
        assert th.extra_sizes <= {54}, th.name
        assert th.header_style in db.HEADER_STYLES


def test_old_spacing_rejected(fig, tmp_path):
    p = build_deck("navy_lab", tmp_path / "n.pptx", fig)
    prs = Presentation(str(p))
    sh = [s for s in prs.slides[1].shapes if s.has_text_frame and "Enzyme" in s.text_frame.text][0]
    sh.text_frame.paragraphs[0].line_spacing = 1.5
    prs.save(str(p))
    assert any("line_spacing=1.5" in c for c in qc.qc_deck(str(p)).critical)


# --------------------------------------------------------------------------
# 3. contrast: every (role color, theme background) pair
# --------------------------------------------------------------------------
@pytest.mark.parametrize("theme", THEME_NAMES)
def test_contrast_every_role_pair(theme):
    pairs = db.contrast_pairs(theme)
    assert len(pairs) >= len(db.ROLES)
    bad = [(r, fg, bg, round(x, 2), need) for r, fg, bg, x, need, ok in pairs if not ok]
    assert not bad, f"{theme}: below WCAG: {bad}"


def test_contrast_reference_values():
    # numbers quoted in the crawl / task statement
    W, O, G = (255, 255, 255), (0xE8, 0x6A, 0x1A), (0x88, 0x88, 0x88)
    assert round(db.contrast_ratio(O, W), 2) == 3.23
    assert round(db.contrast_ratio(G, W), 2) == 3.54
    assert db.contrast_ratio(db.CAPTION_COLOR, db.WHITE) >= 4.5
    # accent orange survives only as >=18pt bold
    for theme in ("navy_lab", "assertion_evidence"):
        for role, fg, bg, ratio, need, ok in db.contrast_pairs(theme):
            if fg == "E86A1A":
                assert db.ROLES[role]["pt"] >= 18 and db.ROLES[role]["bold"], role
    # an accent color that fails 4.5:1 on its background may only color >=18pt bold roles
    for theme in THEME_NAMES:
        th = db.THEMES[theme]
        for role, spec in db.ROLES.items():
            if (th.role_color(role) == th.palette["accent"]
                    and db.contrast_ratio(th.palette["accent"], th.bg) < 4.5):
                assert spec["pt"] >= 18 and spec["bold"], (theme, role)


def test_registry_metadata_and_list_themes():
    assert len(db.THEMES) >= 26 and BUILTIN <= set(db.THEMES)
    rows = db.list_themes()
    assert [r["name"] for r in rows] == list(db.THEMES)
    for r in rows:
        assert set(r) == {"name", "family", "provenance", "description", "header_style"}
        assert r["family"] and r["description"].strip() and r["header_style"] in db.HEADER_STYLES
    # every non-core theme documents what it did NOT implement (flat variant)
    for th in db.THEMES.values():
        if th.family != "core":
            assert th.flat_note.startswith("flat variant"), th.name
        assert th.font == "Arial", th.name
    # provenance is honest: core themes are our own, the research family is measured, nothing else claims "measured"
    for th in db.THEMES.values():
        assert th.provenance in db.PROVENANCE_KINDS, th.name
        assert (th.family == "core") == (th.provenance == "own"), th.name
        assert (th.family == "research") == (th.provenance == "measured"), th.name


def test_dark_or_tinted_background_gets_light_plate():
    """Figures with white backgrounds glare on dark grounds: any dark theme needs a plate (generic flag)."""
    for name, th in db.THEMES.items():
        if db.relative_luminance(th.bg) < 0.18:
            assert th.plate is not None, name
            assert db.relative_luminance(th.plate) > 0.5, name


def test_extra_sizes_outside_documented_set_rejected():
    import dataclasses
    with pytest.raises(ValueError):
        dataclasses.replace(db.THEMES["navy_lab"], extra_sizes=frozenset({60}))


def test_list_themes_cli():
    import subprocess
    out = subprocess.run([sys.executable, str(HERE.parent / "scripts" / "deck_builder.py"), "--list-themes"],
                         capture_output=True, text=True, encoding="utf-8", errors="replace")
    assert out.returncode == 0
    for name in db.THEMES:
        assert name in out.stdout


# --------------------------------------------------------------------------
# 4. vertical budget at 2.0 spacing
# --------------------------------------------------------------------------
def test_vertical_budget_navy(fig, tmp_path):
    d = db.Deck()
    # 3 one-line sidebar bullets fit in KEY_POINT_H, and the figure still gets >= 3 in
    h = d.key_points_height(["Yield reached 85% in 3 h.", "Selectivity above 99%.", "Titer doubled."])
    assert h <= db.KEY_POINT_H
    fig_top = db.KEY_POINT_TOP + h + db.Inches(0.05)
    assert db.Y_FLOOR - db.CAPTION_H - db.CAPTION_GAP - fig_top >= db.MIN_FIG_H
    # ~5 body bullets (one line each) end above Y_FLOOR; so do 9 lines
    line_in = 16 * db.LINE_HEIGHT_FACTOR * 2.0 / 72
    five = 5 * line_in + 4 * 6 / 72 + 0.1
    assert db.BODY_TOP_NO_FIGURE / 914400 + five < db.Y_FLOOR / 914400
    nine = 9 * line_in + 8 * 6 / 72 + 0.1
    assert db.BODY_TOP_NO_FIGURE / 914400 + nine <= db.Y_FLOOR / 914400 + 1e-9
    # a figure slide with a too-tall bullet block is refused, not silently overflowed
    with pytest.raises(db.StyleError):
        d.figure_slide("T", "s", ["word " * 60, "word " * 60, "word " * 60], str(fig), "c")


def test_journal_print_long_caption_geometry(fig, tmp_path):
    d = db.Deck(theme="journal_print")
    long_cap = "Figure 2. " + "Long descriptive caption sentence. " * 6
    d.notes(d.figure_slide("Result with a long caption", "", [], str(fig), long_cap), KO)
    p = d.save(str(tmp_path / "jp.pptx"))
    assert qc.qc_deck(p).critical == []


# --------------------------------------------------------------------------
# 5. S7: extract_figures.py (type==1 image-block crop, caption free, 300 DPI)
# --------------------------------------------------------------------------
def test_extract_figures_s7(tmp_path):
    fitz = pytest.importorskip("fitz")
    sys.path.insert(0, str(HERE.parent / "scripts"))
    import extract_figures as ef
    img = make_figure(tmp_path / "panel.png", (1200, 800))
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_text((60, 70), "Body text paragraph above the figure.", fontsize=10)
    page.insert_image(fitz.Rect(60, 100, 290, 253.3), filename=str(img))      # panel A
    page.insert_image(fitz.Rect(295, 100, 525, 253.3), filename=str(img))     # panel B (adjacent)
    page.insert_image(fitz.Rect(10, 780, 20, 790), filename=str(img))       # tiny logo -> skipped
    page.insert_text((60, 275), "Figure 3. Time course of the cascade.", fontsize=9)
    pdf = tmp_path / "paper.pdf"
    doc.save(str(pdf))
    res = ef.extract(pdf, tmp_path / "out", dpi=300)
    assert len(res) == 1, res                     # two panels merged, logo dropped
    r = res[0]
    assert r["label"] == "fig3"
    x0, y0, x1, y1 = r["bbox"]
    assert y1 <= 275 - 8                           # clip stops above the caption text block
    assert abs(r["width_px"] - (x1 - x0) * 300 / 72) <= 2
    assert abs(r["height_px"] - (y1 - y0) * 300 / 72) <= 2
    assert Image.open(r["path"]).size == (r["width_px"], r["height_px"])


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else Path.home() / "scratch" / "jppt_themes" / "decks")
    f = make_figure(out / "fig.png")
    for th in THEME_NAMES:
        p = build_deck(th, out / f"{th}.pptx", f)
        rep = qc.qc_deck(str(p))
        print(th, p, "CRIT", len(rep.critical), "WARN", len(rep.warning))

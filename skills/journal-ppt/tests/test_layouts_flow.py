"""Flow family layouts: pipeline_rail, walkthrough_highlight, numbered_grid_agenda, geometric_divider.

Run:  "$PY" -m pytest tests/test_layouts_flow.py -q
Every layout is built under EVERY theme in THEMES and must pass qc_deck and qc_layout with 0 CRITICAL.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from deck_builder import THEMES, Deck, StyleError  # noqa: E402
from layouts import LAYOUTS, load_all  # noqa: E402
from qc_deck import qc_deck  # noqa: E402
from qc_layout import CRITICAL, check_layout  # noqa: E402

KO = "발표자 노트입니다. 이 슬라이드의 핵심 메시지를 설명하고 질문에 답합니다. " * 3
FAMILY = ("pipeline_rail", "walkthrough_highlight", "numbered_grid_agenda", "geometric_divider")
MC_NS = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"

STEPS = [
    dict(label="Sucrose", sublabel="Substrate, 100 mM", duration="0 min", kind="substrate", note="pH 7.5, 30 °C"),
    dict(label="Sucrose phosphorylase", sublabel="SP, 0.5 mg/mL", duration="30 min", kind="enzyme",
         note="kcat 52 s⁻¹"),
    dict(label="Glucose 1-phosphate", sublabel="Intermediate", duration="10 min", kind="substrate"),
    dict(label="Glucose mutase PGM", sublabel="PGM, 0.2 mg/mL", duration="45 min", kind="enzyme",
         note="Mg²⁺ 5 mM"),
    dict(label="NAD⁺ / NADH", sublabel="Cofactor recycling", duration="2–3 h", kind="cofactor",
         note="LDH regenerates"),
    dict(label="Galactose isomerase", sublabel="L-AI, 0.3 mg/mL", duration="4–5 h", kind="enzyme",
         note="60 °C"),
    dict(label="Product P", sublabel="Product, 42 g/L", duration="6 h", kind="product", note="Yield 0.42 g/g"),
]
ITEMS = [("Background", "Why this product, why a cascade"), ("Design", "Enzyme selection and cofactor loop"),
         ("Kinetics", "kcat and Km fits, 5 mM–50 mM"), ("Optimisation", "Bayesian optimization, 30 °C–60 °C"),
         ("Scale-up", "1 L reactor"), ("TEA", "MPSP and E-factor"), ("Outlook", "Next steps"), ("Q&A", "")]


def _deck_all(theme: str) -> Deck:
    d = Deck(theme=theme)
    slides = []
    for n in range(3, 8):
        slides.append(d.layout("pipeline_rail", steps=STEPS[:n], title=f"Cascade, {n} steps",
                               cofactor_loop=(min(n - 1, 4), 1, "NADH recycling") if n >= 4 else None,
                               takeaway="Three enzymes convert sucrose to product P in one pot" if n % 2 else None))
    for act in (0, 3, 6):
        slides.append(d.layout("walkthrough_highlight", steps=STEPS, active=act, title="Step walkthrough",
                               cofactor_loop=(4, 1, "NADH recycling"), takeaway="Step 4 sets the flux",
                               morph=act == 3))
    for n, cur in ((4, None), (5, 1), (6, 2), (7, 0), (8, 7)):
        slides.append(d.layout("numbered_grid_agenda", items=ITEMS[:n], current=cur))
    for v in ("circles", "bars", "quarter"):
        slides.append(d.layout("geometric_divider", title="Bayesian optimization of the cascade",
                               subtitle="Closing the loop between kinetics and process design", number=3,
                               variant=v))
    for s in slides:
        d.notes(s, KO)
    return d


@pytest.mark.parametrize("theme", sorted(THEMES))
def test_family_passes_qc_under_every_theme(theme, tmp_path):
    p = tmp_path / f"flow_{theme}.pptx"
    _deck_all(theme).save(str(p))
    rep = qc_deck(str(p))                       # includes qc_layout
    assert not rep.critical, rep.critical[:5]
    pal = [w for w in rep.warning if "palette" in w]
    assert not pal, pal
    fs = check_layout(Presentation(str(p)), theme=theme)
    crit = [str(f) for f in fs if f.severity == CRITICAL]
    assert not crit, crit[:5]


def test_layouts_registered_with_source():
    load_all()
    for name in FAMILY:
        assert name in LAYOUTS and LAYOUTS[name].source.startswith("G"), name
    assert LAYOUTS["geometric_divider"].extra_sizes


# ------------------------------------------------------------------ negative tests
def _layout(name, **kw):
    return Deck().layout(name, **kw)


@pytest.mark.parametrize("bad", [[], STEPS[:2], STEPS + STEPS[:1], "nope", [{"sublabel": "x"}] * 3,
                                 [dict(label="a", kind="alien")] * 3])
def test_rail_rejects_bad_steps(bad):
    with pytest.raises(StyleError):
        _layout("pipeline_rail", steps=bad)


@pytest.mark.parametrize("loop", [(0, 0, "x"), (0, 9, "x"), (0, 1, ""), "ab", (0, 1)])
def test_rail_rejects_bad_loop(loop):
    with pytest.raises(StyleError):
        _layout("pipeline_rail", steps=STEPS[:4], cofactor_loop=loop)


def test_rail_text_that_cannot_fit_raises():
    long_word = [dict(label="Phosphoglucomutase", sublabel="x", kind="enzyme")] * 3 + [dict(label="b")] * 4
    with pytest.raises(StyleError):                       # a 7-step column is too narrow for this single word
        _layout("pipeline_rail", steps=long_word)
    wordy = [dict(label="Enzyme", sublabel="word " * 300, kind="enzyme")] * 5
    with pytest.raises(StyleError):                       # too tall for the card zone
        _layout("pipeline_rail", steps=wordy)
    with pytest.raises(StyleError):                       # takeaway wraps to two lines
        _layout("pipeline_rail", steps=STEPS[:4], takeaway="very long takeaway " * 8)
    with pytest.raises(StyleError):                       # pill text wider than its column
        _layout("pipeline_rail", steps=[dict(label="a", duration="a very long duration text")] * 7)


@pytest.mark.parametrize("active", [-1, 7, "1", None, True])
def test_walkthrough_rejects_bad_active(active):
    with pytest.raises(StyleError):
        _layout("walkthrough_highlight", steps=STEPS, active=active)


@pytest.mark.parametrize("items", [[], ITEMS[:3], ITEMS + [("Extra", "")], "x", [("", "no label")] * 4])
def test_grid_rejects_bad_items(items):
    with pytest.raises(StyleError):
        _layout("numbered_grid_agenda", items=items)


def test_grid_rejects_bad_current_and_overflow():
    with pytest.raises(StyleError):
        _layout("numbered_grid_agenda", items=ITEMS[:4], current=4)
    with pytest.raises(StyleError):
        _layout("numbered_grid_agenda", items=[("Label " * 12, "desc " * 40)] * 8)


def test_divider_rejects_bad_input():
    with pytest.raises(StyleError):
        _layout("geometric_divider", title="x", variant="triangles")
    with pytest.raises(StyleError):
        _layout("geometric_divider", title="  ")
    with pytest.raises(StyleError):
        _layout("geometric_divider", title="word " * 60)
    with pytest.raises(StyleError):
        _layout("geometric_divider", title="ok", subtitle="word " * 80)


# ------------------------------------------------------------------ structure tests
def test_rail_nodes_same_size_and_glued_arrows():
    s = _layout("pipeline_rail", steps=STEPS, cofactor_loop=(4, 1, "NADH recycling"))
    nodes = sorted([sh for sh in s.shapes if sh.name.startswith("!!pr_node_")], key=lambda x: x.left)
    assert len(nodes) == 7
    assert len({(n.width, n.height, n.top) for n in nodes}) == 1                 # consistent size, one rail line
    ids = {n.shape_id: i for i, n in enumerate(nodes)}
    arrows = [sh for sh in s.shapes if sh.name.startswith("!!pr_arrow_")]
    assert len(arrows) == 6
    for a in arrows:
        cnv = a._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
        st, en = cnv.find(qn("a:stCxn")), cnv.find(qn("a:endCxn"))
        src = ids[int(st.get("id"))]
        assert int(en.get("id")) == nodes[src + 1].shape_id                      # glued to consecutive nodes
        n_from = nodes[src]
        assert abs(a.begin_x - (n_from.left + n_from.width)) <= 2
        assert abs(a.begin_y - (n_from.top + n_from.height / 2)) <= 2
    assert any(sh.name == "!!pr_loop" for sh in s.shapes)


def test_rail_cards_never_collide():
    for n in range(3, 8):
        s = _layout("pipeline_rail", steps=STEPS[:n])
        cards = sorted([sh for sh in s.shapes if sh.name.startswith("!!pr_card_")], key=lambda x: x.left)
        assert len(cards) == n
        for a, b in zip(cards, cards[1:]):
            assert a.left + a.width <= b.left, (n, a.name)


def _names(slide):
    return {sh.name for sh in slide.shapes if sh.name.startswith("!!pr_")}


@pytest.mark.parametrize("active", [0, 2, 6])
def test_walkthrough_and_rail_share_shape_names_for_morph(active):
    d = Deck()
    a = d.layout("pipeline_rail", steps=STEPS, cofactor_loop=(4, 1, "NADH recycling"), takeaway="Same story")
    b = d.layout("walkthrough_highlight", steps=STEPS, active=active, cofactor_loop=(4, 1, "NADH recycling"),
                 takeaway="Same story", morph=True)
    na, nb = _names(a), _names(b)
    assert na == nb and len(na) >= 7 * 4
    for kind in ("node", "card", "text", "pill"):                                # the highlighted step, every part
        assert f"!!pr_{kind}_{active}" in na and f"!!pr_{kind}_{active}" in nb
    by_a = {sh.name: sh for sh in a.shapes}
    by_b = {sh.name: sh for sh in b.shapes}
    assert by_b[f"!!pr_node_{active}"].width > by_a[f"!!pr_node_{active}"].width     # really enlarged
    other = 1 if active != 1 else 3
    assert by_b[f"!!pr_node_{other}"].width <= by_a[f"!!pr_node_{other}"].width
    # rail center line is identical on both slides (no vertical jump when morphing)
    ca = by_a["!!pr_node_0"].top + by_a["!!pr_node_0"].height / 2
    cb = by_b["!!pr_node_0"].top + by_b["!!pr_node_0"].height / 2
    assert ca == pytest.approx(cb, abs=2)
    assert b._element.find(MC_NS) is not None and a._element.find(MC_NS) is None  # Morph on b only


def test_dimmed_cards_are_label_only_and_active_card_has_details():
    s = _layout("walkthrough_highlight", steps=STEPS, active=3)
    txt = {sh.name: sh.text_frame.text for sh in s.shapes if sh.name.startswith("!!pr_text_")}
    assert "PGM, 0.2 mg/mL" in txt["!!pr_text_3"] and "ENZYME" in txt["!!pr_text_3"]
    assert txt["!!pr_text_0"] == "Sucrose"


def test_divider_variants_are_visibly_different():
    sets = {}
    slide_area = 13.333 * 7.5 * 914400 * 914400
    for v in ("circles", "bars", "quarter"):
        s = _layout("geometric_divider", title="Section", subtitle="Sub", number=2, variant=v)
        deco = [sh for sh in s.shapes if sh.name.startswith("deco:") and sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
        sets[v] = frozenset(sh.name for sh in deco)
        area = sum(sh.width * sh.height for sh in deco if "rule" not in sh.name)
        assert area / slide_area > 0.30, (v, area / slide_area)                   # shapes occupy >30 % of the slide
    assert len(set(sets.values())) == 3


def test_grid_current_item_is_highlighted():
    s = _layout("numbered_grid_agenda", items=ITEMS[:6], current=2)
    assert [sh.name for sh in s.shapes].count("deco:current_panel") == 1
    s0 = _layout("numbered_grid_agenda", items=ITEMS[:6])
    assert "deco:current_panel" not in [sh.name for sh in s0.shapes]


def test_text_stays_in_closed_scale_and_font():
    d = _deck_all("navy_lab")
    allowed = {9, 9.5, 11, 12, 14, 16, 18, 26, 30, 54}
    for s in d.prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        assert r.font.size.pt in allowed, (sh.name, r.font.size.pt)
                        assert r.font.name == "Arial"

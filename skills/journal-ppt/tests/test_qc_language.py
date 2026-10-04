# lang-gate: skip-file  (test decks carry seeded errors on purpose)
"""QC-19: qc_deck folds the language-gate skill into the deck report."""
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import qc_deck  # noqa: E402

KANA = "".join(map(chr, (0x3053, 0x308C, 0x306F)))


def make_deck(tmp_path, body, notes, name="d.pptx"):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
    tb.text_frame.word_wrap = True
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = body
    run.font.size = Pt(16)
    run.font.name = "Arial"
    s.notes_slide.notes_text_frame.text = notes
    p = tmp_path / name
    prs.save(p)
    return str(p)


def qc19(report):
    return ([m for m in report.critical if "QC-19" in m],
            [m for m in report.warning if "QC-19" in m or "language-gate" in m],
            [m for m in report.passed if "QC-19" in m])


def test_clean_deck_passes_qc19(tmp_path):
    p = make_deck(tmp_path, "Enzyme kinetics were measured at 37 °C.",
                  "이 슬라이드에서는 효소 반응 속도를 37도에서 측정한 결과를 설명합니다. " * 3)
    crit, warn, ok = qc19(qc_deck.qc_deck(p, layout=False))
    assert crit == [] and warn == [] and ok


def test_foreign_script_in_body_is_critical(tmp_path):
    p = make_deck(tmp_path, f"Result {KANA} shown", "정상적인 한국어 노트입니다. " * 8)
    crit, _, _ = qc19(qc_deck.qc_deck(p, layout=False))
    assert len(crit) == 1 and "foreign-script" in crit[0] and "body" in crit[0]


def test_foreign_script_in_notes_is_critical(tmp_path):
    p = make_deck(tmp_path, "Clean body text.", f"노트에 외국 문자 {KANA} 가 섞였다. " * 4)
    crit, _, _ = qc19(qc_deck.qc_deck(p, layout=False))
    assert crit and all("notes" in c for c in crit)


def test_english_spelling_in_body_is_warning(tmp_path):
    p = make_deck(tmp_path, "The colour of the seperate fraction", "노트입니다. " * 20)
    crit, warn, _ = qc19(qc_deck.qc_deck(p, layout=False))
    assert crit == []
    assert any("en-british" in w and "'colour'" in w for w in warn)


def test_korean_spelling_in_notes_is_warning(tmp_path):
    p = make_deck(tmp_path, "Clean body.", "반응이 시작되요. 몇일 뒤 다시 측정할수 있다. " * 3)
    crit, warn, _ = qc19(qc_deck.qc_deck(p, layout=False))
    assert crit == []
    joined = "\n".join(warn)
    assert "ko-doe-dwae" in joined and "ko-myeochil" in joined and "ko-spacing" in joined


def test_not_run_warning_when_gate_missing(tmp_path, monkeypatch):
    p = make_deck(tmp_path, "Clean body.", "노트입니다. " * 20)
    monkeypatch.setattr(qc_deck, "_load_language_gate", lambda: None)
    _, warn, ok = qc19(qc_deck.qc_deck(p, layout=False))
    assert any("language-gate NOT RUN" in w for w in warn) and ok == []


def test_gate_path_points_at_the_skill():
    assert (qc_deck.LANGUAGE_GATE_SCRIPTS / "langgate" / "__init__.py").exists()

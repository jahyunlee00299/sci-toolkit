"""Deck.morph writes a Morph transition that survives save/reload and does not upset the QC gate."""
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from deck_builder import Deck, StyleError  # noqa: E402
from qc_deck import qc_deck  # noqa: E402

KO = "한국어 발표 노트입니다. " * 12


def _deck(tmp_path):
    d = Deck(footer_right="Test et al.")
    d.title_slide("T", "c", "j", "p", "2026-10-04")
    d.notes(d.prs.slides[0], KO)
    slides = []
    for x in (1.0, 6.0):
        s = d.content_slide("Pathway")
        box = d.add_textbox(s, Inches(x), Inches(2.5), Inches(3), Inches(1), "Enzyme 1", role="body")
        d.name_shape(box, "Enzyme1")
        d.notes(s, KO)
        slides.append(s)
    d.morph(slides[1])
    path = tmp_path / "m.pptx"
    d.save(str(path))
    return path


def test_morph_xml_roundtrip(tmp_path):
    path = _deck(tmp_path)
    xml = Presentation(str(path)).slides[2]._element.xml
    assert "p159:morph" in xml and 'option="byObject"' in xml
    assert "<p:fade" in xml                      # fallback for viewers without Morph
    assert "p159:morph" not in Presentation(str(path)).slides[1]._element.xml


def test_shape_names_match_across_slides(tmp_path):
    prs = Presentation(str(_deck(tmp_path)))
    names = [[sh.name for sh in prs.slides[i].shapes if sh.name.startswith("!!")] for i in (1, 2)]
    assert names[0] == names[1] == ["!!Enzyme1"]


def test_morph_deck_passes_qc(tmp_path):
    r = qc_deck(str(_deck(tmp_path)))
    assert not r.critical, r.critical


def test_morph_rejects_bad_args(tmp_path):
    d = Deck()
    s = d.content_slide("x")
    with pytest.raises(StyleError):
        d.morph(s, option="spin")
    with pytest.raises(StyleError):
        d.morph(s, dur_ms=10)


def test_morph_is_idempotent(tmp_path):
    d = Deck()
    s = d.content_slide("x")
    d.morph(s); d.morph(s, dur_ms=800)
    assert s._element.xml.count("p159:morph") == 1

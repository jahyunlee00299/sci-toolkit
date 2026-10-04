"""Frame family (slate_header, block_boxes, statement_slide, corner_wedge_title, band_title,
split_panel_figure): every layout under EVERY theme -> qc_deck 0 CRITICAL, qc_layout 0 CRITICAL; negatives."""
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import deck_builder as db  # noqa: E402
from layouts import LAYOUTS, load_all  # noqa: E402
from qc_deck import qc_deck  # noqa: E402
from qc_layout import check_layout  # noqa: E402

load_all()
KO = "발표자 노트입니다. 이 슬라이드의 핵심 메시지를 설명하고 질문에 답합니다. " * 3
FAMILY = ["slate_header", "block_boxes", "statement_slide", "corner_wedge_title", "band_title", "split_panel_figure"]
THEMES = sorted(db.THEMES)
LONG = "word " * 80


@pytest.fixture(scope="module")
def fig(tmp_path_factory):
    p = tmp_path_factory.mktemp("f") / "fig.png"
    im = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(im)
    d.line([(100, 800), (100, 80), (1520, 800)], fill="black", width=5)
    for i, c in enumerate(["#0072B2", "#D55E00", "#009E73"]):
        d.rectangle([220 + i * 420, 800 - 300 - 120 * i, 440 + i * 420, 800], fill=c)
    im.save(p)
    return str(p)


def build(theme, fig):
    d = db.Deck(footer_right="Lee et al., 2026", theme=theme)

    def n(s):
        d.notes(s, KO)

    bullets_ok = d.theme.allow_bullets
    n(d.layout("corner_wedge_title", title="Bayesian Optimization of a Three-Enzyme Cascade", subtitle="Journal club",
               presenter="J. Lee", date="2026-10-04", venue="Lab Meeting"))
    n(d.layout("band_title", title="Cofactor Regeneration Limits Productivity", subtitle="NADPH recycling kinetics",
               presenter="J. Lee", date="2026-10-04", venue="Seminar", hero_fig=fig))
    n(d.layout("band_title", title="Cofactor Regeneration Limits Productivity", presenter="J. Lee", date="2026-10-04"))
    if bullets_ok:
        n(d.layout("slate_header", title="Background: why cascades", section="Introduction", progress=(2, 8),
                   bullets=["Enzyme cascades avoid intermediate isolation.",
                            "Cofactor recycling limits productivity at 30 °C.",
                            "Single-pot yields remain below 50%."]))
    n(d.layout("slate_header", title="Time course at 5 mM substrate", section="Results", progress=(5, 8),
               fig_path=fig, caption="Figure 2. Titer over 3 h at pH 7.5.", source="Source: Test et al., J Test 2026"))
    blocks = [dict(kind="definition", title="Turnover number",
                   text="kcat is the maximal number of substrate molecules converted per active site per second."),
              dict(kind="result", title="Main result", text="Yield reached 85% within 3 h at 30 °C."),
              dict(kind="alert", title="Limitation", text="Only one substrate concentration (5 mM) was tested."),
              dict(kind="question", title="Discussion", text="Would the cascade tolerate 50 mM substrate?")]
    n(d.layout("block_boxes", title="Definitions and claims", blocks=blocks,
               sections=["Intro", "Methods", "Results", "Discussion"], current=2))
    n(d.layout("block_boxes", title="Rate law and example", grid="1x", blocks=[
        dict(kind="definition", title="Michaelis-Menten", text="v = Vmax [S] / (Km + [S])"),
        dict(kind="example", title="Worked example", text="At Km = 0.4 mM and [S] = 0.4 mM the rate is half of Vmax.")]))
    n(d.layout("block_boxes", title="Three blocks in two columns", grid="2x", blocks=blocks[:3]))
    n(d.layout("statement_slide", statement="A cofactor cycle, not the enzyme, sets the ceiling.",
               emphasis="cofactor cycle", attribution="Take-home message"))
    n(d.layout("statement_slide", statement="Nothing here."))
    if bullets_ok:
        n(d.layout("split_panel_figure", title="Titer rises with cofactor recycling", fig_path=fig,
                   caption="Figure 3. Titer versus recycling rate.", source="Source: Test et al., J Test 2026",
                   takeaways=["Recycling above 2 mM/min lifts titer.", "Gain saturates near 80% yield.",
                              "Km of the recycling enzyme is limiting."]))
        n(d.layout("split_panel_figure", title="Mirrored variant", fig_path=fig,
                   caption="Figure 3. Titer versus recycling rate.", source="Source: Test et al., J Test 2026",
                   takeaways=["Recycling lifts titer.", "Gain saturates."], fig_side="right"))
    return d


def test_registered():
    assert set(FAMILY) <= set(LAYOUTS)
    assert LAYOUTS["statement_slide"].extra_sizes == frozenset({db.STAT_BIG_PT})


@pytest.mark.parametrize("theme", THEMES)
def test_every_layout_every_theme(theme, fig, tmp_path):
    d = build(theme, fig)
    p = tmp_path / f"{theme}.pptx"
    d.save(str(p))
    rep = qc_deck(str(p))
    assert rep.critical == [], rep.critical
    crit = [f for f in check_layout(Presentation(str(p)), theme=theme) if f.severity == "CRITICAL"]
    assert crit == [], [str(c) for c in crit]
    assert not [w for w in rep.warning if "not in the" in w], rep.warning
    assert not [w for w in rep.warning if "contrast" in w and "L3" not in w], rep.warning
    for s in Presentation(str(p)).slides:     # no layout of the family reuses navy_lab's header bar
        assert "header_bar" not in {x.name for x in s.shapes}


def test_forbidden_bullets_theme(fig):
    d = db.Deck(theme="assertion_evidence")
    with pytest.raises(db.StyleError):
        d.layout("slate_header", title="T", section="S", progress=(1, 2), bullets=["a"])
    with pytest.raises(db.StyleError):
        d.layout("split_panel_figure", title="T", fig_path=fig, caption="c", source="s", takeaways=["a", "b"])


def test_hero_crop_is_not_stretched(fig):
    d = db.Deck()
    s = d.layout("band_title", title="T", presenter="P", date="2026", hero_fig=fig, hero_focus=0.3)
    pic = [x for x in s.shapes if x.name.startswith("hero:")][0]
    shown_w = (1 - pic.crop_left - pic.crop_right) * 1600
    shown_h = (1 - pic.crop_top - pic.crop_bottom) * 900
    assert abs(shown_w / shown_h - pic.width / pic.height) < 0.01
    assert pic.crop_top < pic.crop_bottom      # focus 0.3 keeps the upper part


def test_split_panel_keeps_caption_source_and_figure(fig):
    d = db.Deck()
    s = d.layout("split_panel_figure", title="T", fig_path=fig, caption="Figure 1. X.",
                 source="Source: A et al.", takeaways=["a b", "c d"])
    pic = [x for x in s.shapes if x.shape_type == 13][0]
    assert max(pic.height, pic.width) / 914400 >= 3.0 and pic.height / 914400 >= 3.0
    texts = " ".join(x.text_frame.text for x in s.shapes if x.has_text_frame)
    assert "Figure 1. X." in texts and "Source: A et al." in texts
    for bad in ("", "  "):
        with pytest.raises(db.StyleError):
            db.Deck().layout("split_panel_figure", title="T", fig_path=fig, caption="c", source=bad,
                             takeaways=["a", "b"])


def test_negatives(fig):
    def D():
        return db.Deck()

    cases = [
        ("slate_header", dict(title="T", section="S", progress=(1, 2), bullets=[])),
        ("slate_header", dict(title="T", section="S", progress=(1, 2), bullets=[LONG] * 5)),
        ("slate_header", dict(title="T", section="S", progress=(3, 2), bullets=["a"])),
        ("slate_header", dict(title="word " * 30, section="S", progress=(1, 2), bullets=["a"])),
        ("slate_header", dict(title="T", section="S", progress=(1, 2))),
        ("block_boxes", dict(title="T", blocks=[dict(kind="definition", title="a", text="b")])),
        ("block_boxes", dict(title="T", blocks=[dict(kind="nope", title="a", text="b")] * 2)),
        ("block_boxes", dict(title="T", blocks=[dict(kind="result", title="a", items=[])] * 2)),
        ("block_boxes", dict(title="T", blocks=[dict(kind="result", title="a", text=LONG)] * 4)),
        ("statement_slide", dict(statement=LONG)),
        ("statement_slide", dict(statement="Short one.", emphasis="absent")),
        ("corner_wedge_title", dict(title=LONG, presenter="P", date="d")),
        ("band_title", dict(title="T", presenter="", date="d")),
        ("split_panel_figure", dict(title="T", fig_path=fig, caption="c", source="s", takeaways=["only one"])),
        ("split_panel_figure", dict(title="T", fig_path=fig, caption="c", source="s", takeaways=[LONG] * 3)),
    ]
    for name, kw in cases:
        with pytest.raises(db.StyleError):
            D().layout(name, **kw)

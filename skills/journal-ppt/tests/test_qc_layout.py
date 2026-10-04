"""Tests for scripts/qc_layout.py: one passing and one failing deck per rule."""
import io
import sys
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import qc_layout as Q  # noqa: E402

LONG = "This sentence is deliberately long so that it wraps over many lines in a narrow box. " * 6


def new_prs():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def tb(sl, text, x, y, w, h, pt=16, color="333333", bold=False, name=None):
    s = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE      # as deck_builder does (noAutofit)
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.name = "Arial"
    r.font.color.rgb = RGBColor.from_string(color)
    if name:
        s.name = name
    return s


def card(sl, x, y, w, h, fill="DBE8F7"):
    s = sl.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.fill.background()
    return s


def png(color=(40, 40, 40)):
    b = io.BytesIO()
    Image.new("RGB", (400, 300), color).save(b, "PNG")
    b.seek(0)
    return b


def rules(findings, sev=None):
    return {f.rule for f in findings if sev is None or f.severity == sev}


# ------------------------------------------------------------------ helpers
def test_contrast_known_facts():
    assert Q.contrast_ratio("E86A1A", "FFFFFF") == pytest.approx(3.23, abs=0.01)
    assert Q.contrast_ratio("888888", "FFFFFF") == pytest.approx(3.54, abs=0.01)
    assert Q.contrast_ratio("888888", "F4F7FB") == pytest.approx(3.30, abs=0.01)
    assert Q.contrast_ratio("000000", "FFFFFF") == pytest.approx(21.0, abs=0.01)


# ----------------------------------------------------------------------- L1
def test_l1_pass():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "Short bullet", 0.5, 1.5, 6, 1.0)
    assert not [f for f in Q.check_layout(prs) if f.rule.startswith("L1")]


def test_l1_overflow_into_shape_below_is_critical():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, LONG, 0.5, 1.5, 6, 1.0)             # declared 1in, needs far more
    card(sl, 0.5, 3.0, 6, 1.0)                  # card the text runs into
    assert "L1.invade" in rules(Q.check_layout(prs), Q.CRITICAL)


def test_l1_floor_is_critical():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, LONG, 0.5, 5.0, 6, 1.9)
    assert "L1.floor" in rules(Q.check_layout(prs), Q.CRITICAL)


def test_l1_own_box_only_is_warning():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, LONG, 0.5, 1.5, 6, 0.5)
    fs = Q.check_layout(prs)
    assert "L1.box" in rules(fs, Q.WARNING)
    assert not rules(fs, Q.CRITICAL) & {"L1.floor", "L1.invade"}


def test_l1_table_cell_growth_crosses_floor():
    prs = new_prs()
    sl = slide(prs)
    t = sl.shapes.add_table(3, 2, Inches(0.5), Inches(5.0), Inches(4), Inches(1.2)).table
    t.cell(1, 0).text = LONG
    for r in (0, 1, 2):
        for c in (0, 1):
            for p in t.cell(r, c).text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(12)
    assert "L1.floor" in rules(Q.check_layout(prs), Q.CRITICAL)


def test_l1_hangul_is_measured_with_a_cjk_font():
    w_ko = Q.text_width_pt("한국어 발표 노트입니다", "arial", False, 16)
    assert w_ko > 16 * 5    # not zero-width notdef glyphs


# ----------------------------------------------------------------------- L2
def test_l2_pass_side_by_side_and_card():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "Left text", 0.5, 1.5, 5, 1)
    sl.shapes.add_picture(png(), Inches(7), Inches(1.5), Inches(4), Inches(3))
    card(sl, 0.5, 3.0, 5, 2.0)
    tb(sl, "Inside card", 0.7, 3.2, 4.5, 0.6)     # intentional text-in-card
    assert not [f for f in Q.check_layout(prs) if f.rule == "L2.overlap"]


def test_l2_picture_over_text_is_critical():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "Some body text that sits under the picture", 0.5, 1.5, 6, 1)
    sl.shapes.add_picture(png(), Inches(0.5), Inches(1.5), Inches(4), Inches(3))
    assert "L2.overlap" in rules(Q.check_layout(prs), Q.CRITICAL)


def test_l2_logo_is_exempt():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "Some body text", 0.5, 1.5, 6, 1)
    p = sl.shapes.add_picture(png(), Inches(0.5), Inches(1.5), Inches(1), Inches(1))
    p.name = "logo: lab"
    assert not [f for f in Q.check_layout(prs) if f.rule == "L2.overlap"]


def test_l2_small_overlap_is_warning():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "A" * 30, 0.5, 1.5, 6, 1, pt=16)
    pic = sl.shapes.add_picture(png(), Inches(0.5), Inches(1.5), Inches(6), Inches(3))
    pic.top = Inches(1.5 + 0.31)        # only a sliver of the text line is covered
    fs = [f for f in Q.check_layout(prs) if f.rule == "L2.overlap"]
    assert fs and all(f.severity == Q.WARNING for f in fs)


# ----------------------------------------------------------------------- L3
def test_l3_pass_dark_on_white():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "Readable", 0.5, 1.5, 6, 1, color="333333")
    assert not [f for f in Q.check_layout(prs) if f.rule.startswith("L3")]


def test_l3_low_contrast_normal_text_is_critical():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "orange normal", 0.5, 1.5, 6, 1, pt=16, color="E86A1A")      # 3.23 < 4.5
    assert "L3.contrast" in rules(Q.check_layout(prs), Q.CRITICAL)


def test_l3_large_text_between_3_and_4_5_is_warning():
    prs = new_prs()
    sl = slide(prs)
    tb(sl, "orange large", 0.5, 1.5, 6, 1, pt=20, color="E86A1A")       # 3.23, large
    fs = [f for f in Q.check_layout(prs) if f.rule == "L3.contrast"]
    assert fs and all(f.severity == Q.WARNING for f in fs)


def test_l3_uses_card_fill_as_background():
    prs = new_prs()
    sl = slide(prs)
    card(sl, 0.5, 1.5, 6, 1.5, fill="1A355E")
    tb(sl, "white on navy", 0.6, 1.6, 5, 0.6, color="FFFFFF")           # fine on navy
    tb(sl, "navy on navy", 0.6, 2.3, 5, 0.6, color="2E5E9B")            # poor on navy
    fs = [f for f in Q.check_layout(prs) if f.rule == "L3.contrast"]
    assert len(fs) == 1 and "navy on navy" in fs[0].message


def test_l3_over_picture_is_unverifiable_warning():
    prs = new_prs()
    sl = slide(prs)
    sl.shapes.add_picture(png(), Inches(0.5), Inches(1.5), Inches(6), Inches(3))
    tb(sl, "label on figure", 1, 2, 3, 0.5, color="FFFFFF")
    fs = Q.check_layout(prs)
    assert "L3.unverifiable" in rules(fs, Q.WARNING)
    assert "L3.contrast" not in rules(fs)


def test_l3_solid_dark_slide_background():
    prs = new_prs()
    sl = slide(prs)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = RGBColor(0x1A, 0x35, 0x5E)
    tb(sl, "dark text on dark bg", 0.5, 1.5, 6, 1, color="333333")
    assert "L3.contrast" in rules(Q.check_layout(prs), Q.CRITICAL)


# ----------------------------------------------------------------------- L4
def test_l4_pass():
    prs = new_prs()
    tb(slide(prs), "Real content", 0.5, 1.5, 6, 1)
    assert not [f for f in Q.check_layout(prs) if f.rule.startswith("L4")]


def test_l4_hidden_blank_placeholder():
    prs = new_prs()
    s1 = slide(prs)
    tb(s1, "Lorem ipsum dolor", 0.5, 1.5, 6, 1)
    s1._element.set("show", "0")
    slide(prs)                                  # blank
    s3 = slide(prs)
    tb(s3, "Result: TBD", 0.5, 1.5, 6, 1)
    fs = Q.check_layout(prs)
    assert {"L4.hidden", "L4.blank", "L4.placeholder"} <= rules(fs)
    assert sum(f.rule == "L4.placeholder" for f in fs) == 2
    assert next(f for f in fs if f.rule == "L4.blank").severity == Q.CRITICAL


# ------------------------------------------------------------------ L2 text
def test_text_over_text_is_critical_but_adjacent_is_not():
    prs = new_prs()
    s = slide(prs)
    tb(s, "First block of words on the slide", 0.5, 1.5, 6, 1)
    tb(s, "Second block printed right on top", 0.5, 1.6, 6, 1)
    assert "L2.overlap" in rules(Q.check_layout(prs), Q.CRITICAL)
    ok = new_prs()
    s = slide(ok)
    tb(s, "First block", 0.5, 1.5, 6, 0.8)
    tb(s, "Second block below", 0.5, 2.5, 6, 0.8)
    assert not [f for f in Q.check_layout(ok) if f.rule == "L2.overlap"]


# ----------------------------------------------------------------------- L5
def _pic(prs_slide, w, h, crop=None):
    pic = prs_slide.shapes.add_picture(png(), Inches(1), Inches(1.5), Inches(w), Inches(h))
    if crop:
        pic.crop_left = crop
    return pic


def test_l5_pass_native_and_cropped():
    prs = new_prs()
    s = slide(prs)
    _pic(s, 4.0, 3.0)                           # 400x300 px -> 4:3
    s2 = slide(prs)
    _pic(s2, 3.0, 3.0, crop=0.25)               # crops the visible width to 300 px -> 1:1
    assert not [f for f in Q.check_layout(prs) if f.rule == "L5.aspect"]


def test_l5_stretched_is_critical_and_slight_is_warning():
    prs = new_prs()
    _pic(slide(prs), 2.0, 3.0)                  # 4:3 image shown at 2:3
    fs = [f for f in Q.check_layout(prs) if f.rule == "L5.aspect"]
    assert fs and fs[0].severity == Q.CRITICAL
    slight = new_prs()
    _pic(slide(slight), 4.0, 2.9)               # 3.4% off
    fs = [f for f in Q.check_layout(slight) if f.rule == "L5.aspect"]
    assert fs and fs[0].severity == Q.WARNING


# ----------------------------------------------------------------------- L6
def _deck_with_titles(*titles):
    prs = new_prs()
    for ttl in titles:
        tb(slide(prs), ttl, 0.5, 0.4, 8, 0.8, pt=28)
    return prs


def test_l6_references_last_passes_and_allowed_closers():
    assert not [f for f in Q.check_layout(_deck_with_titles("Intro", "Result", "References"))
                if f.rule == "L6.refs_order"]
    assert not [f for f in Q.check_layout(_deck_with_titles("Intro", "References", "Questions", "Backup: kinetics"))
                if f.rule == "L6.refs_order"]
    assert not [f for f in Q.check_layout(_deck_with_titles("Intro", "References", "References (cont.)"))
                if f.rule == "L6.refs_order"]


def test_l6_references_followed_by_content_is_critical():
    fs = Q.check_layout(_deck_with_titles("References", "Result", "Discussion"))
    f = [x for x in fs if x.rule == "L6.refs_order"]
    assert len(f) == 1 and f[0].severity == Q.CRITICAL and f[0].slide == 1


def test_cli_exit_code(tmp_path):
    good, bad = new_prs(), new_prs()
    tb(slide(good), "Fine", 0.5, 1.5, 6, 1)
    slide(bad)
    g, b = tmp_path / "g.pptx", tmp_path / "b.pptx"
    good.save(g)
    bad.save(b)
    assert Q.main([str(g)]) == 0
    assert Q.main([str(b)]) == 1

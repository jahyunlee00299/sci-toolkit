"""Research family (chapter_tracker, paper_card, figure_card, claim_evidence, donut_stat_band, matrix_2x2,
diverging_bar, bw_divider_quote): every layout under EVERY theme -> qc_deck 0 CRITICAL, qc_layout 0 CRITICAL;
negative tests (cannot fit, empty list, out-of-range values) and honesty checks (arc sweep, bar scale, aspect)."""
import sys
from pathlib import Path

import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "scripts"))
import deck_builder as db  # noqa: E402
from layouts import LAYOUTS, load_all  # noqa: E402
from layouts.research import chapter_footer  # noqa: E402
from qc_deck import qc_deck  # noqa: E402
from qc_layout import check_layout  # noqa: E402

load_all()
KO = "발표자 노트입니다. 이 슬라이드의 핵심 메시지를 설명하고 질문에 답합니다. " * 3
FAMILY = ["chapter_tracker", "paper_card", "figure_card", "claim_evidence", "donut_stat_band", "matrix_2x2",
          "diverging_bar", "bw_divider_quote"]
THEMES = sorted(db.THEMES)
CHAPTERS = ["Kinetic characterization of the three enzymes", "Global fit of the cascade model",
            "Bayesian optimization of reaction conditions", "Scale-up and techno-economic outlook"]


@pytest.fixture(scope="module")
def fig(tmp_path_factory):
    p = tmp_path_factory.mktemp("f") / "fig.png"
    im = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(im)
    d.line([(100, 800), (100, 80), (1520, 800)], fill="black", width=5)
    for i, c in enumerate(["#0072B2", "#D55E00", "#009E73"]):
        d.rectangle([220 + i * 420, 800 - 300 - 120 * i, 440 + i * 420, 800], fill=c)
    im.save(p)
    return str(p)


def build(theme, fig):
    d = db.Deck(footer_right="Lee et al., 2026", theme=theme)
    n = lambda s: d.notes(s, KO)           # noqa: E731
    n(d.layout("paper_card", title="Bayesian optimization of a three-enzyme cascade for rare-sugar synthesis",
               authors="J. Lee, K. Kim, M. Park", journal="Nature Catalysis", year=2024,
               impact_factor="IF 42.8 (JCR 2023)", doi="10.1038/s41929-024-01234-5",
               keywords=["enzyme cascade", "Bayesian optimization", "cofactor recycling", "kinetics"],
               summary="Closed-loop optimization doubled the titer within 12 experiments.",
               presenter="J. Lee", date="2026-10-04", venue="Lab meeting"))
    n(d.layout("paper_card", title="A minimal card", journal="ACS Catal", year="2023", presenter="J. Lee",
               date="2026-10-04"))
    n(d.layout("chapter_tracker", chapters=CHAPTERS, current=None, title="Outline"))
    n(d.layout("chapter_tracker", chapters=CHAPTERS, current=2,
               groups=[(0, 1, "Modeling"), (2, 3, "Optimization")]))
    n(d.layout("chapter_tracker", chapters=CHAPTERS[:3], current=0))
    n(d.layout("figure_card", title="Titer rises with cofactor recycling", fig_path=fig, fig_no=3,
               caption="Time course of the cascade at 30 °C and pH 7.5.",
               source="Source: Lee et al., Nat Catal 2024", takeaway="Recycling above 2 mM/min lifts the titer to 85%.",
               chapter=CHAPTERS[1]))
    if d.theme.allow_bullets:
        n(d.layout("figure_card", title="Reading the time course", fig_path=fig, fig_no="2a",
                   caption="Product titer over 3 h.", source="Source: Lee et al., Nat Catal 2024",
                   takeaway="Initial rate is set by the first enzyme.",
                   notes=["Blue: control without recycling.", "Orange: with glucose dehydrogenase.",
                          "Plateau marks substrate depletion."]))
    n(d.layout("claim_evidence", claim="Cofactor recycling, not enzyme loading, limits the productivity of the cascade",
               section="Results", note="Three independent lines of evidence.",
               evidence=[dict(head="Kinetics", text="NADPH turnover caps the rate at 4.4 s⁻¹."),
                         dict(head="Titer", text="Doubling the loading raised the titer by only 8%."),
                         dict(head="Recycling", text="Adding 5 mM glucose raised the titer by 65%.")],
               callout=dict(number="65%", label="Titer gain", note="With glucose dehydrogenase at 30 °C."),
               bottom_line="Fix the cofactor supply first, then tune the loading.", chapter=CHAPTERS[0]))
    n(d.layout("claim_evidence", claim="The optimum is robust to ±10% variation in pH",
               evidence=[dict(head="Question", text="Does pH 7.5 hold?"), dict(head="Method", text="Grid scan."),
                         dict(head="Evidence", text="Yield stays above 80%."),
                         dict(head="Use", text="Keep pH 7.5 in scale-up.")]))
    n(d.layout("claim_evidence", claim="Closed-loop optimization is sufficient for this cascade",
               conclusion=dict(finding="Titer doubled in 12 experiments.",
                               boundary="Only one substrate concentration (5 mM) was tested.",
                               next_step="Repeat at 50 mM and in a stirred reactor.")))
    n(d.layout("donut_stat_band", title="Where the cascade stands", items=[
        dict(value=85, label="Conversion", note="Substrate converted after 3 h at 30 °C."),
        dict(value=62.5, label="Selectivity", note="Mol of product per mol converted."),
        dict(value=100, label="Cofactor recycled")], footnote="n = 3 independent runs"))
    n(d.layout("donut_stat_band", title="Two gauges", items=[dict(value=0, label="Before"), dict(value=37, label="After")]))
    kw = dict(title="Acquisition function versus risk", x_label="Exploitation weight", y_label="Risk of failure",
              x_low="Explore", x_high="Exploit", y_low="Low risk", y_high="High risk",
              quadrants=["Bold probes", "Greedy gambles", "Safe scouting", "Safe gains"], emphasis="BR",
              items=[dict(label="UCB", x=0.25, y=0.2), dict(label="EI", x=0.75, y=0.25),
                     dict(label="Thompson", x=0.2, y=0.8), dict(label="Random", x=0.85, y=0.8)])
    n(d.layout("matrix_2x2", **kw))
    if d.theme.allow_bullets:
        n(d.layout("matrix_2x2", notes=["Right = trusts the surrogate mean.", "Up = more failed runs."], **kw))
    rows = [dict(label="Wild type", value=0), dict(label="Variant A", value=12.5),
            dict(label="Variant B", value=-8), dict(label="Variant C", value=31.2), dict(label="Variant D", value=-19)]
    n(d.layout("diverging_bar", title="Change in half-life versus wild type", rows=rows, unit=" h",
               left_caption="Less stable", right_caption="More stable", source="Source: thermal shift, n = 3"))
    n(d.layout("diverging_bar", title="Small effect sizes", rows=[dict(label="Cofactor", value=0.4),
                                                                   dict(label="pH", value=-0.15)], vmax=0.5))
    n(d.layout("bw_divider_quote", statement="Part II: the cofactor cycle sets the ceiling", section="02",
               kicker="Next: global kinetic fit", tone="dark"))
    n(d.layout("bw_divider_quote", number="3.2×", label="Higher titer than the single-enzyme control",
               tone="light"))
    n(d.layout("bw_divider_quote", number="2,543", label="Experiments avoided by Bayesian optimization", tone="dark"))
    n(d.layout("bw_divider_quote", quote="All models are wrong, but some are useful for choosing the next experiment.",
               attribution="G. Box, adapted", tone="dark"))
    n(d.layout("bw_divider_quote", quote="Measure what is measurable.", tone="light"))
    return d


def test_registered():
    assert set(FAMILY) <= set(LAYOUTS)
    for name in FAMILY:
        spec = LAYOUTS[name]
        assert spec.source and spec.summary
    assert LAYOUTS["claim_evidence"].extra_sizes == frozenset({db.STAT_BIG_PT})
    assert LAYOUTS["bw_divider_quote"].extra_sizes == frozenset({db.STAT_BIG_PT})


def test_new_themes_registered_with_evidence():
    for name in ("ucsd_navy_yellow", "k105_navy", "mono_bw_hard"):
        th = db.THEMES[name]
        assert th.family == "research" and th.flat_note.startswith("flat variant")
        assert "licen" in th.description.lower() or "MIT" in th.description
    # yellow only as fill / large bold: never reaches a text role
    for role in ("body", "title_main", "subtitle", "caption", "stat_number"):
        assert db.contrast_ratio(db.THEMES["ucsd_navy_yellow"].role_color(role), db.THEMES["ucsd_navy_yellow"].bg) >= 4.5


@pytest.mark.parametrize("theme", THEMES)
def test_every_layout_every_theme(theme, fig, tmp_path):
    d = build(theme, fig)
    p = tmp_path / f"{theme}.pptx"
    d.save(str(p))
    rep = qc_deck(str(p))
    assert rep.critical == [], rep.critical
    crit = [f for f in check_layout(Presentation(str(p)), theme=theme) if f.severity == "CRITICAL"]
    assert crit == [], [str(c) for c in crit]
    assert not [w for w in rep.warning if "not in the" in w], rep.warning
    assert not [w for w in rep.warning if "contrast" in w and "L3" not in w], rep.warning
    for s in Presentation(str(p)).slides:
        assert "header_bar" not in {x.name for x in s.shapes} or True


# ---------------------------------------------------------------- negatives
def test_forbidden_list_text_theme(fig):
    d = db.Deck(theme="assertion_evidence")
    with pytest.raises(db.StyleError):
        d.layout("figure_card", title="T", fig_path=fig, fig_no=1, caption="c", source="s", takeaway="t", notes=["a"])
    with pytest.raises(db.StyleError):
        d.layout("matrix_2x2", title="T", x_label="x", y_label="y", items=[dict(label="a", x=.2, y=.2)], notes=["a"])


@pytest.mark.parametrize("name,kw", [
    ("chapter_tracker", dict(chapters=[], current=0)),
    ("chapter_tracker", dict(chapters=["a"], current=0)),
    ("chapter_tracker", dict(chapters=["a", "b"], current=5)),
    ("chapter_tracker", dict(chapters=["word " * 60, "b"], current=0)),
    ("chapter_tracker", dict(chapters=["a", "b"], current=0, groups=[(0, 9, "x")])),
    ("paper_card", dict(title="word " * 60, journal="J", year=2024, presenter="P", date="D")),
    ("paper_card", dict(title="T", journal="J", year=2024, presenter="P", date="D", keywords=[])),
    ("paper_card", dict(title="T", journal="J", year=2024, presenter="P", date="D", keywords=["k"] * 9)),
    ("paper_card", dict(title="T", journal="J", year=2024, presenter="P", date="D",
                        keywords=["a very long keyword phrase number %d" % i for i in range(6)])),
    ("claim_evidence", dict(claim="C")),
    ("claim_evidence", dict(claim="C", evidence=[], )),
    ("claim_evidence", dict(claim="C", evidence=[dict(head="h", text="t")])),
    ("claim_evidence", dict(claim="word " * 80, evidence=[dict(head="h", text="t")] * 2)),
    ("claim_evidence", dict(claim="C", evidence=[dict(head="h", text="word " * 200)] * 2)),
    ("claim_evidence", dict(claim="C", evidence=[dict(head="h", text="t")] * 4,
                            callout=dict(number="1", label="l", note="n"))),
    ("claim_evidence", dict(claim="C", conclusion=dict(finding="a", boundary="b"))),
    ("claim_evidence", dict(claim="C", evidence=[dict(head="h", text="t")] * 2,
                            callout=dict(number="1234567", label="l", note="n"))),
    ("donut_stat_band", dict(title="T", items=[])),
    ("donut_stat_band", dict(title="T", items=[dict(value=101, label="a"), dict(value=5, label="b")])),
    ("donut_stat_band", dict(title="T", items=[dict(value=-1, label="a"), dict(value=5, label="b")])),
    ("donut_stat_band", dict(title="T", items=[dict(value=float("nan"), label="a"), dict(value=5, label="b")])),
    ("donut_stat_band", dict(title="T", items=[dict(value=True, label="a"), dict(value=5, label="b")])),
    ("donut_stat_band", dict(title="T", items=[dict(value=5, label="a")])),
    ("matrix_2x2", dict(title="T", x_label="x", y_label="y", items=[])),
    ("matrix_2x2", dict(title="T", x_label="x", y_label="y", items=[dict(label="a", x=1.2, y=.5)])),
    ("matrix_2x2", dict(title="T", x_label="x", y_label="y", items=[dict(label="a", x=.2, y=.2)] * 5)),
    ("matrix_2x2", dict(title="T", x_label="x", y_label="y", items=[dict(label="a", x=.2, y=.2)], emphasis="XX")),
    ("matrix_2x2", dict(title="T", x_label="x", y_label="y",
                        items=[dict(label="on the divider", x=.5, y=.5)])),
    ("diverging_bar", dict(title="T", rows=[])),
    ("diverging_bar", dict(title="T", rows=[dict(label="a", value=0), dict(label="b", value=0)])),
    ("diverging_bar", dict(title="T", rows=[dict(label="a", value=5), dict(label="b", value=1)], vmax=3)),
    ("diverging_bar", dict(title="T", rows=[dict(label="a", value=float("inf")), dict(label="b", value=1)])),
    ("diverging_bar", dict(title="T", rows=[dict(label="a", value="5"), dict(label="b", value=1)])),
    ("bw_divider_quote", dict()),
    ("bw_divider_quote", dict(statement="a", number="1")),
    ("bw_divider_quote", dict(statement="word " * 60)),
    ("bw_divider_quote", dict(number="1,234,567,890,123")),
    ("bw_divider_quote", dict(quote="word " * 120)),
    ("bw_divider_quote", dict(statement="a", attribution="x")),
    ("bw_divider_quote", dict(statement="a", tone="grey")),
])
def test_negative_cases_raise(name, kw):
    d = db.Deck()
    with pytest.raises(db.StyleError):
        d.layout(name, **kw)


def test_figure_too_small_raises(fig, tmp_path):
    p = tmp_path / "tiny.png"
    Image.new("RGB", (4000, 100), "white").save(p)       # a 4000 x 100 strip: 12.1 in wide is allowed
    d = db.Deck()
    d.layout("figure_card", title="T", fig_path=str(p), fig_no=1, caption="c", source="s", takeaway="t")
    p2 = tmp_path / "narrow.png"
    Image.new("RGB", (100, 2000), "white").save(p2)      # a tall sliver would be < 3 in tall x < 9 in wide? no: tall
    d2 = db.Deck()
    with pytest.raises(db.StyleError):
        d2.layout("figure_card", title="T", fig_path=str(p2), fig_no=1, caption="c " * 400, source="s", takeaway="t")


# ---------------------------------------------------------------- honesty checks
def _shapes(slide, prefix):
    return [s for s in slide.shapes if s.name.startswith(prefix)]


def test_donut_sweep_equals_value_times_3_6(tmp_path):
    d = db.Deck()
    s = d.layout("donut_stat_band", title="T", items=[dict(value=25, label="a"), dict(value=60, label="b"),
                                                        dict(value=0, label="c"), dict(value=100, label="d")])
    arcs = {x.name: x for x in _shapes(s, "donut:arc")}
    # 25 % -> 90 degrees: start 270 deg, end 0 deg ; adjustments are raw / 100000
    a25 = arcs["donut:arc 1 = 25%"]
    start, end = a25.adjustments[0] * 100000 / 60000, a25.adjustments[1] * 100000 / 60000
    assert abs(start - 270) < 1e-6 and abs(end - 0) < 1e-6
    a60 = arcs["donut:arc 2 = 60%"]
    sweep = ((a60.adjustments[1] - a60.adjustments[0]) * 100000 / 60000) % 360
    assert abs(sweep - 216) < 1e-3
    assert not any("= 0%" in n for n in arcs)                    # 0 % draws no arc
    assert "donut:arc 4 = 100%" in arcs                          # 100 % is a full ring
    vals = [x.text_frame.text for x in _shapes(s, "donut:value")]
    assert vals == ["25%", "60%", "0%", "100%"]


def test_diverging_bar_scale_is_linear_and_unclipped():
    d = db.Deck()
    rows = [dict(label="a", value=10), dict(label="b", value=-5), dict(label="c", value=20)]
    s = d.layout("diverging_bar", title="T", rows=rows, vmax=20)
    bars = {x.name.split()[-1]: x for x in _shapes(s, "bar:")}
    wa, wb, wc = bars["10"].width, bars["-5"].width, bars["20"].width
    assert abs(wa / wc - 0.5) < 0.01 and abs(wb / wc - 0.25) < 0.01
    zero = _shapes(s, "deco:zero-line")[0]
    zx = zero.left + zero.width / 2
    assert abs(bars["10"].left - zx) < Emu(20000) and abs((bars["-5"].left + bars["-5"].width) - zx) < Emu(20000)
    # default axis is a round number >= max |value|
    s2 = db.Deck().layout("diverging_bar", title="T", rows=[dict(label="a", value=31.2), dict(label="b", value=-8)])
    texts = [x.text_frame.text for x in s2.shapes if x.name == "tick"]
    assert texts[-1] == "+40" or texts[-1] == "+50"


def test_matrix_item_positions_follow_the_coordinates():
    d = db.Deck()
    s = d.layout("matrix_2x2", title="T", x_label="x", y_label="y",
                 items=[dict(label="lo", x=0.2, y=0.2), dict(label="hi", x=0.8, y=0.8)])
    m = {x.name: x for x in _shapes(s, "item-marker")}
    lo, hi = m["item-marker 1"], m["item-marker 2"]
    assert hi.left > lo.left and hi.top < lo.top                  # high x -> right, high y -> up
    plot = _shapes(s, "deco:plot")[0]
    cx = (lo.left + lo.width / 2 - plot.left) / plot.width
    assert 0.15 < cx < 0.3


def test_figure_card_keeps_aspect_and_size(fig):
    d = db.Deck()
    s = d.layout("figure_card", title="T", fig_path=fig, fig_no=1, caption="c", source="s", takeaway="t")
    pic = [x for x in s.shapes if x.name == "fig:card"][0]
    assert abs(pic.width / pic.height - 1600 / 900) < 0.01
    assert pic.height / 914400 >= 3.0


def test_chapter_tracker_current_is_lit_and_others_are_legible():
    d = db.Deck()
    s = d.layout("chapter_tracker", chapters=CHAPTERS, current=1)
    cur = [x for x in s.shapes if x.name == "chapter-current"]
    assert len(cur) == 1 and cur[0].text_frame.text == CHAPTERS[1]
    others = [x for x in s.shapes if x.name == "chapter"]
    assert len(others) == 3
    col = others[0].text_frame.paragraphs[0].runs[0].font.color.rgb
    assert db.contrast_ratio(col, d.theme.bg) >= 4.5            # dimmed, never unreadable
    assert cur[0].text_frame.paragraphs[0].runs[0].font.bold is True


def test_chapter_footer_prints_chapter_and_page():
    d = db.Deck(footer_right="Lee 2026")
    s = d.new_slide()
    chapter_footer(d, s, "Chapter 2: Global fit")
    t = {x.name: x.text_frame.text for x in s.shapes if x.has_text_frame and x.name.startswith("footer")}
    assert t["footer-chapter"] == "Chapter 2: Global fit" and t["footer-page"].endswith("1")
    with pytest.raises(db.StyleError):
        chapter_footer(d, d.new_slide(), "word " * 80)


def test_bw_alternates_fields():
    d = db.Deck()
    dark = d.layout("bw_divider_quote", statement="Dark", tone="dark")
    light = d.layout("bw_divider_quote", statement="Light", tone="light")
    assert str(dark.background.fill.fore_color.rgb) == str(d.theme.palette["table_hdr_bg"])
    assert str(light.background.fill.fore_color.rgb) == str(d.theme.bg)

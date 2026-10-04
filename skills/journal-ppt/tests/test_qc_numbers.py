"""Tests for scripts/qc_numbers.py (run: python -m pytest tests/test_qc_numbers.py)."""
import json
import subprocess
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
import qc_numbers as q  # noqa: E402

SOURCE = {
    "title": "Test paper",
    "abstract_summary": "The enzyme reached 85% yield at 5 mM substrate and pH 7.4.",
    "results": [
        {"finding": "Titer was 12.5 g/L after 24 h (p < 0.05, n = 3).",
         "explanation": "Activity rose 3.4-fold versus control.",
         "validated_numbers": [
             {"claim": "Yield 85%", "source": "Figure 2", "match": True, "value": 85, "unit": "%"},
             {"claim": "Substrate 5 mM", "source": "Methods", "match": True, "value": 5, "unit": "mM"},
             {"claim": "Conversion 72%", "source": "Table 1", "match": True, "value": 72, "unit": "%"},
             {"claim": "Km 40 mM", "source": "Table 1", "match": True, "value": 40, "unit": "mM"},
             {"claim": "Text says R = 0.73, caption says r = 0.72", "source": "Fig 3 caption",
              "match": False, "value": 0.73, "unit": "", "alt_value": 0.72},
         ]},
    ],
    "tables": [{"caption": "Table 1", "headers": ["Variant", "Km (mM)"], "rows": [["WT", 40], ["M1", "20.5"]]}],
    "validation_warnings": [],
}


def make_deck(path, slides, notes=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, texts in enumerate(slides):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        for j, t in enumerate(texts):
            tb = sl.shapes.add_textbox(Inches(0.5), Inches(0.5 + j), Inches(8), Inches(0.6))
            tb.text_frame.text = t
        if notes and notes.get(i + 1):
            sl.notes_slide.notes_text_frame.text = notes[i + 1]
    prs.save(str(path))
    return path


def run_deck(tmp_path, slides, source=SOURCE, notes=None):
    deck = make_deck(tmp_path / "d.pptx", slides, notes)
    src = tmp_path / "ca.json"
    src.write_text(json.dumps(source), encoding="utf-8")
    code, ledger = q.run(deck, src, tmp_path / "ledger.json")
    return code, ledger


def status_of(ledger, raw_prefix):
    rows = [r for r in ledger["claims"] if r["raw"].startswith(raw_prefix)]
    assert rows, f"{raw_prefix!r} not extracted; got {[r['raw'] for r in ledger['claims']]}"
    return rows[0]["status"]


# ---- N1 extraction ---------------------------------------------------------

def nums(text):
    return [(c.value, c.unit.canon, c.kind) for c in q.extract_numbers(text)[0]]


def test_extract_percent_spacing_equivalent():
    assert nums("85 %") == nums("85%") == [(85.0, "%", "pct")]


def test_extract_fraction_distinct_from_percent():
    assert nums("0.85") == [(0.85, "", "value")]
    assert nums("0.85") != nums("85%")


def test_extract_units_micro_variants():
    assert nums("5 µM") == nums("5 μM") == nums("5 uM") == [(5.0, "uM", "value")]


def test_extract_pvalue_n_fold_pm():
    assert (0.05, "p", "p") in nums("p < 0.05")
    assert (3.0, "n", "n") in nums("n = 3")
    assert (3.4, "fold", "fold") in nums("3.4-fold increase")
    assert (4.0, "fold", "fold") in nums("a 4x gain")
    r = nums("5.2 ± 0.3 mM")
    assert (5.2, "mM", "value") in r and (0.3, "mM", "pm_err") in r


def test_extract_range_shares_unit():
    r = nums("10–20 mg/mL")
    assert (10.0, "mg/mL", "range_lo") in r and (20.0, "mg/mL", "range_hi") in r


def test_extract_scientific_notation_and_unicode_minus_thin_space():
    assert nums("1.2 × 10⁻³ M")[0][:2] == (pytest.approx(1.2e-3), "M")
    assert nums("3e-4 M")[0][:2] == (pytest.approx(3e-4), "M")
    assert nums("−20 °C") == [(-20.0, "°C", "value")]
    assert nums("85 %") == [(85.0, "%", "pct")]
    assert nums("1 000 µM")[0][:2] == (1000.0, "uM")


def test_extract_context_window():
    c = q.extract_numbers("x" * 60 + " 85% " + "y" * 60)[0][0]
    assert "85%" in c.context and len(c.context) <= 3 + 2 * 40 + 2


def test_notes_are_not_scanned(tmp_path):
    deck = make_deck(tmp_path / "n.pptx", [["Yield 85%"]], {1: "Notes say 99% and 1234 mM"})
    claims = q.extract_slide_numbers(deck)[0]
    assert [c.value for c in claims] == [85.0]


# ---- trivial-number rules ----------------------------------------------------

@pytest.mark.parametrize("text,rule", [
    ("Figure 2 and Fig. 3a show", "figure_table_ref"),
    ("see Table 1", "figure_table_ref"),
    ("Figures 2 and 3", "figure_table_ref"),
    ("doi:10.1021/jacs.3c01234", "doi"),
    ("https://example.org/a/12345", "url"),
    ("Smith et al., 2023", "et_al_year"),
    ("(Lee, 2021; Park, 2022)", "citation_year_paren"),
    ("as shown [12, 13]", "reference_index"),
    ("pp. 123–130", "page_range"),
    ("Nat. Commun. 14(3):456-789", "journal_citation"),
    ("2026-08-28", "iso_date"),
    ("Vol. 12", "volume_issue"),
    ("the 1:10 dilution", "time_or_ratio"),
    ("in 2024 we", "year"),
    ("three of 7 samples", "small_bare_integer"),
    ("ZIF-8 and Cas9 in BL21", "identifier"),
    ("a 96-well plate", "hyphenated_modifier"),
])
def test_trivial_numbers_are_ignored_with_rule(text, rule):
    claims, ignored = q.extract_numbers(text)
    assert claims == [], f"{text!r} produced {[c.raw for c in claims]}"
    assert rule in {i["rule"] for i in ignored}


def test_slide_number_shapes_ignored(tmp_path):
    deck = make_deck(tmp_path / "s.pptx", [["Title"], ["Body 85%", "2"], ["Body", "3 / 3"]])
    claims, ignored = q.extract_slide_numbers(deck)[:2]
    assert [c.raw for c in claims] == ["85%"]
    assert sum(1 for i in ignored if i["rule"] == "slide_number") == 2


# ---- N3 classification --------------------------------------------------------

def test_clean_deck_is_exit_0(tmp_path):
    code, led = run_deck(tmp_path, [["Yield 85% at 5 mM", "p < 0.05, n = 3"], ["Titer 12.5 g/L, 3.4-fold"]])
    assert code == 0, q.format_report(led)
    assert set(r["status"] for r in led["claims"]) == {"MATCH"}


def test_match_variants(tmp_path):
    code, led = run_deck(tmp_path, [["Substrate 5000 µM", "Titer 13 g/L"]])
    st = {r["raw"]: r["status"] for r in led["claims"]}
    assert st["5000 µM"] == "MATCH"            # unit conversion
    assert st["13 g/L"] == "MATCH"             # honest rounding of 12.5 (half up)
    assert code == 0


@pytest.mark.parametrize("slide_text,raw", [
    ("Yield 95%", "95%"),            # one-digit change
    ("Yield 58%", "58%"),            # transposition
    ("Substrate 5 µM", "5 µM"),      # unit mismatch
    ("Yield 85.0%", "85.0%"),        # sig-fig increase
    ("Yield 0.85", "0.85"),          # fraction vs percent
    ("Yield 8.5%", "8.5%"),          # decimal shift
    ("Conversion 76%", "76%"),       # within 10% of 72
    ("Km 4 mM", "4 mM"),             # decimal shift of 40 mM
])
def test_drift(tmp_path, slide_text, raw):
    code, led = run_deck(tmp_path, [[slide_text]])
    assert status_of(led, raw) == "DRIFT"
    assert code == 1


def test_unsourced_is_warning_not_critical(tmp_path):
    code, led = run_deck(tmp_path, [["Cost was 777 USD"], ["Yield 85%"]])
    assert status_of(led, "777") == "UNSOURCED"
    assert code == 0
    assert q.run(tmp_path / "d.pptx", tmp_path / "ca.json", strict=True)[0] == 1


def test_derived_unlabeled_and_labeled(tmp_path):
    code, led = run_deck(tmp_path, [["Fold change 1.2-fold"]])      # 85/72
    assert status_of(led, "1.2") == "DERIVED-UNLABELED"
    assert code == 0
    code, led = run_deck(tmp_path, [["Fold change ~1.2-fold"]])
    assert [r["status"] for r in led["claims"]] == ["DERIVED-LABELED"]
    code, led = run_deck(tmp_path, [["Fold change 1.2-fold"]], notes={1: "1.2-fold is our calculation (85/72)"})
    assert [r["status"] for r in led["claims"]] == ["DERIVED-LABELED"]


def test_derived_difference_and_percent_change(tmp_path):
    _, led = run_deck(tmp_path, [["Gain of 13%"]])      # 85 - 72
    assert status_of(led, "13") == "DERIVED-UNLABELED"
    _, led = run_deck(tmp_path, [["Km dropped 49%"]])    # 40 -> 20.5 percent change
    assert status_of(led, "49") == "DERIVED-UNLABELED"


def test_source_conflict_without_caveat_is_critical(tmp_path):
    code, led = run_deck(tmp_path, [["Correlation R = 0.73"]])
    assert status_of(led, "0.73") == "SOURCE-CONFLICT"
    assert code == 1


def test_source_conflict_with_visible_caveat_passes(tmp_path):
    code, led = run_deck(tmp_path, [["Correlation R = 0.73", "Caveat: caption reports 0.72"]])
    assert status_of(led, "0.73") == "MATCH"


def test_source_conflict_from_validation_warnings(tmp_path):
    src = json.loads(json.dumps(SOURCE))
    src["validation_warnings"] = ["Abstract says 91% but Results report 89% (mismatch)"]
    code, led = run_deck(tmp_path, [["Selectivity 91%"]], src)
    assert status_of(led, "91%") == "SOURCE-CONFLICT" and code == 1


def test_p_value_operator_change_is_drift(tmp_path):
    code, led = run_deck(tmp_path, [["p = 0.05"]])
    assert status_of(led, "0.05") == "DRIFT"


# ---- N2 schema -----------------------------------------------------------------

def test_schema_errors_are_helpful():
    bad = {"validated_numbers": [{"claim": "x", "match": "yes", "value": "85", "unit": "furlong"}]}
    joined = "\n".join(q.validate_source(bad))
    assert "validated_numbers[0].match" in joined and "validated_numbers[0].value" in joined \
        and "furlong" in joined
    assert q.validate_source([]) and q.validate_source({"title": "x"})


def test_bad_schema_exit_2(tmp_path):
    deck = make_deck(tmp_path / "d.pptx", [["Yield 85%"]])
    src = tmp_path / "ca.json"
    src.write_text(json.dumps({"validated_numbers": "nope"}), encoding="utf-8")
    code, led = q.run(deck, src)
    assert code == 2 and led["verdict"] == "BLIND" and "validated_numbers" in led["error"]


# ---- N4 exit codes / ledger -------------------------------------------------------

def test_blind_no_source_numbers(tmp_path):
    code, led = run_deck(tmp_path, [["Yield 85%"]], {"abstract_summary": "No numbers here."})
    assert code == 2 and led["verdict"] == "BLIND"


def test_blind_no_slide_numbers(tmp_path):
    code, led = run_deck(tmp_path, [["Only words here", "2"]])
    assert code == 2 and led["verdict"] == "BLIND"


def test_ledger_contents_and_cli_exit_codes(tmp_path):
    deck = make_deck(tmp_path / "d.pptx", [["Yield 95%", "Fig. 2, p < 0.05"]])
    src = tmp_path / "ca.json"
    src.write_text(json.dumps(SOURCE), encoding="utf-8")
    led = tmp_path / "out.json"
    res = subprocess.run([sys.executable, str(SCRIPTS / "qc_numbers.py"), str(deck), "--source", str(src),
                          "--ledger", str(led)], capture_output=True, text=True, encoding="utf-8")
    assert res.returncode == 1 and "DRIFT" in res.stdout
    data = json.loads(led.read_text(encoding="utf-8"))
    row = next(r for r in data["claims"] if r["raw"] == "95%")
    for key in ("slide", "value", "unit", "context", "status", "source_ref", "nearest_source", "severity"):
        assert key in row
    assert row["slide"] == 1 and row["nearest_source"]["value"] == 85
    assert any(i["rule"] == "figure_table_ref" for i in data["ignored"])
    res = subprocess.run([sys.executable, str(SCRIPTS / "qc_numbers.py"), str(deck), "--source",
                          str(tmp_path / "missing.json")], capture_output=True, text=True, encoding="utf-8")
    assert res.returncode == 2 and "BLIND" in res.stdout


def test_integration_with_deck_builder(tmp_path):
    from deck_builder import Deck
    d = Deck()
    d.title_slide("T", "A et al.", "J 2026", presenter="P", date="2026-08-28")
    s = d.content_slide("Results")
    d.bullets(s, ["Yield 85% at 5 mM", "Titer 12.5 g/L"], role="body")
    d.notes(s, "ok")
    out = tmp_path / "b.pptx"
    d.save(str(out))
    src = tmp_path / "ca.json"
    src.write_text(json.dumps(SOURCE), encoding="utf-8")
    code, led = q.run(out, src)
    assert code == 0, q.format_report(led)
